"""Build a .pptx from a deck plan (content.py deck_plan) with python-pptx, the second producer.

What an agent writes with python-pptx's public API on its default template (4:3, Office theme):
the template's layouts and placeholders, native charts from CategoryChartData, tables, pictures,
shapes for stat callouts, speaker notes on every slide. python-pptx has no API for sections, so
this deck has none. For reproducibility only, the zip entries get a fixed timestamp after saving
(python-pptx stamps them with the current time), and so do the chart workbooks' core.xml.

    .cache/venv/bin/python gen_pptx_pythonpptx.py plan.json out.pptx
"""

from __future__ import annotations

import datetime as dt
import io
import json
import sys
import zipfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ACCENT = RGBColor(0x0B, 0x7A, 0x75)
DARK = RGBColor(0x1B, 0x2A, 0x3A)
LIGHT = RGBColor(0xEE, 0xF3, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FIXED = dt.datetime(2026, 1, 15, 9, 0, 0)

# python-pptx writes each chart's data as an embedded workbook with XlsxWriter, which stamps
# docProps/core.xml with the current time: give it the fixed one (reproducibility only).
import xlsxwriter.workbook as _xw  # noqa: E402

_xw_init = _xw.Workbook.__init__


def _xw_fixed_init(self, *args, **kwargs):
    _xw_init(self, *args, **kwargs)
    self.set_properties({"created": FIXED})


_xw.Workbook.__init__ = _xw_fixed_init


def fill_bullets(tf, items, size=None, head=None):
    tf.clear()
    first = True
    if head:
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = head
        r.font.bold = True
        r.font.color.rgb = ACCENT
        first = False
    for t in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = t
        if size:
            r.font.size = Pt(size)


def remove_placeholder(ph):
    el = ph._element
    el.getparent().remove(el)


def build(plan: dict, out: str) -> None:
    prs = Presentation()
    L_TITLE, L_CONTENT, L_SECTION, L_TWO, L_TITLE_ONLY = (prs.slide_layouts[i] for i in (0, 1, 2, 3, 5))
    cp = prs.core_properties
    cp.title = plan["meta"]["title"]
    cp.author = "Corporate Affairs"
    cp.last_modified_by = "Corporate Affairs"
    cp.created = cp.modified = FIXED
    cp.revision = 1

    for s in plan["slides"]:
        t = s["type"]
        if t == "title":
            slide = prs.slides.add_slide(L_TITLE)
            slide.shapes.title.text = s["title"]
            slide.placeholders[1].text = s["subtitle"]
        elif t == "section":
            slide = prs.slides.add_slide(L_SECTION)
            slide.shapes.title.text = s["title"]
            slide.placeholders[1].text = s["subtitle"]
        elif t in ("bullets", "closing"):
            slide = prs.slides.add_slide(L_CONTENT)
            slide.shapes.title.text = s["title"]
            fill_bullets(slide.placeholders[1].text_frame, s["bullets"])
        elif t == "long_text":
            slide = prs.slides.add_slide(L_CONTENT)
            slide.shapes.title.text = s["title"]
            fill_bullets(slide.placeholders[1].text_frame, s["text"], size=14)
        elif t == "bullets_pic":
            slide = prs.slides.add_slide(L_TWO)
            slide.shapes.title.text = s["title"]
            fill_bullets(slide.placeholders[1].text_frame, s["bullets"], size=16)
            right = slide.placeholders[2]
            x, y, w = right.left, right.top, right.width
            remove_placeholder(right)
            im = s["image"]
            slide.shapes.add_picture(im["path"], x, y, width=w, height=int(w * im["h"] / im["w"]))
        elif t == "two_col":
            slide = prs.slides.add_slide(L_TWO)
            slide.shapes.title.text = s["title"]
            fill_bullets(slide.placeholders[1].text_frame, s["bullets"], size=15, head=s["left_head"])
            fill_bullets(slide.placeholders[2].text_frame, s["right"], size=15, head=s["right_head"])
        elif t == "chart":
            slide = prs.slides.add_slide(L_TITLE_ONLY)
            slide.shapes.title.text = s["title"]
            spec = s["chart"]
            data = CategoryChartData()
            data.categories = spec["series"][0]["labels"]
            for ser in spec["series"]:
                data.add_series(ser["name"], ser["values"])
            kind = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}[spec["type"]]
            chart = slide.shapes.add_chart(kind, Inches(0.4), Inches(1.6), Inches(6.2), Inches(5.0), data).chart
            chart.has_title = True
            chart.chart_title.text_frame.text = spec["title"]
            chart.has_legend = len(spec["series"]) > 1 or spec["type"] == "pie"
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.include_in_layout = False
            plot = chart.plots[0]
            if spec["type"] == "bar":
                plot.has_data_labels = True
                plot.data_labels.font.size = Pt(9)
                plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
            elif spec["type"] == "pie":
                plot.has_data_labels = True
                plot.data_labels.show_percentage = True
                plot.data_labels.show_value = False
                plot.data_labels.number_format = "0%"
                plot.data_labels.number_format_is_linked = False
            box = slide.shapes.add_textbox(Inches(6.8), Inches(1.8), Inches(2.8), Inches(4.8))
            box.text_frame.word_wrap = True
            fill_bullets(box.text_frame, s["commentary"], size=14)
        elif t == "table":
            slide = prs.slides.add_slide(L_TITLE_ONLY)
            slide.shapes.title.text = s["title"]
            rows, cols = len(s["table"]), len(s["table"][0])
            table = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(1.7), Inches(9.0), Inches(0.4) * rows).table
            for i, row in enumerate(s["table"]):
                for j, text in enumerate(row):
                    cell = table.cell(i, j)
                    cell.text = text
                    para = cell.text_frame.paragraphs[0]
                    para.font.size = Pt(14)
                    para.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
        elif t == "stats":
            slide = prs.slides.add_slide(L_TITLE_ONLY)
            slide.shapes.title.text = s["title"]
            for k, st in enumerate(s["stats"]):
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5 + k * 3.1), Inches(1.8), Inches(2.8), Inches(2.9))
                card.fill.solid()
                card.fill.fore_color.rgb = LIGHT
                card.line.fill.background()
                tf = card.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = st["value"]
                r.font.size, r.font.bold, r.font.color.rgb = Pt(44), True, ACCENT
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.CENTER
                r2 = p2.add_run()
                r2.text = st["label"]
                r2.font.size, r2.font.color.rgb = Pt(14), DARK
            cap = slide.shapes.add_textbox(Inches(0.5), Inches(5.1), Inches(9.0), Inches(1.2))
            cap.text_frame.word_wrap = True
            r = cap.text_frame.paragraphs[0].add_run()
            r.text = s["caption"]
            r.font.size, r.font.italic = Pt(16), True
        else:
            raise ValueError(f"unknown slide type {t}")
        slide.notes_slide.notes_text_frame.text = s["notes"]

    buf = io.BytesIO()
    prs.save(buf)
    # Same entries, same order, same compression; only the zip timestamps are fixed.
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=FIXED.timetuple()[:6])
            zi.compress_type = info.compress_type
            zi.external_attr = info.external_attr
            dst.writestr(zi, src.read(info.filename))


if __name__ == "__main__":
    build(json.loads(open(sys.argv[1], encoding="utf-8").read()), sys.argv[2])
