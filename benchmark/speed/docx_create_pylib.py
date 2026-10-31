"""Create the corpus report from its plan with python-docx (the pylib stack's creation path).

    .cache/venv/bin/python docx_create_pylib.py plan.json out.docx

The same content and features as corpus/gen_docx.js where python-docx has an API; lxml (python-docx's own
OxmlElement / parse_xml) where it has none, as an agent writes it: list restarts (a new w:num), cell shading,
coloured table borders, cell margins, repeated header rows, picture alt text, footnotes (a new part), tracked
insertions and deletions, the TOC field, PAGE / NUMPAGES fields, comment dates.
"""
import json
import sys
from datetime import datetime, timezone
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Emu, Pt, RGBColor, Twips

plan_path, out = sys.argv[1:3]
plan = json.load(open(plan_path, encoding="utf-8"))
HEAD_FILL, ZEBRA, TOTAL_FILL = "1F4E79", "EAF1F8", "D9E2F3"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Twips(11906), Twips(16838)
sec.top_margin = sec.bottom_margin = Twips(1440)
sec.left_margin = sec.right_margin = Twips(1300)
styles = doc.styles


def style(name, font=None, size=None, bold=None, italic=None, color=None, before=None, after=None, keep=None):
    try:
        s = styles[name]
    except KeyError:
        s = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = styles["Normal"]
    if font:
        s.font.name = font
    if size:
        s.font.size = Pt(size)
    if bold is not None:
        s.font.bold = bold
    if italic is not None:
        s.font.italic = italic
    if color:
        s.font.color.rgb = RGBColor.from_string(color)
    if before is not None:
        s.paragraph_format.space_before = Twips(before)
    if after is not None:
        s.paragraph_format.space_after = Twips(after)
    if keep is not None:
        s.paragraph_format.keep_with_next = keep
    return s


style("Normal", font="Calibri", size=11, after=120).paragraph_format.line_spacing = 1.1
style("Title", font="Cambria", size=28, color=HEAD_FILL)
style("Heading 1", font="Cambria", size=16, bold=True, color=HEAD_FILL, before=360, after=120, keep=True)
style("Heading 2", font="Cambria", size=13, bold=True, color="2E74B5", before=240, after=80, keep=True)
style("Heading 3", font="Calibri", size=11, bold=True, italic=True, color="404040", before=200, after=60, keep=True)
style("Subtitle", font="Cambria", size=16, color="595959", after=480)
style("Caption", size=9, italic=True, color="404040", before=60, after=160, keep=True)
style("TOC Heading", font="Cambria", size=16, bold=True, color=HEAD_FILL, after=240)

# Header and "Page X of Y" footer
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hp.add_run(plan["meta"]["header"])
r.font.size, r.font.color.rgb = Pt(8), RGBColor.from_string("595959")


def add_field(paragraph, instruction, shown="1", size=None, dirty=False):
    for kind, text in (("begin", None), (None, instruction), ("separate", None), (None, shown), ("end", None)):
        run = paragraph.add_run()
        if size:
            run.font.size = Pt(size)
        if kind:
            char = OxmlElement("w:fldChar")
            char.set(qn("w:fldCharType"), kind)
            if dirty and kind == "begin":  # Word rebuilds the field when it opens the file
                char.set(qn("w:dirty"), "true")
            run._r.append(char)
        elif text == instruction:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = instruction
            run._r.append(instr)
        else:
            run.text = text


fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
fp.add_run("Page ").font.size = Pt(9)
add_field(fp, "PAGE", size=9)
fp.add_run(" of ").font.size = Pt(9)
add_field(fp, "NUMPAGES", size=9)

# Footnotes: python-docx has no API, so the part is written as Word writes it
notes = "".join(
    f'<w:footnote w:id="{fid}"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr><w:r><w:rPr>'
    f'<w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r><w:r><w:t xml:space="preserve"> {escape(text)}</w:t>'
    f'</w:r></w:p></w:footnote>' for fid, text in plan["footnotes"].items())
NOTES = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:footnotes {nsdecls("w")}>'
         '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
         '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p>'
         f'</w:footnote>{notes}</w:footnotes>')
if plan["footnotes"]:
    part = Part(PackURI("/word/footnotes.xml"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml", NOTES.encode(),
                doc.part.package)
    doc.part.relate_to(part, RT.FOOTNOTES)

# Lists: List Bullet / List Bullet 2 / List Number; each numbered list restarts on its own w:num
numbering = doc.part.numbering_part.element


def list_num_id(style_name):
    return styles[style_name].element.pPr.numPr.numId.val


instances = {}


def restart(p, style_name, key):
    if key not in instances:
        abstract = numbering.num_having_numId(list_num_id(style_name)).abstractNumId.val
        num = numbering.add_num(abstract)
        override = parse_xml(f'<w:lvlOverride {nsdecls("w")} w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride>')
        num.append(override)
        instances[key] = num.numId
    numpr = p._p.get_or_add_pPr().get_or_add_numPr()
    numpr.get_or_add_ilvl().val = 0
    numpr.get_or_add_numId().val = instances[key]


comments = {c["id"]: c for c in plan["comments"]}


def paragraph(inl, style_name=None):
    p = doc.add_paragraph(style=style_name)
    opened = {}
    for x in inl:
        k = x["k"]
        if k == "r":
            run = p.add_run(x["text"])
            run.bold, run.italic = bool(x.get("b")) or None, bool(x.get("i")) or None
            for cid in list(opened):
                opened[cid].append(run)
        elif k in ("ins", "del"):
            tag = "w:t" if k == "ins" else "w:delText"
            rpr = ("<w:rPr>" + ("<w:b/>" if x.get("b") else "") + ("<w:i/>" if x.get("i") else "") + "</w:rPr>")
            p._p.append(parse_xml(
                f'<w:{k} {nsdecls("w")} w:id="{x["id"]}" w:author="{escape(x["author"])}" w:date="{x["date"]}">'
                f'<w:r>{rpr}<{tag} xml:space="preserve">{escape(x["text"])}</{tag}></w:r></w:{k}>'))
        elif k == "fn":
            p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
                                  f'<w:footnoteReference w:id="{x["id"]}"/></w:r>'))
        elif k == "cs":
            opened[x["id"]] = []
        elif k == "ce":
            runs = opened.pop(x["id"])
            c = comments[x["id"]]
            if runs:  # python-docx anchors a comment on whole runs only
                comment = doc.add_comment(runs, text=c["text"], author=c["author"], initials=c["initials"])
                comment._comment_elm.date = datetime.fromisoformat(c["date"].replace("Z", "+00:00"))
    return p


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().insert_element_before(shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                                                     "w:vAlign", "w:hideMark")


BORDERS = (f'<w:tblBorders {nsdecls("w")}>' + "".join(
    f'<w:{e} w:val="single" w:sz="4" w:space="0" w:color="8EAADB"/>'
    for e in ("top", "left", "bottom", "right", "insideH", "insideV")) + "</w:tblBorders>")
MARGINS = (f'<w:tblCellMar {nsdecls("w")}><w:top w:w="40" w:type="dxa"/><w:left w:w="80" w:type="dxa"/>'
           '<w:bottom w:w="40" w:type="dxa"/><w:right w:w="80" w:type="dxa"/></w:tblCellMar>')


def table(spec):
    widths, ncols, nrows = spec["widths"], len(spec["widths"]), len(spec["rows"])
    t = doc.add_table(rows=nrows, cols=ncols)
    tblpr = t._tbl.tblPr
    tblpr.insert_element_before(parse_xml(BORDERS), "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    tblpr.insert_element_before(parse_xml(MARGINS), "w:tblLook")
    for c, w in enumerate(widths):
        t.columns[c].width = Twips(w)
        for cell in t.columns[c].cells:
            cell.width = Twips(w)
    grid = [[None] * ncols for _ in range(nrows)]
    merges = []
    carry = [0] * ncols
    for r, row in enumerate(spec["rows"]):
        col = 0
        for c in row["cells"]:
            while col < ncols and carry[col] > 0:
                carry[col] -= 1
                col += 1
            span = c.get("colSpan") or 1
            grid[r][col] = c
            rows_down = c.get("rowSpan") or 1
            if span > 1 or rows_down > 1:
                merges.append((r, col, r + rows_down - 1, col + span - 1))
            if c.get("rowSpan"):
                for k in range(col, col + span):
                    carry[k] = c["rowSpan"] - 1
            col += span
        while col < ncols and carry[col] > 0:
            carry[col] -= 1
            col += 1
    for r0, c0, r1, c1 in merges:  # merge first: a merged cell is then formatted once
        t.cell(r0, c0).merge(t.cell(r1, c1))
    seen = set()
    for r, row in enumerate(spec["rows"]):
        cells = t.rows[r].cells
        fill = HEAD_FILL if row.get("header") else TOTAL_FILL if row.get("total") else ZEBRA if row.get("shade") else None
        if row.get("header"):
            t.rows[r]._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for c in range(ncols):
            cell = cells[c]
            if cell._tc in seen:  # python-docx repeats a merged cell in every grid column it covers
                continue
            seen.add(cell._tc)
            if fill:
                shade(cell, fill)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            spec_cell = grid[r][c]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if spec_cell is None:
                continue
            p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if row.get("header") else
                           WD_ALIGN_PARAGRAPH.RIGHT if spec_cell.get("align") == "right" else WD_ALIGN_PARAGRAPH.LEFT)
            run = p.add_run(spec_cell["text"])
            run.font.size = Pt(9)
            if row.get("header") or row.get("total"):
                run.bold = True
            if row.get("header"):
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


for s, section in enumerate(plan["sections"]):
    if s > 0:
        new = doc.add_section(WD_SECTION.NEW_PAGE)
        if section.get("orientation") == "landscape":
            new.orientation = WD_ORIENT.LANDSCAPE
            new.page_width, new.page_height = new.page_height, new.page_width
    for b in section["blocks"]:
        t = b["t"]
        if t == "title":
            doc.add_heading(b["text"], 0).paragraph_format.space_before = Twips(2400)
        elif t == "subtitle":
            doc.add_paragraph(b["text"], style="Subtitle")
        elif t == "pagebreak":
            doc.add_page_break()
        elif t == "toc":
            doc.add_paragraph("Contents", style="TOC Heading")
            add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "Update the field to build the table of contents.",
                      dirty=True)
        elif t == "h":
            doc.add_heading(b["text"], b["level"])
        elif t == "p":
            paragraph(b["inl"])
        elif t == "li":
            name = ("List Number" if b["list"] == "num" else "List Bullet" if b["level"] == 0 else "List Bullet 2")
            p = paragraph(b["inl"], name)
            if b["list"] == "num":
                restart(p, name, (b["list"], b["instance"]))
        elif t == "table":
            table(b)
        elif t == "caption":
            doc.add_paragraph(b["text"], style="Caption")
        elif t == "img":
            doc.add_picture(b["path"], width=Emu(b["w"] * 9525), height=Emu(b["h"] * 9525))
            p = doc.paragraphs[-1]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            docpr = p._p.xpath(".//wp:docPr")[0]  # alt text: no API
            docpr.set("descr", b["alt"])
            docpr.set("title", b["alt"])
        else:
            raise ValueError(f"unknown block {t}")

cp = doc.core_properties
cp.title, cp.author, cp.last_modified_by = plan["meta"]["title"], plan["meta"]["creator"], plan["meta"]["creator"]
cp.comments = "Synthetic benchmark document (fictional organisation)."
cp.created = cp.modified = datetime(2026, 1, 15, 9, 0, tzinfo=timezone.utc)
doc.save(out)
