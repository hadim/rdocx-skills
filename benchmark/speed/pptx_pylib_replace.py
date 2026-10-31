"""Counted replacement with python-pptx (borrowed from features/pptx P03): slides (groups, tables) and notes,
across runs; refuses, writing nothing, when the count differs.
    python pptx_pylib_replace.py IN OUT OLD NEW COUNT"""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

src, out, old, new, expect = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4], int(sys.argv[5])


def frames(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from frames(sh.shapes)
        elif getattr(sh, "has_table", False) and sh.has_table:
            yield from (cell.text_frame for row in sh.table.rows for cell in row.cells)
        elif sh.has_text_frame:
            yield sh.text_frame


def replace(p):
    runs = p.runs
    chars = [(k, c) for k, r in enumerate(runs) for c in r.text]
    text = "".join(c for _, c in chars)
    n, i, kept = 0, 0, []
    while i < len(chars):
        if text.startswith(old, i):
            kept += [(chars[i][0], c) for c in new]
            i, n = i + len(old), n + 1
        else:
            kept.append(chars[i])
            i += 1
    if n:
        for k, r in enumerate(runs):
            r.text = "".join(c for j, c in kept if j == k)
    return n


prs = Presentation(src)
count = 0
for slide in prs.slides:
    tfs = list(frames(slide.shapes))
    if slide.has_notes_slide:
        tfs.append(slide.notes_slide.notes_text_frame)
    count += sum(replace(p) for tf in tfs for p in tf.paragraphs)
if count != expect:
    sys.exit(f"refused: {count} occurrences of {old!r}, expected {expect}; nothing written")
prs.save(out)
