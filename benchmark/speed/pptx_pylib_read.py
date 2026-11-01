"""Slide text and speaker notes of a deck with python-pptx (groups and tables included), printed slide by slide."""
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def texts(shapes):
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from texts(sh.shapes)
        elif sh.has_text_frame and sh.text_frame.text:
            yield sh.text_frame.text
        elif getattr(sh, "has_table", False) and sh.has_table:
            yield from (cell.text for row in sh.table.rows for cell in row.cells)


prs = Presentation(sys.argv[1])
out = []
for k, slide in enumerate(prs.slides, 1):
    out.append(f"--- slide {k}")
    out.extend(texts(slide.shapes))
    if slide.has_notes_slide:
        out.append("Notes: " + slide.notes_slide.notes_text_frame.text)
sys.stdout.write("\n".join(out) + "\n")
