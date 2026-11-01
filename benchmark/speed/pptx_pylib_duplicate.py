"""Duplicate slide N (one-based) right after itself with python-pptx, which has no API for it: the usual recipe.
A new slide on the same layout, its placeholders removed, the source's shapes deep-copied, the source's
relationships (pictures, charts, media; not the notes) added to the copy with their ids remapped, the speaker
notes' text copied, then the copy moved after the source in <p:sldIdLst>.
    python pptx_pylib_duplicate.py IN OUT N"""
import copy
import sys

from pptx import Presentation

src, out, n = sys.argv[1], sys.argv[2], int(sys.argv[3])
prs = Presentation(src)
source = prs.slides[n - 1]
dup = prs.slides.add_slide(source.slide_layout)
for shape in list(dup.shapes):
    shape._element.getparent().remove(shape._element)
rid_map = {}
for rid, rel in source.part.rels.items():
    if rel.reltype.endswith("/notesSlide") or rel.reltype.endswith("/slideLayout"):
        continue
    rid_map[rid] = dup.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref) if rel.is_external \
        else dup.part.relate_to(rel.target_part, rel.reltype)
tree = dup.shapes._spTree
for el in source.shapes._spTree.iterchildren():
    if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
        continue
    new = copy.deepcopy(el)
    for node in new.iter():
        for key, value in node.attrib.items():
            if key.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}") and value in rid_map:
                node.set(key, rid_map[value])
    tree.append(new)
if source.has_notes_slide:
    dup.notes_slide.notes_text_frame.text = source.notes_slide.notes_text_frame.text
ids = prs.slides._sldIdLst
entry = ids[-1]
ids.remove(entry)
ids.insert(n, entry)
prs.save(out)
