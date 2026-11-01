"""rdocx / rpptx rendering findings of the visual section, each reduced to a neutral reproduction.

Every case builds its own input (python-docx, python-pptx, pptxgenjs, zlib; no corpus file), renders it with the
pinned rdocx / rpptx build and with LibreOffice (the reference that runs here; Word and PowerPoint do not), and
reads the PDFs with Poppler. ``reproduced`` is True when rdocx / rpptx still shows the fault: the day a pin fixes
one, this says so. The same scripts are in NOTES.md.
"""

from __future__ import annotations

import json
import subprocess
from pathlib import Path

import numpy as np
from PIL import Image

import common

WORK = common.WORK / "visual" / "findings"

BUILD = r'''
import struct, sys, zlib, re, zipfile
import docx, pptx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


def chunk(t, d):
    return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)


def png(path, interlaced, w=64, h=48, rgb=(200, 30, 30)):
    """A one-colour RGB PNG, plain or Adam7-interlaced (Pillow cannot write interlaced PNG)."""
    if interlaced:
        raw = b""
        for x0, y0, dx, dy in [(0, 0, 8, 8), (4, 0, 8, 8), (0, 4, 4, 8), (2, 0, 4, 4), (0, 2, 2, 4), (1, 0, 2, 2), (0, 1, 1, 2)]:
            cols, rows = len(range(x0, w, dx)), len(range(y0, h, dy))
            if cols and rows:
                raw += (b"\x00" + bytes(rgb) * cols) * rows
    else:
        raw = (b"\x00" + bytes(rgb) * w) * h
    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 1 if interlaced else 0)
    open(path, "wb").write(b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def red_box(shapes, x):
    s = shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(1), Inches(1.5), Inches(1.5))
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor(200, 30, 30)
    return s


def swap_png_for_emf(src, dst, emf):
    """Replace the package's PNG picture by an EMF picture (same part, .emf name, EMF content type)."""
    zin, zout = zipfile.ZipFile(src), zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        name, data = item.filename, zin.read(item.filename)
        if "/media/" in name and name.endswith(".png"):
            name, data = name[:-4] + ".emf", open(emf, "rb").read()
        elif name.endswith(".rels"):
            data = re.sub(rb"(media/image\d+)\.png", rb"\1.emf", data)
        elif name == "[Content_Types].xml":
            data = data.replace(b'<Default Extension="png" ContentType="image/png"/>',
                                b'<Default Extension="png" ContentType="image/png"/><Default Extension="emf" ContentType="image/x-emf"/>')
        zout.writestr(name, data)
    zout.close()


# F1: three words in a text box
p = pptx.Presentation(); s = p.slides.add_slide(p.slide_layouts[6])
s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Alpha beta gamma"
p.save("words.pptx")

# F2: the same red picture, plain then interlaced
png("plain.png", False); png("adam7.png", True)
p = pptx.Presentation(); s = p.slides.add_slide(p.slide_layouts[6])
s.shapes.add_picture("plain.png", Inches(1), Inches(1), Inches(3)); s.shapes.add_picture("adam7.png", Inches(5), Inches(1), Inches(3))
p.save("interlaced.pptx")

# F3: an EMF picture (box.emf, made by LibreOffice from an SVG red rectangle) in a document and in a deck
d = docx.Document(); d.add_picture("plain.png", width=docx.shared.Inches(2)); d.save("_png.docx")
swap_png_for_emf("_png.docx", "emf.docx", "box.emf")
p = pptx.Presentation(); p.slides.add_slide(p.slide_layouts[6]).shapes.add_picture("plain.png", Inches(1), Inches(1), Inches(3))
p.save("_png.pptx"); swap_png_for_emf("_png.pptx", "emf.pptx", "box.emf")

# F4: bullets as Word writes them: a private-use code point of a symbol font in w:lvlText
d = docx.Document(); numbering = d.part.numbering_part.element
for k, (font, ch) in enumerate((("Wingdings 2", "\uf0a1"), ("Wingdings", "\uf0a7"), ("Symbol", "\uf0b7"))):
    numbering.insert(0, parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{900 + k}"><w:multiLevelType w:val="singleLevel"/>'
        f'<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="{ch}"/><w:lvlJc w:val="left"/>'
        f'<w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr><w:rPr><w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:hint="default"/></w:rPr>'
        f'</w:lvl></w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{900 + k}"><w:abstractNumId w:val="{900 + k}"/></w:num>'))
    para = d.add_paragraph(f"{font} bullet U+{ord(ch):04X}")
    para._p.get_or_add_pPr().append(parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{900 + k}"/></w:numPr>'))
d.save("bullets.docx")

# F5: a "Table Grid" table whose own w:tblBorders turn every border off
d = docx.Document(); t = d.add_table(rows=2, cols=2, style="Table Grid")
for i, c in enumerate(t._cells):
    c.text = f"cell {i}"
t._tbl.tblPr.find(qn("w:tblLook")).addprevious(parse_xml(f'<w:tblBorders {nsdecls("w")}>' + "".join(
    f'<w:{e} w:val="none" w:sz="0" w:space="0" w:color="auto"/>' for e in ("top", "left", "bottom", "right", "insideH", "insideV"))
    + "</w:tblBorders>"))
d.save("borders.docx")

# F6: a 40-line table row (no w:cantSplit) after 15 paragraphs, cells vertically centred
d = docx.Document()
for i in range(15):
    d.add_paragraph(f"Filler paragraph {i + 1}.")
t = d.add_table(rows=1, cols=1, style="Table Grid"); cell = t.cell(0, 0); cell.text = "Row line 1"
for i in range(2, 41):
    cell.add_paragraph(f"Row line {i}")
cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))
d.save("row_split.docx")

# F7: automatic text colour in a black cell
d = docx.Document(); c = d.add_table(rows=1, cols=1).cell(0, 0); c.text = "Auto colour on black"
c._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="000000"/>'))
d.save("auto_colour.docx")

# F8: a hidden shape and a hidden group on a slide, then a visible shape
p = pptx.Presentation(); s = p.slides.add_slide(p.slide_layouts[6])
red_box(s.shapes, 0.5)._element.nvSpPr.cNvPr.set("hidden", "1")
g = s.shapes.add_group_shape(); red_box(g.shapes, 4); g._element.nvGrpSpPr.cNvPr.set("hidden", "1")
red_box(s.shapes, 7.5)
p.save("hidden.pptx")

# F10: two shapes sharing an id
p = pptx.Presentation(); s = p.slides.add_slide(p.slide_layouts[6])
a = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)); a.text_frame.text = "First"
b = s.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1)); b.text_frame.text = "Second"
b._element.nvSpPr.cNvPr.set("id", a._element.nvSpPr.cNvPr.get("id"))
p.save("dup_id.pptx")

# F11: two bookmarks sharing a w:id (as Word 2013 saved two of the GOV.UK reports)
d = docx.Document()
for text, name in (("First heading", "_Toc378755770"), ("Some text", "_GoBack")):
    para = d.add_paragraph(text)
    para._p.insert(0, parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="2" w:name="{name}"/>'))
    para._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="2"/>'))
d.save("dup_bookmark.docx")
'''

PPTXGENJS = r'''
// F9: one pptxgenjs bar chart and one line chart, default options
const pptxgen = require("pptxgenjs");
const p = new pptxgen();
const data = [{ name: "Series A", labels: ["North", "South"], values: [3, 5] }];
p.addSlide().addChart(p.charts.BAR, data, { x: 1, y: 1, w: 6, h: 4 });
p.addSlide().addChart(p.charts.LINE, data, { x: 1, y: 1, w: 6, h: 4 });
p.writeFile({ fileName: "charts_pptxgenjs.pptx" }).then(() => console.log("ok"));
'''

SVG = '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100"><rect x="10" y="10" width="180" height="80" fill="#c81e1e"/></svg>'


def sh(cmd, **kw):
    return subprocess.run([str(c) for c in cmd], capture_output=True, text=True, **kw)


def rd(tool, name, out=None, extra=()):
    cli = common.RDOCX if tool == "rdocx" else common.RPPTX
    out = out or (WORK / (Path(name).stem + f".{tool}.pdf"))
    r = common.run([cli, "convert", WORK / name, "--to", "pdf", "-o", out, "--force", "--font-dir", common.FONTS_DIR, *extra],
                   env=common.env_rdocx(), cwd=WORK)
    msgs = [ln for ln in (r["stderr"] + r["stdout"]).splitlines() if ln.strip() and not ln.startswith("Written to")]
    return out, r["rc"], msgs


def lo(name):
    kind = "docx" if name.endswith(".docx") else "pptx"
    skill = common.BUILTIN_DOCX if kind == "docx" else common.BUILTIN_PPTX
    outdir = WORK / f"lo_{kind}"
    outdir.mkdir(exist_ok=True)
    common.run([common.PY, skill / "scripts" / "office" / "soffice.py", "--headless", "--convert-to", "pdf", "--outdir", outdir,
                WORK / name], env=common.env_builtin(), cwd=outdir)
    return outdir / (Path(name).stem + ".pdf")


def red_runs(pdf: Path, dpi=20):
    """x (inches) where red areas start on page 1."""
    prefix = WORK / f"{pdf.stem}-r"
    sh(["pdftoppm", "-r", dpi, "-f", 1, "-l", 1, "-png", pdf, prefix])
    a = np.asarray(Image.open(next(WORK.glob(prefix.name + "-*.png"))).convert("RGB")).astype(int)
    red = (a[:, :, 0] > 150) & (a[:, :, 1] < 80) & (a[:, :, 2] < 80)
    cols = np.where(red.any(0))[0]
    if not len(cols):
        return []
    return [round(float(r[0]) / dpi, 1) for r in np.split(cols, np.where(np.diff(cols) > 1)[0] + 1)]


def h_lines(pdf: Path):
    prefix = WORK / f"{pdf.stem}-g"
    sh(["pdftoppm", "-r", 72, "-gray", "-f", 1, "-l", 1, "-png", pdf, prefix])
    a = np.asarray(Image.open(next(WORK.glob(prefix.name + "-*.png"))).convert("L"))[:200]
    return [y for y in range(a.shape[0]) if (a[y] < 128).sum() > 200]


def text_on(pdf: Path, page: int, needle: str) -> int:
    return sh(["pdftotext", "-f", page, "-l", page, pdf, "-"]).stdout.count(needle)


def run() -> list[dict]:
    if WORK.exists():
        import shutil
        shutil.rmtree(WORK)
    WORK.mkdir(parents=True)
    (WORK / "box.svg").write_text(SVG)
    common.run([common.PY, common.BUILTIN_DOCX / "scripts" / "office" / "soffice.py", "--headless", "--convert-to", "emf",
                "--outdir", WORK, WORK / "box.svg"], env=common.env_builtin(), cwd=WORK, check=True)
    common.run([common.PY, "-c", BUILD], env=common.env_pylib(), cwd=WORK, check=True)
    (WORK / "charts.js").write_text(PPTXGENJS)
    common.run([common.NODE, WORK / "charts.js"], env=common.env_builtin(), cwd=WORK, check=True)
    out = []

    def add(fid, tool, title, expected, observed, reproduced, **kw):
        out.append({"id": fid, "tool": tool, "title": title, "expected": expected, "observed": observed,
                    "reproduced": bool(reproduced), **kw})

    pdf, _, _ = rd("rpptx", "words.pptx")
    words = [ln.split(">")[1].split("<")[0] for ln in sh(["pdftotext", "-bbox", pdf, "-"]).stdout.splitlines() if "<word" in ln]
    add("F1", "rpptx", "PDF text: a run's whole text sits on its first word",
        "three words with their own boxes (pdftotext -bbox: 'Alpha', 'beta', 'gamma')", f"pdftotext -bbox words: {words}",
        words == ["Alpha beta gamma"])

    pdf, _, msgs = rd("rpptx", "interlaced.pptx")
    imgs = len(sh(["pdfimages", "-list", pdf]).stdout.strip().splitlines()) - 2
    add("F2", "rpptx", "an interlaced (Adam7) PNG is not drawn, reported as a missing relationship",
        "both pictures drawn (rdocx draws the same interlaced PNG; LibreOffice too)",
        f"{imgs} of 2 pictures in the PDF; diagnostics: {msgs}", imgs < 2)

    pd, _, md = rd("rdocx", "emf.docx")
    pp, _, mp = rd("rpptx", "emf.pptx")
    ld = lo("emf.docx")
    add("F3", "rdocx, rpptx", "EMF pictures are not drawn",
        f"the red rectangle drawn (LibreOffice: red starts at x={red_runs(ld)} in)",
        f"rdocx: red at {red_runs(pd)}, diagnostics {md}; rpptx: red at {red_runs(pp)}, diagnostics {mp}",
        not red_runs(pd) or not red_runs(pp))

    pdf, _, msgs = rd("rdocx", "bullets.docx")
    xml = sh(["pdftohtml", "-xml", "-i", "-stdout", "-q", pdf]).stdout
    import re
    glyphs = [(m.group(1), m.group(2)) for m in re.finditer(r'font="(\d+)">([^<]{1})</text>', xml)]
    fams = dict(re.findall(r'<fontspec id="(\d+)" size="[^"]*" family="([^"]*)"', xml))
    got = [(fams.get(f), f"U+{ord(c):04X}") for f, c in glyphs]
    add("F4", "rdocx", "a Wingdings 2 bullet draws a missing-glyph box, with no diagnostic",
        "a black square, as Word and LibreOffice draw it (the Wingdings and Symbol bullets are mapped)",
        f"bullet glyphs (font, code point): {got}; diagnostics: {msgs}",
        any(cp == "U+F0A1" for _, cp in got))

    # F12: the same build, CLI and Python binding, same font folder
    py_pdf = WORK / "bullets.python.pdf"
    common.run([common.R_PY, "-c", "import rdocx, sys; open(sys.argv[2], 'wb').write(rdocx.Document(sys.argv[1])"
                ".to_pdf(font_dir=sys.argv[3]))", WORK / "bullets.docx", py_pdf, common.FONTS_DIR], env=common.env_rdocx(),
               check=True)
    xml_py = sh(["pdftohtml", "-xml", "-i", "-stdout", "-q", py_pdf]).stdout
    fams_py = dict(re.findall(r'<fontspec id="(\d+)" size="[^"]*" family="([^"]*)"', xml_py))
    got_py = [(fams_py.get(f), f"U+{ord(c):04X}") for f, c in re.findall(r'font="(\d+)">([^<]{1})</text>', xml_py)]
    add("F12", "rdocx", "the CLI and the Python binding draw the same document with different fonts",
        "the same glyphs from `rdocx convert --font-dir DIR` and `Document.to_pdf(font_dir=DIR)`",
        f"bullet glyphs, CLI: {got}; Python: {got_py} (FontAwesome, a system icon font, has a 'bullhorn' at U+F0A1). "
        "Checked once by hand on the Taylor past-schemes report: 48 pages (CLI) against 47 (Python)",
        got != got_py)

    pdf, _, _ = rd("rdocx", "borders.docx")
    add("F5", "rdocx", "a table style's borders are drawn although the table turns them off",
        f"no border (LibreOffice: horizontal lines at y={h_lines(lo('borders.docx'))})",
        f"horizontal lines at y={h_lines(pdf)} (72 dpi)", bool(h_lines(pdf)))

    pdf, _, _ = rd("rdocx", "row_split.docx")
    lpdf = lo("row_split.docx")
    r1, l1 = text_on(pdf, 1, "Row line"), text_on(lpdf, 1, "Row line")
    add("F6", "rdocx", "a table row with vertically centred cells is moved whole to the next page instead of split",
        f"the row starts on page 1 and breaks (LibreOffice: {l1} of its 40 lines on page 1; rdocx splits the same row "
        "without vAlign)", f"{r1} of 40 lines on page 1", r1 == 0)

    pdf, _, _ = rd("rdocx", "auto_colour.docx")
    col = re.findall(r'color="(#[0-9a-f]{6})"', sh(["pdftohtml", "-xml", "-i", "-stdout", "-q", pdf]).stdout)
    lcol = re.findall(r'color="(#[0-9a-f]{6})"', sh(["pdftohtml", "-xml", "-i", "-stdout", "-q", lo("auto_colour.docx")]).stdout)
    add("F7", "rdocx", "automatic text colour on a black cell stays black (invisible)",
        f"white text, as Word draws automatic colour on a dark fill (LibreOffice: {lcol})", f"text colour {col}",
        col and all(c == "#000000" for c in col))

    pdf, _, _ = rd("rpptx", "hidden.pptx")
    add("F8", "rpptx", "hidden shapes and groups (cNvPr hidden=\"1\") are drawn",
        f"only the visible box (LibreOffice: red starts at x={red_runs(lo('hidden.pptx'))} in)",
        f"red starts at x={red_runs(pdf)} in", len(red_runs(pdf)) > 1)

    pdf, _, msgs = rd("rpptx", "charts_pptxgenjs.pptx")
    add("F9", "rpptx", "pptxgenjs bar and line charts are drawn as 'Unsupported chart' placeholders",
        "both charts drawn, as LibreOffice draws them (pptxgenjs writes a third c:axId for a series axis it never "
        "declares; the ids are distinct)", f"diagnostics: {sorted(set(msgs))}",
        any("unsupported chart" in m for m in msgs))

    r = common.run([common.RPPTX, "fit", "--json", WORK / "dup_id.pptx"], env=common.env_rdocx())
    rr = common.run([common.RPPTX, "render", WORK / "dup_id.pptx", "-o", WORK / "dup_render"], env=common.env_rdocx())
    add("F10", "rpptx", "rpptx fit refuses a deck with a duplicate shape id (render accepts it)",
        "a fit report (PowerPoint repairs duplicate ids; render works on the same deck)",
        f"fit exit {r['rc']}: {r['stderr'].strip()[-110:]}; render exit {rr['rc']}", r["rc"] == 2)

    r = common.run([common.RDOCX, "render", WORK / "dup_bookmark.docx", "-o", WORK / "dup_bm"], env=common.env_rdocx())
    add("F11", "rdocx", "a document whose bookmarks share a w:id is refused (render, text, inspect, validate)",
        "the pages rendered (the id only pairs bookmarkStart with bookmarkEnd); see corpus/NOTES.md",
        f"render exit {r['rc']}: {r['stderr'].strip()[-120:]}", r["rc"] != 0)
    (WORK / "findings.json").write_text(json.dumps(out, indent=1))
    return out
