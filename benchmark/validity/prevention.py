"""Part 2, prevention: does a generator let an agent write a broken file?

Every footgun the built-in SKILL.md files document for docx-js and pptxgenjs, and the analogous mistakes on rdocx /
rpptx and python-docx / python-pptx, written as the minimal generator call an agent would write. Each case runs
twice, the mistake (``call``) and the same script done right (``control``), so that what the mistake changes is
isolated from whatever the generator writes anyway. For each case:

- ``refused_at_call``: the library raised at the call (exit 3), or at the save (exit 4: ``refused_at_save``);
- ``output_invalid``: the oracle (``oracle.py``) finds an issue in the mistake's output that the control's
  output does not have, or the case's own spec note says the written value breaks a rule the oracle cannot see;
- ``written``: what the call wrote, read back with zipfile and lxml (``inspect``);
- ``caught_by``: the checkers of Part 1 that flag the mistake's output and pass the control's.

Also ``atomic_save``: each stack's save over its input (or over an existing output, for the generators that
only create) while the process may not write more than 4 KiB (RLIMIT_FSIZE): is the input left intact?
"""

from __future__ import annotations

import json
import re
import resource
import signal
import subprocess
import sys
import time
import zipfile
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))
import common  # noqa: E402
import oracle  # noqa: E402
from run import norm  # noqa: E402

SKILL_DOCX = "built-in docx SKILL.md"
SKILL_PPTX = "built-in pptx SKILL.md"

# --- script templates -----------------------------------------------------------------------------------------
TEMPLATES = {
    "docx-js": ("node", """const fs = require("fs");
const docx = require("docx");
const { Document, Packer, Paragraph, TextRun, PageBreak, Table, TableRow, TableCell, WidthType, ShadingType,
        ImageRun, LevelFormat, AlignmentType, CommentRangeStart, CommentRangeEnd, CommentReference } = docx;
const OUT = process.argv[2], PNG = process.argv[3];
let doc;
try {
%CALL%
} catch (e) { console.log("REFUSED_AT_CALL: " + e.message); process.exit(3); }
Packer.toBuffer(doc).then((b) => fs.writeFileSync(OUT, b))
  .catch((e) => { console.log("REFUSED_AT_SAVE: " + e.message); process.exit(4); });
"""),
    "pptxgenjs": ("node", """const pptxgen = require("pptxgenjs");
const OUT = process.argv[2], PNG = process.argv[3];
(async () => {
  let pres = new pptxgen();
  pres.layout = "LAYOUT_16x9";
  try {
%CALL%
  } catch (e) { console.log("REFUSED_AT_CALL: " + e.message); process.exit(3); }
  try { await pres.writeFile({ fileName: OUT }); }
  catch (e) { console.log("REFUSED_AT_SAVE: " + e.message); process.exit(4); }
})();
"""),
    "rdocx": ("rdocx", """import sys
import docx_ops, rdocx
OUT, PNG = sys.argv[1], sys.argv[2]
doc = rdocx.Document()
doc.add_paragraph("Hello world")
try:
%CALL%
except Exception as e:
    print(f"REFUSED_AT_CALL: {type(e).__name__}: {e}"); sys.exit(3)
try:
    docx_ops.save_atomic(doc, OUT)
except Exception as e:
    print(f"REFUSED_AT_SAVE: {type(e).__name__}: {e}"); sys.exit(4)
"""),
    "rpptx": ("rdocx", """import sys
import pptx_ops, rpptx
from rpptx import RGBColor, Inches, Pt
OUT, PNG = sys.argv[1], sys.argv[2]
prs = rpptx.Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Title"
box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
box.text_frame.text = "Hello world"
try:
%CALL%
except Exception as e:
    print(f"REFUSED_AT_CALL: {type(e).__name__}: {e}"); sys.exit(3)
try:
    pptx_ops.save_atomic(prs, OUT)
except Exception as e:
    print(f"REFUSED_AT_SAVE: {type(e).__name__}: {e}"); sys.exit(4)
"""),
    "python-docx": ("pylib", """import sys
import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor
OUT, PNG = sys.argv[1], sys.argv[2]
doc = docx.Document()
p = doc.add_paragraph("Hello world")
try:
%CALL%
except Exception as e:
    print(f"REFUSED_AT_CALL: {type(e).__name__}: {e}"); sys.exit(3)
try:
    doc.save(OUT)
except Exception as e:
    print(f"REFUSED_AT_SAVE: {type(e).__name__}: {e}"); sys.exit(4)
"""),
    "python-pptx": ("pylib", """import sys
import pptx
from pptx.dml.color import RGBColor
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Inches, Pt
OUT, PNG = sys.argv[1], sys.argv[2]
prs = pptx.Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Title"
box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
box.text_frame.text = "Hello world"
try:
%CALL%
except Exception as e:
    print(f"REFUSED_AT_CALL: {type(e).__name__}: {e}"); sys.exit(3)
try:
    prs.save(OUT)
except Exception as e:
    print(f"REFUSED_AT_SAVE: {type(e).__name__}: {e}"); sys.exit(4)
"""),
}
STACK_OF = {"docx-js": "builtin", "pptxgenjs": "builtin", "rdocx": "rdocx", "rpptx": "rdocx",
            "python-docx": "pylib", "python-pptx": "pylib"}
FMT_OF = {"docx-js": "docx", "pptxgenjs": "pptx", "rdocx": "docx", "rpptx": "pptx", "python-docx": "docx",
          "python-pptx": "pptx"}


@dataclass
class Case:
    id: str
    generator: str
    footgun: str
    source: str
    call: str
    control: str | None
    inspect: Callable[[Path], str] | None = None
    spec: str = ""  # the rule at stake, and the verdict on what is written when the oracle cannot see it
    spec_invalid: Callable[[str], bool] | None = None  # verdict from ``written`` for rules the oracle cannot see
    cli: list | None = None  # a CLI case: argv with OUT / IN placeholders instead of a script

    @property
    def stack(self):
        return STACK_OF[self.generator]


def _xml(path: Path, part: str) -> str:
    with zipfile.ZipFile(path) as zf:
        return zf.read(part).decode("utf-8", "replace")


def _all_xml(path: Path, prefix: str) -> str:
    with zipfile.ZipFile(path) as zf:
        return "\n".join(zf.read(n).decode("utf-8", "replace") for n in zf.namelist()
                         if n.startswith(prefix) and n.endswith(".xml"))


def grep(prefix: str, pattern: str, limit=4):
    """An inspector: the matches of ``pattern`` in the XML parts under ``prefix``."""
    def inspect(path: Path) -> str:
        found = re.findall(pattern, _all_xml(path, prefix))
        return "; ".join(found[:limit]) if found else f"no match for {pattern}"
    return inspect


def first_text_contains(text):
    def inspect(path: Path) -> str:
        body = _all_xml(path, "word/") + _all_xml(path, "ppt/slides/")
        return f"{text!r} {'present' if text in body else 'absent'} in the written XML"
    return inspect


def ind(code: str, n: int) -> str:
    return "\n".join(" " * n + line if line.strip() else line for line in code.strip("\n").splitlines())


W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
P_NS = ('xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"')
RPPTX_SP = ('<p:sp {ns}><p:nvSpPr><p:cNvPr id="3" name="TextBox 2"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="914400" y="1828800"/><a:ext cx="3657600" cy="914400"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr>'
            '<p:txBody><a:bodyPr/><a:lstStyle/><a:p>{ppr}<a:r><a:t>Hello world</a:t></a:r></a:p></p:txBody></p:sp>')


def sp_xml(ppr="", order="normal"):
    xml = RPPTX_SP.format(ns=P_NS, ppr=ppr)
    if order == "spPr-first":
        nv = re.search(r"<p:nvSpPr>.*?</p:nvSpPr>", xml).group(0)
        sppr = re.search(r"<p:spPr>.*?</p:spPr>", xml).group(0)
        xml = xml.replace(nv, "@@NV@@").replace(sppr, nv).replace("@@NV@@", sppr)
    return xml


CASES: list[Case] = [
    # ---------------- docx-js: the built-in docx SKILL.md gotchas --------------------------------------------
    Case("docxjs.pagebreak_outside_paragraph", "docx-js", "PageBreak outside a Paragraph",
         f'{SKILL_DOCX}: "PageBreak must be inside a Paragraph."',
         'doc = new Document({ sections: [{ children: [new Paragraph("a"), new PageBreak(), new Paragraph("b")] }] });',
         'doc = new Document({ sections: [{ children: [new Paragraph("a"), new Paragraph({ children: [new PageBreak()] }), new Paragraph("b")] }] });',
         grep("word/document.xml", r"<w:body><w:p>.{0,80}|</w:p><w:r>.{0,40}"),
         "CT_Body holds block-level elements only: a w:r directly in w:body breaks wml.xsd."),
    Case("docxjs.newline_in_text", "docx-js", "\\n inside a TextRun",
         f'{SKILL_DOCX}: "Never use \\\\n -- use separate Paragraph elements."',
         'doc = new Document({ sections: [{ children: [new Paragraph({ children: [new TextRun("line one\\nline two")] })] }] });',
         'doc = new Document({ sections: [{ children: [new Paragraph("line one"), new Paragraph("line two")] }] });',
         grep("word/document.xml", r"<w:t[^>]*>line one[^<]*</w:t>"),
         "A line feed is a legal character of w:t (spec-valid); Word lays it out as white space, not a break."),
    Case("docxjs.shading_solid", "docx-js", "ShadingType.SOLID in a table cell",
         f'{SKILL_DOCX}: "Table shading: use ShadingType.CLEAR, never SOLID (renders black)."',
         'doc = new Document({ sections: [{ children: [new Table({ width: { size: 9360, type: WidthType.DXA }, columnWidths: [9360], rows: [new TableRow({ children: [new TableCell({ width: { size: 9360, type: WidthType.DXA }, shading: { type: ShadingType.SOLID, color: "D9E2F3" }, children: [new Paragraph("cell")] })] })] })] }] });',
         'doc = new Document({ sections: [{ children: [new Table({ width: { size: 9360, type: WidthType.DXA }, columnWidths: [9360], rows: [new TableRow({ children: [new TableCell({ width: { size: 9360, type: WidthType.DXA }, shading: { type: ShadingType.CLEAR, color: "auto", fill: "D9E2F3" }, children: [new Paragraph("cell")] })] })] })] }] });',
         grep("word/document.xml", r"<w:shd [^>]*/>"),
         'w:shd w:val="solid" with a pattern colour is spec-valid: a 100 % pattern in w:color (the cell paints in that colour, black when "auto").'),
    Case("docxjs.table_width_percentage", "docx-js", "WidthType.PERCENTAGE for table and cell widths",
         f'{SKILL_DOCX}: "set columnWidths on the table AND width on every cell, both in WidthType.DXA (PERCENTAGE breaks in Google Docs)."',
         'doc = new Document({ sections: [{ children: [new Table({ width: { size: 100, type: WidthType.PERCENTAGE }, rows: [new TableRow({ children: [new TableCell({ width: { size: 50, type: WidthType.PERCENTAGE }, children: [new Paragraph("a")] }), new TableCell({ width: { size: 50, type: WidthType.PERCENTAGE }, children: [new Paragraph("b")] })] })] })] }] });',
         'doc = new Document({ sections: [{ children: [new Table({ width: { size: 9360, type: WidthType.DXA }, columnWidths: [4680, 4680], rows: [new TableRow({ children: [new TableCell({ width: { size: 4680, type: WidthType.DXA }, children: [new Paragraph("a")] }), new TableCell({ width: { size: 4680, type: WidthType.DXA }, children: [new Paragraph("b")] })] })] })] }] });',
         grep("word/document.xml", r"<w:(?:tblW|tcW|gridCol) [^>]*/>"),
         'ST_MeasurementOrPercent accepts "100%" with w:type="pct" (ISO/IEC 29500 transitional): spec-valid.'),
    Case("docxjs.imagerun_without_type", "docx-js", "ImageRun without type:",
         f'{SKILL_DOCX}: "ImageRun requires type: (\\"png\\", \\"jpg\\", ...)."',
         'doc = new Document({ sections: [{ children: [new Paragraph({ children: [new ImageRun({ data: fs.readFileSync(PNG), transformation: { width: 100, height: 75 } })] })] }] });',
         'doc = new Document({ sections: [{ children: [new Paragraph({ children: [new ImageRun({ type: "png", data: fs.readFileSync(PNG), transformation: { width: 100, height: 75 } })] })] }] });',
         grep("", r'Target="media/[^"]*"|Extension="[^"]*"', 8),
         "The picture part needs a content type (Part 2); see the oracle."),
    Case("docxjs.literal_bullet", "docx-js", "a literal bullet character instead of a numbering config",
         f'{SKILL_DOCX}: "Lists: never insert \\u2022 literally; use a numbering config with LevelFormat.BULLET."',
         'doc = new Document({ sections: [{ children: [new Paragraph("\\u2022 first item"), new Paragraph("\\u2022 second item")] }] });',
         'doc = new Document({ numbering: { config: [{ reference: "b", levels: [{ level: 0, format: LevelFormat.BULLET, text: "\\u2022", alignment: AlignmentType.LEFT }] }] }, sections: [{ children: [new Paragraph({ text: "first item", numbering: { reference: "b", level: 0 } }), new Paragraph({ text: "second item", numbering: { reference: "b", level: 0 } })] }] });',
         grep("word/document.xml", r"<w:t[^>]*>[^<]*item</w:t>|<w:numId [^>]*/>"),
         "Cosmetic (not a list for Word); spec-valid."),
    Case("docxjs.column_widths_mismatch", "docx-js", "columnWidths that do not sum to the table width",
         f'{SKILL_DOCX}: "Column widths must sum to the table width."',
         'doc = new Document({ sections: [{ children: [new Table({ width: { size: 9360, type: WidthType.DXA }, columnWidths: [2000, 2000], rows: [new TableRow({ children: [new TableCell({ width: { size: 2000, type: WidthType.DXA }, children: [new Paragraph("a")] }), new TableCell({ width: { size: 2000, type: WidthType.DXA }, children: [new Paragraph("b")] })] })] })] }] });',
         'doc = new Document({ sections: [{ children: [new Table({ width: { size: 9360, type: WidthType.DXA }, columnWidths: [4680, 4680], rows: [new TableRow({ children: [new TableCell({ width: { size: 4680, type: WidthType.DXA }, children: [new Paragraph("a")] }), new TableCell({ width: { size: 4680, type: WidthType.DXA }, children: [new Paragraph("b")] })] })] })] }] });',
         grep("word/document.xml", r"<w:(?:tblW|gridCol) [^>]*/>"),
         "Layout only (the grid wins or is recomputed); spec-valid."),
    Case("docxjs.colour_hash", "docx-js", "a colour with '#' (analogue of the pptxgenjs fault)",
         "analogue", 'doc = new Document({ sections: [{ children: [new Paragraph({ children: [new TextRun({ text: "red", color: "#FF0000" })] })] }] });',
         'doc = new Document({ sections: [{ children: [new Paragraph({ children: [new TextRun({ text: "red", color: "FF0000" })] })] }] });',
         grep("word/document.xml", r"<w:color [^>]*/>"), "w:color/@w:val is ST_HexColor (six hex digits or auto)."),
    Case("docxjs.colour_8digit", "docx-js", "an 8-digit (alpha) hex colour (analogue)",
         "analogue", 'doc = new Document({ sections: [{ children: [new Paragraph({ children: [new TextRun({ text: "red", color: "FF000080" })] })] }] });',
         'doc = new Document({ sections: [{ children: [new Paragraph({ children: [new TextRun({ text: "red", color: "FF0000" })] })] }] });',
         grep("word/document.xml", r"<w:color [^>]*/>"), "w:color/@w:val is ST_HexColor (six hex digits or auto)."),
    Case("docxjs.comment_reference_missing", "docx-js", "a comment reference to an id no comment has (analogue)",
         "analogue",
         'doc = new Document({ comments: { children: [{ id: 0, author: "A", date: new Date("2026-10-01T00:00:00Z"), children: [new Paragraph("note")] }] }, sections: [{ children: [new Paragraph({ children: [new CommentRangeStart(5), new TextRun("text"), new CommentRangeEnd(5), new TextRun({ children: [new CommentReference(5)] })] })] }] });',
         'doc = new Document({ comments: { children: [{ id: 0, author: "A", date: new Date("2026-10-01T00:00:00Z"), children: [new Paragraph("note")] }] }, sections: [{ children: [new Paragraph({ children: [new CommentRangeStart(0), new TextRun("text"), new CommentRangeEnd(0), new TextRun({ children: [new CommentReference(0)] })] })] }] });',
         grep("word/", r'<w:comment(?:Reference)? [^>]*w:id="\d+"', 3),
         "Part 1 17.13.4.5: the reference names comment 5, which does not exist (dangling reference).",
         spec_invalid=lambda written: 'w:id="5"' in written),
    Case("docxjs.negative_image_size", "docx-js", "a negative picture width (analogue)",
         "analogue", 'doc = new Document({ sections: [{ children: [new Paragraph({ children: [new ImageRun({ type: "png", data: fs.readFileSync(PNG), transformation: { width: -100, height: 75 } })] })] }] });',
         'doc = new Document({ sections: [{ children: [new Paragraph({ children: [new ImageRun({ type: "png", data: fs.readFileSync(PNG), transformation: { width: 100, height: 75 } })] })] }] });',
         grep("word/document.xml", r"<wp:extent [^>]*/>"), "wp:extent/@cx is ST_PositiveCoordinate (>= 0)."),

    # ---------------- pptxgenjs: the built-in pptx SKILL.md gotchas -------------------------------------------
    Case("pptxgenjs.colour_hash_fill", "pptxgenjs", "a hex colour with '#' (shape fill)",
         f'{SKILL_PPTX}: "Hex colors: never #, never 8 digits ... Both \\"#FF0000\\" and alpha baked into the hex ... corrupt the file."',
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "#FF0000" } });', 4),
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "FF0000" } });', 4),
         grep("ppt/slides/", r'<a:srgbClr val="[^"]*"'), "a:srgbClr/@val is ST_HexColorRGB (six hex digits)."),
    Case("pptxgenjs.colour_hash_shadow", "pptxgenjs", "a hex colour with '#' (shadow colour)",
         f'{SKILL_PPTX}: (same) "never #"',
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "FFFFFF" }, shadow: { type: "outer", blur: 3, offset: 2, angle: 45, color: "#000000", opacity: 0.4 } });', 4),
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "FFFFFF" }, shadow: { type: "outer", blur: 3, offset: 2, angle: 45, color: "000000", opacity: 0.4 } });', 4),
         grep("ppt/slides/", r'<a:outerShdw[\s\S]{0,200}?<a:srgbClr val="[^"]*"'), "a:srgbClr/@val is ST_HexColorRGB (six hex digits)."),
    Case("pptxgenjs.colour_hash_gridline", "pptxgenjs", "a hex colour with '#' (chart gridline colour)",
         f'{SKILL_PPTX}: (same) "never #"',
         ind('pres.addSlide().addChart(pres.charts.BAR, [{ name: "A", labels: ["x", "y"], values: [1, 2] }], { x: 1, y: 1, w: 6, h: 4, valGridLine: { color: "#E0E0E0", size: 1 } });', 4),
         ind('pres.addSlide().addChart(pres.charts.BAR, [{ name: "A", labels: ["x", "y"], values: [1, 2] }], { x: 1, y: 1, w: 6, h: 4, valGridLine: { color: "E0E0E0", size: 1 } });', 4),
         grep("ppt/charts/", r'<c:majorGridlines>[\s\S]{0,200}?<a:srgbClr val="[^"]*"'), "a:srgbClr/@val is ST_HexColorRGB (six hex digits)."),
    Case("pptxgenjs.colour_8digit_fill", "pptxgenjs", "an 8-digit hex colour, alpha baked in (shape fill)",
         f'{SKILL_PPTX}: "alpha baked into the hex (\\"00000020\\") corrupt[s] the file. For translucency: transparency: 0-100 on fills."',
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "00000020" } });', 4),
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "000000", transparency: 87 } });', 4),
         grep("ppt/slides/", r'<a:srgbClr val="[^"]*"(?:/>|>\s*<a:alpha val="[^"]*")'), "a:srgbClr/@val is ST_HexColorRGB (six hex digits)."),
    Case("pptxgenjs.colour_8digit_shadow", "pptxgenjs", "an 8-digit hex colour, alpha baked in (shadow colour)",
         f'{SKILL_PPTX}: (same) "never 8 digits ... opacity: 0.0-1.0 on shadows"',
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "FFFFFF" }, shadow: { type: "outer", blur: 3, offset: 2, angle: 45, color: "00000020" } });', 4),
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "FFFFFF" }, shadow: { type: "outer", blur: 3, offset: 2, angle: 45, color: "000000", opacity: 0.13 } });', 4),
         grep("ppt/slides/", r'<a:outerShdw[\s\S]{0,200}?<a:srgbClr val="[^"]*"'), "a:srgbClr/@val is ST_HexColorRGB (six hex digits)."),
    Case("pptxgenjs.shared_option_object", "pptxgenjs", "one shadow options object shared by two addShape calls",
         f'{SKILL_PPTX}: "pptxgenjs mutates option objects in place (converts values to EMU on first use). Never share one shadow/options object across two add* calls."',
         ind('const s = pres.addSlide(); const sh = { type: "outer", blur: 3, offset: 2, angle: 45, color: "000000", opacity: 0.4 };\n'
             's.addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 2, h: 1, fill: { color: "FFFFFF" }, shadow: sh });\n'
             's.addShape(pres.shapes.RECTANGLE, { x: 4, y: 1, w: 2, h: 1, fill: { color: "FFFFFF" }, shadow: sh });', 4),
         ind('const s = pres.addSlide(); const sh = () => ({ type: "outer", blur: 3, offset: 2, angle: 45, color: "000000", opacity: 0.4 });\n'
             's.addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 2, h: 1, fill: { color: "FFFFFF" }, shadow: sh() });\n'
             's.addShape(pres.shapes.RECTANGLE, { x: 4, y: 1, w: 2, h: 1, fill: { color: "FFFFFF" }, shadow: sh() });', 4),
         grep("ppt/slides/", r"<a:outerShdw [^>]*>"),
         "Spec-valid either way; the second shadow gets values converted twice (wrong, not invalid)."),
    Case("pptxgenjs.negative_shadow_offset", "pptxgenjs", "a negative shadow offset",
         f'{SKILL_PPTX}: "Shadow offset must be >= 0 -- a negative offset corrupts the file."',
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "FFFFFF" }, shadow: { type: "outer", blur: 3, offset: -2, angle: 90, color: "000000", opacity: 0.4 } });', 4),
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "FFFFFF" }, shadow: { type: "outer", blur: 3, offset: 2, angle: 270, color: "000000", opacity: 0.4 } });', 4),
         grep("ppt/slides/", r"<a:outerShdw [^>]*>"), "a:outerShdw/@dist is ST_PositiveCoordinate (>= 0)."),
    Case("pptxgenjs.letter_spacing", "pptxgenjs", "letterSpacing instead of charSpacing",
         f'{SKILL_PPTX}: "letterSpacing is silently ignored -- the real option is charSpacing."',
         ind('pres.addSlide().addText("spaced", { x: 1, y: 1, w: 4, h: 1, letterSpacing: 4, isTextBox: true });', 4),
         ind('pres.addSlide().addText("spaced", { x: 1, y: 1, w: 4, h: 1, charSpacing: 4, isTextBox: true });', 4),
         grep("ppt/slides/", r"<a:rPr [^>]*>"), "Spec-valid; the option is dropped."),
    Case("pptxgenjs.stacked_outend", "pptxgenjs", "dataLabelPosition outEnd on a stacked bar chart",
         f'{SKILL_PPTX}: "On a stacked bar or column chart, dataLabelPosition must be ctr, inEnd, or inBase. outEnd corrupts the file."',
         ind('pres.addSlide().addChart(pres.charts.BAR, [{ name: "A", labels: ["x", "y"], values: [1, 2] }, { name: "B", labels: ["x", "y"], values: [3, 4] }], { x: 1, y: 1, w: 6, h: 4, barDir: "col", barGrouping: "stacked", showValue: true, dataLabelPosition: "outEnd" });', 4),
         ind('pres.addSlide().addChart(pres.charts.BAR, [{ name: "A", labels: ["x", "y"], values: [1, 2] }, { name: "B", labels: ["x", "y"], values: [3, 4] }], { x: 1, y: 1, w: 6, h: 4, barDir: "col", barGrouping: "stacked", showValue: true, dataLabelPosition: "ctr" });', 4),
         grep("ppt/charts/", r'<c:grouping val="[^"]*"/>|<c:dLblPos val="[^"]*"/>', 3),
         "Spec-valid (ST_DLblPos has outEnd and no clause restricts it by grouping); PowerPoint refuses it per the SKILL.md.",
         spec_invalid=lambda written: False),
    Case("pptxgenjs.secondary_axis_without_axes", "pptxgenjs", "a combo chart on secondary axes without valAxes/catAxes",
         f'{SKILL_PPTX}: "A combo series using secondaryValAxis/secondaryCatAxis needs both valAxes and catAxes ... Without them pptxgenjs writes axis ids it never declares, and PowerPoint discards that chart and reports the file as corrupt."',
         ind('pres.addSlide().addChart([{ type: pres.charts.BAR, data: [{ name: "A", labels: ["x", "y"], values: [1, 2] }], options: { barDir: "col" } }, { type: pres.charts.LINE, data: [{ name: "B", labels: ["x", "y"], values: [30, 40] }], options: { secondaryValAxis: true, secondaryCatAxis: true } }], { x: 1, y: 1, w: 6, h: 4 });', 4),
         ind('pres.addSlide().addChart([{ type: pres.charts.BAR, data: [{ name: "A", labels: ["x", "y"], values: [1, 2] }], options: { barDir: "col" } }, { type: pres.charts.LINE, data: [{ name: "B", labels: ["x", "y"], values: [30, 40] }], options: { secondaryValAxis: true, secondaryCatAxis: true } }], { x: 1, y: 1, w: 6, h: 4, valAxes: [{ showValAxisTitle: false }, { showValAxisTitle: false }], catAxes: [{ catAxisTitle: "x" }, { catAxisHidden: true }] });', 4),
         grep("ppt/charts/", r'<c:axId val="\d+"/>', 12),
         "Part 1 21.2.2.9: each chart group's c:axId names an axis the plot area declares (judgement: an undeclared id is a dangling reference).",
         spec_invalid=None),
    Case("pptxgenjs.company_ampersand", "pptxgenjs", "'&' in pres.company",
         f'{SKILL_PPTX}: "pres.company and a defineSlideMaster() title are written into the XML unescaped -- an & in either (\\"Harrow & Sons\\") leaves docProps/app.xml or the layout unparseable, and validate.py fails."',
         ind('pres.company = "Harrow & Sons"; pres.addSlide().addText("x", { x: 1, y: 1, w: 2, h: 1, isTextBox: true });', 4),
         ind('pres.company = "Harrow &amp; Sons"; pres.addSlide().addText("x", { x: 1, y: 1, w: 2, h: 1, isTextBox: true });', 4),
         grep("docProps/app.xml", r"<Company>[^<]*</Company>"), "XML 1.0: an unescaped '&' is not well-formed."),
    Case("pptxgenjs.layout_title_ampersand", "pptxgenjs", "'&' in a defineSlideMaster title",
         f'{SKILL_PPTX}: (same) "an & in ... a defineSlideMaster() title ... leaves ... the layout unparseable"',
         ind('pres.defineSlideMaster({ title: "Harrow & Sons", background: { color: "FFFFFF" } }); pres.addSlide({ masterName: "Harrow & Sons" }).addText("x", { x: 1, y: 1, w: 2, h: 1, isTextBox: true });', 4),
         ind('pres.defineSlideMaster({ title: "Harrow &amp; Sons", background: { color: "FFFFFF" } }); pres.addSlide({ masterName: "Harrow &amp; Sons" }).addText("x", { x: 1, y: 1, w: 2, h: 1, isTextBox: true });', 4),
         grep("ppt/slideLayouts/", r'<p:cSld name="[^"]*"'), "XML 1.0: an unescaped '&' is not well-formed."),
    Case("pptxgenjs.reused_instance", "pptxgenjs", "one pptxgen instance written twice",
         f'{SKILL_PPTX}: "One new pptxgen() per output file -- never reuse an instance."',
         ind('pres.addSlide().addText("first", { x: 1, y: 1, w: 2, h: 1, isTextBox: true });\n'
             'await pres.writeFile({ fileName: OUT + ".first.pptx" });\n'
             'pres.addSlide().addText("second", { x: 1, y: 1, w: 2, h: 1, isTextBox: true });', 4),
         ind('pres.addSlide().addText("first", { x: 1, y: 1, w: 2, h: 1, isTextBox: true });\n'
             'await pres.writeFile({ fileName: OUT + ".first.pptx" });\n'
             'pres = new pptxgen(); pres.layout = "LAYOUT_16x9";\n'
             'pres.addSlide().addText("first", { x: 1, y: 1, w: 2, h: 1, isTextBox: true });\n'
             'pres.addSlide().addText("second", { x: 1, y: 1, w: 2, h: 1, isTextBox: true });', 4),
         grep("", r'PartName="/ppt/slides/slide\d+\.xml"|<p:sldId [^>]*/>', 8),
         "See the oracle (the second file must be a complete, consistent package)."),
    Case("pptxgenjs.rect_radius_on_rectangle", "pptxgenjs", "rectRadius on RECTANGLE",
         f'{SKILL_PPTX}: "rectRadius only works on ROUNDED_RECTANGLE, not RECTANGLE."',
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, rectRadius: 0.2, fill: { color: "FFFFFF" } });', 4),
         ind('pres.addSlide().addShape(pres.shapes.ROUNDED_RECTANGLE, { x: 1, y: 1, w: 3, h: 1, rectRadius: 0.2, fill: { color: "FFFFFF" } });', 4),
         grep("ppt/slides/", r'<a:prstGeom prst="[^"]*">|<a:gd name="adj"[^>]*/>'), "Spec-valid; the radius is dropped."),
    Case("pptxgenjs.scheme_colour_in_hex_option", "pptxgenjs", "a scheme colour in a hex-only option (shadow colour)",
         f'{SKILL_PPTX}: "These take hex only ... shadow.color ... A scheme color there is written as <a:srgbClr val=\\"bg2\\"/> without complaint."',
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "FFFFFF" }, shadow: { type: "outer", blur: 3, offset: 2, angle: 45, color: pres.SchemeColor.background2, opacity: 0.4 } });', 4),
         ind('pres.addSlide().addShape(pres.shapes.RECTANGLE, { x: 1, y: 1, w: 3, h: 1, fill: { color: "FFFFFF" }, shadow: { type: "outer", blur: 3, offset: 2, angle: 45, color: "E7E6E6", opacity: 0.4 } });', 4),
         grep("ppt/slides/", r'<a:outerShdw[\s\S]{0,200}?<a:srgbClr val="[^"]*"'), "a:srgbClr/@val is ST_HexColorRGB (six hex digits)."),
    Case("pptxgenjs.misspelt_placeholder", "pptxgenjs", "a misspelt placeholder name",
         f'{SKILL_PPTX}: "A misspelt placeholder or masterName fails silently (a stray text box at 0,0; the blank layout)."',
         ind('pres.defineSlideMaster({ title: "T", objects: [{ placeholder: { options: { name: "title", type: "title", x: 0.5, y: 0.3, w: 9, h: 0.8 }, text: "" } }] });\n'
             'pres.addSlide({ masterName: "T" }).addText("Heading", { placeholder: "titel" });', 4),
         ind('pres.defineSlideMaster({ title: "T", objects: [{ placeholder: { options: { name: "title", type: "title", x: 0.5, y: 0.3, w: 9, h: 0.8 }, text: "" } }] });\n'
             'pres.addSlide({ masterName: "T" }).addText("Heading", { placeholder: "title" });', 4),
         grep("ppt/slides/", r"<p:nvPr>.{0,60}"), "Spec-valid; the text lands in a stray box."),
    Case("pptxgenjs.text_off_slide", "pptxgenjs", "coordinates past the slide edge",
         f'{SKILL_PPTX}: "Coordinates past the edge are written, not clamped -- the shape just isn\'t on the slide."',
         ind('pres.addSlide().addText("lost", { x: 12, y: 1, w: 3, h: 1, isTextBox: true });', 4),
         ind('pres.addSlide().addText("lost", { x: 6, y: 1, w: 3, h: 1, isTextBox: true });', 4),
         grep("ppt/slides/", r'<a:off x="\d+" y="\d+"/>'), "Spec-valid (ST_Coordinate is signed and unbounded by the slide)."),

    # ---------------- rdocx (docx) analogues ------------------------------------------------------------------
    Case("rdocx.colour_hash", "rdocx", "a colour string with '#'", "analogue",
         ind('doc.paragraphs[0].runs[0].font.color.rgb = "#FF0000"', 4),
         ind('doc.paragraphs[0].runs[0].font.color.rgb = "FF0000"', 4),
         grep("word/document.xml", r"<w:color [^>]*/>"), "w:color/@w:val is ST_HexColor."),
    Case("rdocx.colour_8digit", "rdocx", "an 8-digit hex colour", "analogue",
         ind('doc.paragraphs[0].runs[0].font.color.rgb = "FF000080"', 4),
         ind('doc.paragraphs[0].runs[0].font.color.rgb = "FF0000"', 4),
         grep("word/document.xml", r"<w:color [^>]*/>"), "w:color/@w:val is ST_HexColor."),
    Case("rdocx.cell_shading_hash", "rdocx", "a cell shading colour with '#'", "analogue",
         ind('t = doc.add_table(1, 1, style="Table Grid")\nt.cell(0, 0).shading = "#D9E2F3"', 4),
         ind('t = doc.add_table(1, 1, style="Table Grid")\nt.cell(0, 0).shading = "D9E2F3"', 4),
         grep("word/document.xml", r"<w:shd [^>]*/>"), "w:shd/@w:fill is ST_HexColor."),
    Case("rdocx.negative_picture_width", "rdocx", "a negative picture width", "analogue",
         ind('doc.add_picture(PNG, "p.png", -rdocx.Inches(1), rdocx.Inches(0.75))', 4),
         ind('doc.add_picture(PNG, "p.png", rdocx.Inches(1), rdocx.Inches(0.75))', 4),
         grep("word/document.xml", r"<wp:extent [^>]*/>"), "wp:extent/@cx is ST_PositiveCoordinate (>= 0)."),
    Case("rdocx.bare_int_font_size", "rdocx", "a bare int where a length is expected (font.size = 12)",
         "analogue (rdocx docx SKILL.md rule 9 says it raises)",
         ind('doc.paragraphs[0].runs[0].font.size = 12', 4),
         ind('doc.paragraphs[0].runs[0].font.size = rdocx.Pt(12)', 4),
         grep("word/document.xml", r"<w:sz [^>]*/>"), "A size of 12 EMU would round to w:sz 0 (spec-valid but invisible)."),
    Case("rdocx.replace_xml_malformed", "rdocx", "malformed XML through replace_xml", "analogue",
         ind(f'doc.paragraphs[0].replace_xml(\'<w:p {W_NS}><w:r><w:t>x</w:t></w:p>\')', 4),
         ind(f'doc.paragraphs[0].replace_xml(\'<w:p {W_NS}><w:r><w:t>x</w:t></w:r></w:p>\')', 4),
         None, "XML 1.0 well-formedness."),
    Case("rdocx.replace_xml_misordered", "rdocx", "w:pPr after w:r through replace_xml", "analogue",
         ind(f'doc.paragraphs[0].replace_xml(\'<w:p {W_NS}><w:r><w:t>x</w:t></w:r><w:pPr><w:jc w:val="center"/></w:pPr></w:p>\')', 4),
         ind(f'doc.paragraphs[0].replace_xml(\'<w:p {W_NS}><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:t>x</w:t></w:r></w:p>\')', 4),
         grep("word/document.xml", r"<w:p>.{0,90}"), "CT_P: pPr first."),
    Case("rdocx.replace_xml_unknown_element", "rdocx", "an unknown w: element through replace_xml", "analogue",
         ind(f'doc.paragraphs[0].replace_xml(\'<w:p {W_NS}><w:bogus/><w:r><w:t>x</w:t></w:r></w:p>\')', 4),
         ind(f'doc.paragraphs[0].replace_xml(\'<w:p {W_NS}><w:r><w:t>x</w:t></w:r></w:p>\')', 4),
         grep("word/document.xml", r"<w:p>.{0,60}"), "wml.xsd has no w:bogus."),
    Case("rdocx.replace_xml_bad_enum", "rdocx", "an invalid enum value through replace_xml", "analogue",
         ind(f'doc.paragraphs[0].replace_xml(\'<w:p {W_NS}><w:pPr><w:jc w:val="middle"/></w:pPr><w:r><w:t>x</w:t></w:r></w:p>\')', 4),
         ind(f'doc.paragraphs[0].replace_xml(\'<w:p {W_NS}><w:pPr><w:jc w:val="center"/></w:pPr><w:r><w:t>x</w:t></w:r></w:p>\')', 4),
         grep("word/document.xml", r"<w:jc [^>]*/>"), 'ST_Jc has no "middle".'),
    Case("rdocx.unknown_style", "rdocx", "an unknown style name", "analogue",
         ind('doc.paragraphs[0].style = "No Such Style"', 4),
         ind('doc.paragraphs[0].style = "Heading 1"', 4),
         grep("word/document.xml", r"<w:pStyle [^>]*/>"), "Not a schema rule (judgement: a dangling style id is tolerated)."),
    Case("rdocx.stale_run_handle", "rdocx", "a run handle used after an append retired it", "analogue",
         ind('r = doc.paragraphs[0].runs[0]\ndoc.add_paragraph("appended")\nr.text = "changed"', 4),
         ind('r = doc.paragraphs[0].runs[0]\ndoc.add_paragraph("appended")\ndoc.paragraphs[0].runs[0].text = "changed"', 4),
         first_text_contains("changed"), "Not a validity question: is the edit applied, refused, or lost?"),
    Case("rdocx.comment_on_missing_text", "rdocx", "a comment anchored on text that is not there", "analogue",
         ind('docx_ops.comment_on_text(doc, "Not in the document", "Check", "Reviewer", date="2026-10-01T00:00:00Z")', 4),
         ind('docx_ops.comment_on_text(doc, "Hello", "Check", "Reviewer", date="2026-10-01T00:00:00Z")', 4),
         grep("word/document.xml", r"<w:commentR[^>]*/>", 3), "A comment with no anchor would be an orphan."),
    Case("rdocx.replace_count_mismatch", "rdocx", "a replacement whose expected count is wrong (CLI)",
         "analogue (rdocx docx SKILL.md rule 2)", "rdocx replace IN -p 'Quarterly report' -v 'Annual report' --expect 2 -o OUT",
         "rdocx replace IN -p 'Quarterly report' -v 'Annual report' --expect 1 -o OUT",
         first_text_contains("Annual report"), "Not a validity question: is a wrong count refused?",
         cli=[common.RDOCX, "replace", "IN", "-p", "Quarterly report", "-v", "Annual report", "--expect", "EXPECT", "-o", "OUT"]),

    # ---------------- rpptx (pptx) analogues -----------------------------------------------------------------
    Case("rpptx.colour_hash", "rpptx", "RGBColor.from_string with '#'", "analogue",
         ind('box.fill.solid()\nbox.fill.fore_color.rgb = RGBColor.from_string("#FF0000")', 4),
         ind('box.fill.solid()\nbox.fill.fore_color.rgb = RGBColor.from_string("FF0000")', 4),
         grep("ppt/slides/", r'<a:srgbClr val="[^"]*"'), "a:srgbClr/@val is ST_HexColorRGB."),
    Case("rpptx.colour_8digit", "rpptx", "RGBColor.from_string with 8 hex digits", "analogue",
         ind('box.fill.solid()\nbox.fill.fore_color.rgb = RGBColor.from_string("00000020")', 4),
         ind('box.fill.solid()\nbox.fill.fore_color.rgb = RGBColor.from_string("000000")', 4),
         grep("ppt/slides/", r'<a:srgbClr val="[^"]*"'), "a:srgbClr/@val is ST_HexColorRGB."),
    Case("rpptx.negative_shadow_distance", "rpptx", "a negative shadow distance", "analogue",
         ind('box.shadow.visible = True\nbox.shadow.distance = -Pt(3)', 4),
         ind('box.shadow.visible = True\nbox.shadow.distance = Pt(3)', 4),
         grep("ppt/slides/", r"<a:outerShdw [^>]*>"), "a:outerShdw/@dist is ST_PositiveCoordinate."),
    Case("rpptx.negative_width", "rpptx", "a negative shape width", "analogue",
         ind('slide.shapes.add_textbox(Inches(1), Inches(4), -Inches(3), Inches(1)).text_frame.text = "neg"', 4),
         ind('slide.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1)).text_frame.text = "neg"', 4),
         grep("ppt/slides/", r'<a:ext cx="-?\d+" cy="\d+"/>', 6), "a:ext/@cx is ST_PositiveCoordinate."),
    Case("rpptx.bare_int_font_size", "rpptx", "a bare int where a length is expected (font.size = 12)",
         "analogue (rdocx pptx SKILL.md rule 8 says it raises)",
         ind('box.text_frame.paragraphs[0].runs[0].font.size = 12', 4),
         ind('box.text_frame.paragraphs[0].runs[0].font.size = Pt(12)', 4),
         grep("ppt/slides/", r"<a:rPr [^>]*>"), "a:rPr/@sz is ST_TextFontSize (100 to 400000)."),
    Case("rpptx.replace_xml_malformed", "rpptx", "malformed XML through replace_xml", "analogue",
         ind(f'box.element.replace_xml(\'{sp_xml()[:-8]}\')', 4),
         ind(f'box.element.replace_xml(\'{sp_xml()}\')', 4),
         None, "XML 1.0 well-formedness."),
    Case("rpptx.replace_xml_misordered", "rpptx", "p:spPr before p:nvSpPr through replace_xml", "analogue",
         ind(f'box.element.replace_xml(\'{sp_xml(order="spPr-first")}\')', 4),
         ind(f'box.element.replace_xml(\'{sp_xml()}\')', 4),
         grep("ppt/slides/", r"<p:sp>.{0,40}", 3), "CT_Shape: nvSpPr first."),
    Case("rpptx.replace_xml_unknown_element", "rpptx", "an unknown p: element through replace_xml", "analogue",
         ind(f'box.element.replace_xml(\'{sp_xml().replace("<p:txBody>", "<p:bogus/><p:txBody>")}\')', 4),
         ind(f'box.element.replace_xml(\'{sp_xml()}\')', 4),
         grep("ppt/slides/", r"<p:bogus/>|<p:txBody>", 3), "pml.xsd has no p:bogus."),
    Case("rpptx.replace_xml_bad_enum", "rpptx", "an invalid enum value through replace_xml", "analogue",
         ind("box.element.replace_xml('" + sp_xml(ppr='<a:pPr algn="middle"/>') + "')", 4),
         ind("box.element.replace_xml('" + sp_xml(ppr='<a:pPr algn="ctr"/>') + "')", 4),
         grep("ppt/slides/", r"<a:pPr [^>]*/>"), 'ST_TextAlignType has no "middle".'),
    Case("rpptx.company_ampersand", "rpptx", "'&' in the company property (the pptxgenjs fault, on rpptx)", "analogue",
         ind('prs.app_properties.company = "Harrow & Sons"', 4),
         ind('prs.app_properties.company = "Harrow and Sons"', 4),
         grep("docProps/app.xml", r"<Company>[^<]*</Company>"), "XML 1.0 (must be escaped as &amp;)."),
    Case("rpptx.layout_name_ampersand", "rpptx", "'&' in a layout name (the pptxgenjs fault, on rpptx)", "analogue",
         ind('prs.slide_layouts[5].name = "Harrow & Sons"', 4),
         ind('prs.slide_layouts[5].name = "Harrow and Sons"', 4),
         grep("ppt/slideLayouts/", r'<p:cSld name="Harrow[^"]*"'), "XML 1.0 (must be escaped as &amp;)."),
    Case("rpptx.stale_paragraph_handle", "rpptx", "a paragraph handle used after text_frame.text retired it", "analogue",
         ind('para = box.text_frame.paragraphs[0]\nbox.text_frame.text = "new text"\npara.text = "changed"', 4),
         ind('para = box.text_frame.paragraphs[0]\nbox.text_frame.text = "new text"\nbox.text_frame.paragraphs[0].text = "changed"', 4),
         first_text_contains("changed"), "Not a validity question: is the edit applied, refused, or lost?"),
    Case("rpptx.replace_count_mismatch", "rpptx", "a replacement whose expected count is wrong (CLI)",
         "analogue (rdocx pptx SKILL.md rule 2)", "rpptx replace IN -p 'Quarterly review' -v 'Annual review' --expect 2 -o OUT",
         "rpptx replace IN -p 'Quarterly review' -v 'Annual review' --expect 1 -o OUT",
         first_text_contains("Annual review"), "Not a validity question: is a wrong count refused?",
         cli=[common.RPPTX, "replace", "IN", "-p", "Quarterly review", "-v", "Annual review", "--expect", "EXPECT", "-o", "OUT"]),

    # ---------------- python-docx analogues -------------------------------------------------------------------
    Case("python-docx.colour_hash", "python-docx", "RGBColor.from_string with '#'", "analogue",
         ind('p.runs[0].font.color.rgb = RGBColor.from_string("#FF0000")', 4),
         ind('p.runs[0].font.color.rgb = RGBColor.from_string("FF0000")', 4),
         grep("word/document.xml", r"<w:color [^>]*/>"), "w:color/@w:val is ST_HexColor."),
    Case("python-docx.colour_8digit", "python-docx", "RGBColor.from_string with 8 hex digits", "analogue",
         ind('p.runs[0].font.color.rgb = RGBColor.from_string("FF000080")', 4),
         ind('p.runs[0].font.color.rgb = RGBColor.from_string("FF0000")', 4),
         grep("word/document.xml", r"<w:color [^>]*/>"), "w:color/@w:val is ST_HexColor."),
    Case("python-docx.negative_picture_width", "python-docx", "a negative picture width", "analogue",
         ind('doc.add_picture(PNG, width=-Inches(1))', 4), ind('doc.add_picture(PNG, width=Inches(1))', 4),
         grep("word/document.xml", r"<wp:extent [^>]*/>"), "wp:extent/@cx is ST_PositiveCoordinate (>= 0)."),
    Case("python-docx.bare_int_font_size", "python-docx", "a bare int where a length is expected (font.size = 12)", "analogue",
         ind('p.runs[0].font.size = 12', 4), ind('p.runs[0].font.size = Pt(12)', 4),
         grep("word/document.xml", r"<w:sz [^>]*/>"), "ST_HpsMeasure: w:sz 0 is schema-valid but makes the text invisible (judgement: valid, silently wrong)."),
    Case("python-docx.lxml_malformed", "python-docx", "malformed XML through parse_xml", "analogue",
         ind('p._p.append(parse_xml(\'<w:r %s><w:t>x</w:t>\' % nsdecls("w")))', 4),
         ind('p._p.append(parse_xml(\'<w:r %s><w:t>x</w:t></w:r>\' % nsdecls("w")))', 4),
         None, "XML 1.0 well-formedness."),
    Case("python-docx.lxml_misordered", "python-docx", "w:pPr after w:r through lxml", "analogue",
         ind('p._p.append(parse_xml(\'<w:pPr %s><w:jc w:val="center"/></w:pPr>\' % nsdecls("w")))', 4),
         ind('p._p.insert(0, parse_xml(\'<w:pPr %s><w:jc w:val="center"/></w:pPr>\' % nsdecls("w")))', 4),
         grep("word/document.xml", r"<w:p>.{0,90}"), "CT_P: pPr first."),
    Case("python-docx.lxml_unknown_element", "python-docx", "an unknown w: element through lxml", "analogue",
         ind('p._p.append(OxmlElement("w:bogus"))', 4), ind('p._p.append(OxmlElement("w:r"))', 4),
         grep("word/document.xml", r"<w:bogus/>"), "wml.xsd has no w:bogus."),
    Case("python-docx.unknown_style", "python-docx", "an unknown style name", "analogue",
         ind('doc.add_paragraph("x", style="No Such Style")', 4), ind('doc.add_paragraph("x", style="Heading 1")', 4),
         grep("word/document.xml", r"<w:pStyle [^>]*/>"), "Not a schema rule."),
    Case("python-docx.stale_run_handle", "python-docx", "a run handle used after p.text replaced the runs", "analogue",
         ind('r = p.runs[0]\np.text = "new text"\nr.text = "changed"', 4),
         ind('r = p.runs[0]\np.text = "new text"\np.runs[0].text = "changed"', 4),
         first_text_contains("changed"), "Not a validity question: is the edit applied, refused, or lost?"),

    # ---------------- python-pptx analogues -------------------------------------------------------------------
    Case("python-pptx.colour_hash", "python-pptx", "RGBColor.from_string with '#'", "analogue",
         ind('box.fill.solid()\nbox.fill.fore_color.rgb = RGBColor.from_string("#FF0000")', 4),
         ind('box.fill.solid()\nbox.fill.fore_color.rgb = RGBColor.from_string("FF0000")', 4),
         grep("ppt/slides/", r'<a:srgbClr val="[^"]*"'), "a:srgbClr/@val is ST_HexColorRGB."),
    Case("python-pptx.colour_8digit", "python-pptx", "RGBColor.from_string with 8 hex digits", "analogue",
         ind('box.fill.solid()\nbox.fill.fore_color.rgb = RGBColor.from_string("00000020")', 4),
         ind('box.fill.solid()\nbox.fill.fore_color.rgb = RGBColor.from_string("000000")', 4),
         grep("ppt/slides/", r'<a:srgbClr val="[^"]*"'), "a:srgbClr/@val is ST_HexColorRGB."),
    Case("python-pptx.negative_width", "python-pptx", "a negative shape width", "analogue",
         ind('slide.shapes.add_textbox(Inches(1), Inches(4), -Inches(3), Inches(1)).text_frame.text = "neg"', 4),
         ind('slide.shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1)).text_frame.text = "neg"', 4),
         grep("ppt/slides/", r'<a:ext cx="-?\d+" cy="\d+"/>', 6), "a:ext/@cx is ST_PositiveCoordinate."),
    Case("python-pptx.bare_int_font_size", "python-pptx", "a bare int where a length is expected (font.size = 12)", "analogue",
         ind('box.text_frame.paragraphs[0].runs[0].font.size = 12', 4),
         ind('box.text_frame.paragraphs[0].runs[0].font.size = Pt(12)', 4),
         grep("ppt/slides/", r"<a:rPr [^>]*>"), "a:rPr/@sz is ST_TextFontSize (100 to 400000)."),
    Case("python-pptx.lxml_misordered", "python-pptx", "p:spPr moved before p:nvSpPr through lxml", "analogue",
         ind('sp = box._element\nsp.insert(0, sp.spPr)', 4), ind('sp = box._element\nsp.insert(1, sp.spPr)', 4),
         grep("ppt/slides/", r"<p:sp>.{0,40}", 3), "CT_Shape: nvSpPr first."),
    Case("python-pptx.lxml_unknown_element", "python-pptx", "an unknown p: element through lxml", "analogue",
         ind('box._element.addprevious(parse_xml("<p:bogus %s/>" % nsdecls("p")))', 4), None,
         grep("ppt/slides/", r"<p:bogus/>"), "pml.xsd has no p:bogus."),
    Case("python-pptx.stale_run_handle", "python-pptx", "a run handle used after text_frame.text replaced the runs", "analogue",
         ind('r = box.text_frame.paragraphs[0].runs[0]\nbox.text_frame.text = "new text"\nr.text = "changed"', 4),
         ind('r = box.text_frame.paragraphs[0].runs[0]\nbox.text_frame.text = "new text"\nbox.text_frame.paragraphs[0].runs[0].text = "changed"', 4),
         first_text_contains("changed"), "Not a validity question: is the edit applied, refused, or lost?"),

    # ---------------- a control character in text (XML 1.0 cannot represent U+0001 at all) --------------------
    Case("docxjs.control_char_in_text", "docx-js", "U+0001 in a run's text (analogue)", "analogue",
         'doc = new Document({ sections: [{ children: [new Paragraph({ children: [new TextRun("bad\\u0001char")] })] }] });',
         'doc = new Document({ sections: [{ children: [new Paragraph({ children: [new TextRun("bad char")] })] }] });',
         grep("word/document.xml", r"bad[\s\S]{0,8}char"), "XML 1.0 Char production: U+0001 cannot appear, even as &#1;."),
    Case("pptxgenjs.control_char_in_text", "pptxgenjs", "U+0001 in a text box (analogue)", "analogue",
         ind('pres.addSlide().addText("bad\\u0001char", { x: 1, y: 1, w: 3, h: 1, isTextBox: true });', 4),
         ind('pres.addSlide().addText("bad char", { x: 1, y: 1, w: 3, h: 1, isTextBox: true });', 4),
         grep("ppt/slides/", r"bad[\s\S]{0,8}char"), "XML 1.0 Char production."),
    Case("rdocx.control_char_in_text", "rdocx", "U+0001 in a paragraph's text", "analogue",
         ind('doc.add_paragraph("bad\\x01char")', 4), ind('doc.add_paragraph("bad char")', 4),
         grep("word/document.xml", r"bad[\s\S]{0,8}char"), "XML 1.0 Char production."),
    Case("rpptx.control_char_in_text", "rpptx", "U+0001 in a text frame", "analogue",
         ind('box.text_frame.text = "bad\\x01char"', 4), ind('box.text_frame.text = "bad char"', 4),
         grep("ppt/slides/", r"bad[\s\S]{0,8}char"), "XML 1.0 Char production."),
    Case("python-docx.control_char_in_text", "python-docx", "U+0001 in a paragraph's text", "analogue",
         ind('doc.add_paragraph("bad\\x01char")', 4), ind('doc.add_paragraph("bad char")', 4),
         grep("word/document.xml", r"bad[\s\S]{0,8}char"), "XML 1.0 Char production."),
    Case("python-pptx.control_char_in_text", "python-pptx", "U+0001 in a text frame", "analogue",
         ind('box.text_frame.text = "bad\\x01char"', 4), ind('box.text_frame.text = "bad char"', 4),
         grep("ppt/slides/", r"bad[\s\S]{0,8}char"), "XML 1.0 Char production."),

    # ---------------- the CLI and the input file ----------------------------------------------------------------
    Case("rdocx.cli_output_over_input", "rdocx", "an editing command told to write over its input (CLI)",
         "analogue (rdocx docx SKILL.md rule 1)", "rdocx replace F -p 'Quarterly report' -v 'Annual report' --expect 1 -o F",
         None, first_text_contains("Annual report"), "Not a validity question: does the CLI refuse to overwrite its input?",
         cli=[common.RDOCX, "replace", "OUT", "-p", "Quarterly report", "-v", "Annual report", "--expect", "1", "-o", "OUT"]),
    Case("rpptx.cli_output_over_input", "rpptx", "an editing command told to write over its input (CLI)",
         "analogue (rdocx pptx SKILL.md rule 1)", "rpptx replace F -p 'Quarterly review' -v 'Annual review' --expect 1 -o F",
         None, first_text_contains("Annual review"), "Not a validity question: does the CLI refuse to overwrite its input?",
         cli=[common.RPPTX, "replace", "OUT", "-p", "Quarterly review", "-v", "Annual review", "--expect", "1", "-o", "OUT"]),
]


# --- running a case -------------------------------------------------------------------------------------------
ENVS = {"node": common.env_builtin, "rdocx": common.env_rdocx, "pylib": common.env_pylib}


def _script(case: Case, code: str, folder: Path) -> list:
    runner, tpl = TEMPLATES[case.generator]
    src = tpl.replace("%CALL%", code)
    ext = ".js" if runner == "node" else ".py"
    path = folder / f"script{ext}"
    path.write_text(src)
    exe = {"node": common.NODE, "rdocx": common.R_PY, "pylib": common.PY}[runner]
    return [exe, path], ENVS[runner]()


def _run_variant(case: Case, which: str, folder: Path, png: Path, bases: dict):
    folder.mkdir(parents=True, exist_ok=True)
    out = folder / f"out.{FMT_OF[case.generator]}"
    code = case.call if which == "call" else case.control
    if case.cli:
        expect = "2" if which == "call" else "1"
        src = bases[(FMT_OF[case.generator], "rdocx")]
        inplace = "IN" not in case.cli
        if inplace:  # the command reads and writes OUT: a copy of the base
            out.write_bytes(Path(src).read_bytes())
        before = out.read_bytes() if inplace else None
        argv = [{"IN": src, "OUT": out, "EXPECT": expect}.get(a, a) for a in case.cli]
        r = common.run(argv, env=common.env_rdocx(), timeout=120)
        if inplace:
            status = "refused_at_call" if r["rc"] and out.read_bytes() == before else "written"
        else:
            status = "written" if out.exists() else ("refused_at_call" if r["rc"] else "not_written")
        return {"status": status, "rc": r["rc"], "message": (r["stdout"] + r["stderr"]).strip()[:300], "out": out}
    argv, env = _script(case, code, folder)
    r = common.run(argv + [out, png], env=env, timeout=120)
    text = (r["stdout"] + r["stderr"]).strip()
    if r["rc"] == 3:
        status = "refused_at_call"
    elif r["rc"] == 4:
        status = "refused_at_save"
    elif r["rc"] == 0 and out.exists():
        status = "written"
    else:
        status = "crashed"
    msg = next((ln for ln in text.splitlines() if ln.startswith("REFUSED")), text[-300:])
    return {"status": status, "rc": r["rc"], "message": msg[:300], "out": out}


def run_case(case: Case, work: Path, png: Path, bases: dict):
    folder = work / case.id
    rec = {"id": case.id, "stack": case.stack, "generator": case.generator, "format": FMT_OF[case.generator],
           "footgun": case.footgun, "source": case.source, "call": case.call.strip(),
           "control": (case.control or "").strip(), "spec": case.spec}
    call = _run_variant(case, "call", folder / "call", png, bases)
    ctrl = _run_variant(case, "control", folder / "control", png, bases) if case.control else None
    rec["refused_at_call"] = call["status"] == "refused_at_call"
    rec["refused_at_save"] = call["status"] == "refused_at_save"
    rec["status"] = call["status"]
    rec["message"] = call["message"]
    rec["control_status"] = ctrl["status"] if ctrl else None
    ctrl_written = None
    if ctrl and ctrl["status"] == "written" and case.inspect:
        try:
            ctrl_written = case.inspect(ctrl["out"])
        except Exception as e:
            ctrl_written = f"unreadable: {type(e).__name__}"
    rec["written_by_control"] = ctrl_written
    if call["status"] != "written":
        rec.update(output_invalid=None, written=None, oracle_new=[], checks={}, caught_by=[],
                   outcome=call["status"])
        return rec, None
    out = call["out"]
    base_issues = {norm(i) for i in oracle.audit(ctrl["out"])} if ctrl and ctrl["status"] == "written" else set()
    rec["oracle_new"] = [i for i in oracle.audit(out) if norm(i) not in base_issues]
    try:
        rec["written"] = case.inspect(out) if case.inspect else None
    except Exception as e:  # unreadable package: the oracle says why
        rec["written"] = f"unreadable: {type(e).__name__}: {e}"[:200]
    extra = case.spec_invalid(rec["written"] or "") if case.spec_invalid else False
    rec["output_invalid"] = bool(rec["oracle_new"]) or bool(extra)
    if case.spec_invalid is None and case.id == "pptxgenjs.secondary_axis_without_axes":
        rec["output_invalid"] = None  # filled from the axis check below
    jobs = [(out, ctrl["out"] if ctrl and ctrl["status"] == "written" else None)]
    return rec, jobs


def undeclared_axes(path: Path) -> list[str]:
    """Chart groups whose c:axId names no axis the plot area declares (zipfile + lxml)."""
    from lxml import etree
    C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
    bad = []
    with zipfile.ZipFile(path) as zf:
        for n in zf.namelist():
            if not re.match(r"ppt/charts/chart\d+\.xml$", n):
                continue
            root = etree.fromstring(zf.read(n))
            declared = {ax.find(f"{{{C}}}axId").get("val") for kind in ("catAx", "valAx", "dateAx", "serAx")
                        for ax in root.iter(f"{{{C}}}{kind}")}
            for grp in root.iter():
                if isinstance(grp.tag, str) and grp.tag.endswith("Chart") and grp.tag.startswith(f"{{{C}}}") \
                        and grp.find(f"{{{C}}}axId") is not None:
                    ids = [a.get("val") for a in grp.findall(f"{{{C}}}axId")]
                    missing = [i for i in ids if i not in declared]
                    if missing:
                        bad.append(f"{n}: {etree.QName(grp).localname} axId {missing} not declared")
    return bad


# --- atomic save ----------------------------------------------------------------------------------------------
LIMIT = 4096


def _limited():
    signal.signal(signal.SIGXFSZ, signal.SIG_IGN)
    resource.setrlimit(resource.RLIMIT_FSIZE, (LIMIT, LIMIT))


SAVE_CASES = [
    ("rdocx", "docx", "Document.save over the input", common.R_PY,
     "import rdocx, sys; d = rdocx.Document(sys.argv[1]); d.add_paragraph('more ' * 4000); d.save(sys.argv[1])"),
    ("rdocx", "docx", "docx_ops.save_atomic(doc, out, src) over an existing output", common.R_PY,
     "import rdocx, sys, docx_ops; d = rdocx.Document(sys.argv[2]); d.add_paragraph('more ' * 4000); docx_ops.save_atomic(d, sys.argv[1], sys.argv[2])"),
    ("rdocx", "pptx", "Presentation.save over the input", common.R_PY,
     "import rpptx, sys; p = rpptx.Presentation(sys.argv[1]); p.slides.add_slide(p.slide_layouts[5]).shapes.title.text = 'more ' * 4000; p.save(sys.argv[1])"),
    ("rdocx", "pptx", "pptx_ops.save_atomic(prs, out, src) over an existing output", common.R_PY,
     "import rpptx, sys, pptx_ops; p = rpptx.Presentation(sys.argv[2]); p.slides.add_slide(p.slide_layouts[5]).shapes.title.text = 'more ' * 4000; pptx_ops.save_atomic(p, sys.argv[1], sys.argv[2])"),
    ("pylib", "docx", "python-docx Document.save over the input", common.PY,
     "import docx, sys; d = docx.Document(sys.argv[1]); d.add_paragraph('more ' * 4000); d.save(sys.argv[1])"),
    ("pylib", "pptx", "python-pptx Presentation.save over the input", common.PY,
     "import pptx, sys; p = pptx.Presentation(sys.argv[1]); p.slides.add_slide(p.slide_layouts[5]).shapes.title.text = 'more ' * 4000; p.save(sys.argv[1])"),
    ("builtin", "docx", "docx-js Packer.toBuffer + fs.writeFileSync over an existing output", common.NODE,
     "const fs = require('fs'); const { Document, Packer, Paragraph } = require('docx'); "
     "Packer.toBuffer(new Document({ sections: [{ children: [new Paragraph('more '.repeat(4000))] }] }))"
     ".then((b) => fs.writeFileSync(process.argv[1], b)).catch((e) => { console.log(e.message); process.exit(1); });"),
    ("builtin", "pptx", "pptxgenjs writeFile over an existing output", common.NODE,
     "const pptxgen = require('pptxgenjs'); const p = new pptxgen(); p.addSlide().addText('more '.repeat(4000), { x: 1, y: 1, w: 8, h: 3 }); "
     "p.writeFile({ fileName: process.argv[1] }).catch((e) => { console.log(e.message); process.exit(1); });"),
    ("builtin", "docx", "built-in edit flow: rm -f ../out.docx && zip -Xr ../out.docx . (over an existing output)", "zip", None),
]


def atomic_save(work: Path, bases: dict):
    results = []
    for stack, fmt, what, exe, code in SAVE_CASES:
        folder = work / "atomic" / re.sub(r"\W+", "_", what)[:60]
        folder.mkdir(parents=True, exist_ok=True)
        src = folder / f"input.{fmt}"
        src.write_bytes(bases[(fmt, stack if stack != "builtin" else "builtin")].read_bytes())
        before = src.read_bytes()
        target = src
        if "existing output" in what:
            target = folder / f"previous-output.{fmt}"
            target.write_bytes(before)
        if exe == "zip":
            unpacked = folder / "unpacked"
            with zipfile.ZipFile(src) as zf:
                zf.extractall(unpacked)
            argv, cwd = ["sh", "-c", f"rm -f ../{target.name} && zip -qXr ../{target.name} ."], unpacked
            env = common.env_builtin()
        else:
            argv = [exe, "-c" if exe != common.NODE else "-e", code, target] + ([src] if target != src else [])
            cwd = folder
            env = common.env_rdocx() if exe == common.R_PY else (common.env_builtin() if exe == common.NODE else common.env_pylib())
        t0 = time.perf_counter()
        p = subprocess.run([str(a) for a in argv], cwd=cwd, env=env, capture_output=True, text=True,
                           preexec_fn=_limited, timeout=120)
        after = target.read_bytes() if target.exists() else None
        intact = after == before
        if intact:
            state = "unchanged"
        elif after is None:
            state = "missing"
        else:
            issues = oracle.audit(target)
            state = (f"replaced by a smaller package that opens as a ZIP ({len(after)} bytes) but breaks "
                     f"{len(issues)} rule(s), e.g. {issues[0][:120]}" if issues and not issues[0].startswith("zip")
                     else f"corrupt ({len(after)} bytes: {issues[0][:80] if issues else 'unreadable'})")
        leftovers = sorted(x.name for x in folder.iterdir() if x.name.startswith(".") or x.name.endswith(".tmp"))
        text = (p.stdout + p.stderr).strip()
        results.append({
            "stack": stack, "format": fmt, "save": what, "limit_bytes": LIMIT,
            "failed": p.returncode != 0, "error": text.splitlines()[-1][:200] if text else "",
            "input_intact": src.read_bytes() == before,
            "target": "the input" if target == src else "an existing output file",
            "target_intact": intact, "target_state": state,
            "temp_files_left": leftovers, "seconds": round(time.perf_counter() - t0, 3),
        })
    return results


def run_all(work: Path, png: Path, jobs: int = 12, bases: dict | None = None):
    work.mkdir(parents=True, exist_ok=True)
    if bases is None:  # the base files of Part 1
        bases = {(f, s): common.WORK / "validity" / "bases" / f"{s}.{f}" for f in ("docx", "pptx")
                 for s in ("rdocx", "builtin", "pylib")}
    with ThreadPoolExecutor(max_workers=jobs) as pool:
        done = list(pool.map(lambda c: run_case(c, work, png, bases), CASES))
    files = []
    for rec, pair in done:
        if pair:
            out, ctrl = pair[0]
            files.append(((rec["id"], "call"), rec["format"], out))
            if ctrl:
                files.append(((rec["id"], "control"), rec["format"], ctrl))
    from run import run_checkers
    checked = run_checkers(jobs, files)
    records = []
    for rec, pair in done:
        if pair:
            calls = checked[(rec["id"], "call")]
            ctrls = checked.get((rec["id"], "control"), {})
            rec["checks"] = {k: {**v, "flagged_on_control": ctrls.get(k, {}).get("flagged")} for k, v in calls.items()}
            rec["caught_by"] = [k for k, v in calls.items() if v["flagged"] and not ctrls.get(k, {}).get("flagged", False)]
            if rec["id"] == "pptxgenjs.secondary_axis_without_axes":
                bad = undeclared_axes(pair[0][0])
                rec["written"] = (rec["written"] or "") + (" | undeclared: " + "; ".join(bad) if bad else " | every axId declared")
                rec["output_invalid"] = bool(bad) or bool(rec["oracle_new"])
        if rec["status"] == "written":
            rec["outcome"] = "written_invalid" if rec["output_invalid"] else "written_valid"
            # same as the correct call: the library fixed the mistake silently (or it has no effect on the file);
            # different: an option ignored, a wrong value, a lost edit, or only a cosmetic difference
            rec["same_as_control"] = (rec["written"] == rec["written_by_control"]
                                      if rec["control_status"] == "written" else None)
        else:
            rec["same_as_control"] = None
        records.append(rec)
    saves = atomic_save(work, bases)
    return {"cases": records, "atomic_save": saves}


if __name__ == "__main__":
    out = run_all(common.WORK / "validity" / "prevention", common.WORK / "validity" / "picture.png")
    print(json.dumps(out, indent=1)[:20000])
