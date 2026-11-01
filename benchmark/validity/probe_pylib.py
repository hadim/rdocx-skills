"""Open a file with python-docx / python-pptx and iterate its content (the "pylib_open" checker).

Usage: PY probe_pylib.py FILE   (.docx/.dotx with python-docx, .pptx/.potx with python-pptx)
Exit 0 and print a short summary when the file loads and iterates; exit 1 with the error otherwise.
python-docx refuses a .dotx content type by design; that is recorded as a flag like any other.
"""

import sys

path = sys.argv[1]
try:
    if path.lower().endswith((".docx", ".dotx")):
        import docx

        doc = docx.Document(path)
        n = sum(len(p.text) for p in doc.paragraphs)
        n += sum(len(c.text) for t in doc.tables for row in t.rows for c in row.cells)
        pics = len(doc.inline_shapes)
        comments = len(list(doc.comments))
        for s in doc.sections:
            n += sum(len(p.text) for p in s.header.paragraphs)
        print(f"ok: {len(doc.paragraphs)} paragraphs, {pics} pictures, {comments} comments, {n} chars")
    else:
        import pptx

        prs = pptx.Presentation(path)
        n = 0
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    n += len(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    n += sum(len(c.text) for row in shape.table.rows for c in row.cells)
                if getattr(shape, "has_chart", False) and shape.has_chart:
                    n += len(list(shape.chart.plots[0].categories))
            if slide.has_notes_slide:
                n += len(slide.notes_slide.notes_text_frame.text)
            n += len(slide.slide_layout.name)
        print(f"ok: {len(prs.slides)} slides, {n} chars")
except Exception as e:
    print(f"{type(e).__name__}: {e}"[:600])
    sys.exit(1)
