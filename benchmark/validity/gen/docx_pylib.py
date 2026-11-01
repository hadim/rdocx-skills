"""Base .docx made with python-docx (the pylib stack): styles, a table, a picture, a comment, a tracked change.

Usage: PY docx_pylib.py OUT.docx PICTURE.png
python-docx 1.2 has `add_comment`; it has no API for tracked changes, so the w:del / w:ins pair is built with
its own oxml helpers (lxml), as an agent without a skill would.
"""

import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

out, png = sys.argv[1], sys.argv[2]

doc = Document()
doc.core_properties.title = "Validity base"
doc.core_properties.author = "bench"
note = doc.styles.add_style("Note box", WD_STYLE_TYPE.PARAGRAPH)
note.base_style = doc.styles["Normal"]
note.font.italic = True
note.paragraph_format.left_indent = Inches(0.5)

doc.add_heading("Quarterly report", 1)
doc.add_paragraph("A styled note paragraph.", style="Note box")
p = doc.add_paragraph()
run = p.add_run("Revenue grew by 12 percent over the quarter.")
doc.add_comment(run, text="Check this figure.", author="Reviewer", initials="RV")

p = doc.add_paragraph("The budget is ")


def revision(tag, text, rid):
    wrap = OxmlElement(tag)
    wrap.set(qn("w:id"), str(rid))
    wrap.set(qn("w:author"), "Reviewer")
    wrap.set(qn("w:date"), "2026-10-01T09:00:00Z")
    r = OxmlElement("w:r")
    t = OxmlElement("w:delText" if tag == "w:del" else "w:t")
    t.text = text
    r.append(t)
    wrap.append(r)
    return wrap


p._p.append(revision("w:del", "ten", 101))
p._p.append(revision("w:ins", "twelve", 102))
p.add_run(" thousand euros.")

t = doc.add_table(rows=3, cols=2, style="Table Grid")
for r, row in enumerate([("Item", "Value"), ("North", "41"), ("South", "37")]):
    for c, text in enumerate(row):
        t.cell(r, c).text = text
doc.add_paragraph("Figure 1: a picture.")
doc.add_picture(png, width=Inches(1.25))
doc.save(out)
