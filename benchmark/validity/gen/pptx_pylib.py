"""Base .pptx made with python-pptx (the pylib stack): two layouts, a table, a chart, a picture, notes.

Usage: PY pptx_pylib.py OUT.pptx PICTURE.png
"""

import sys

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

out, png = sys.argv[1], sys.argv[2]

prs = Presentation()
prs.core_properties.title = "Validity base"
prs.core_properties.author = "bench"

s1 = prs.slides.add_slide(prs.slide_layouts[0])
s1.shapes.title.text = "Quarterly review"
s1.placeholders[1].text = "Validity benchmark base deck"
s1.notes_slide.notes_text_frame.text = "Open with the headline number."

s2 = prs.slides.add_slide(prs.slide_layouts[5])
s2.shapes.title.text = "Regional figures"
frame = s2.shapes.add_table(3, 2, Inches(1), Inches(1.6), Inches(6), Inches(1.5))
for r, row in enumerate([("Region", "Value"), ("North", "41"), ("South", "37")]):
    for c, text in enumerate(row):
        frame.table.cell(r, c).text = text
s2.shapes.add_picture(png, Inches(7.5), Inches(1.6), width=Inches(1.5))
s2.notes_slide.notes_text_frame.text = "Mention the photo record."

s3 = prs.slides.add_slide(prs.slide_layouts[5])
s3.shapes.title.text = "Costs"
data = CategoryChartData()
data.categories = ["Option A", "Option B", "Option C"]
data.add_series("Cost (EUR k)", (236, 230, 251))
s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5), data)

prs.save(out)
