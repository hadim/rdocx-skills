"""Comment run positions skip runs inside an inline content control (w:sdt)."""
import json
import re
import subprocess
import zipfile

import docx
import rdocx
from docx.oxml.ns import qn

# One paragraph: "before " | <w:sdt>"TARGET"</w:sdt> | " after"
d = docx.Document()
p = d.add_paragraph()
p.add_run("before ")
sdt = p._p.makeelement(qn("w:sdt"), {})
sdt.append(sdt.makeelement(qn("w:sdtPr"), {}))
content = sdt.makeelement(qn("w:sdtContent"), {})
r = p.add_run("TARGET")._r
p._p.remove(r)
content.append(r)
sdt.append(content)
p._p.append(sdt)
p.add_run(" after")
d.save("in.docx")


def anchored(path):
    xml = zipfile.ZipFile(path).read("word/document.xml").decode()
    m = re.search(r'<w:commentRangeStart w:id="(\d+)"/>(.*?)<w:commentRangeEnd w:id="\1"/>', xml, re.S)
    return "".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", m.group(2)))


doc = rdocx.Document.open("in.docx")
print("Paragraph.runs:", [run.text for run in doc.paragraphs[0].runs])
view = json.loads(subprocess.run(["rdocx", "text", "in.docx", "--json"], check=True, capture_output=True, text=True).stdout)
print("text --json runs:", [run["text"] for run in view["paragraphs"][0]["runs"]])
doc.add_comment(
    rdocx.RunRange(start=rdocx.RunPosition(body_index=0, run_index=1), end=rdocx.RunPosition(body_index=0, run_index=2)),
    author="A",
    text="x",
)
doc.save("py.docx")
print("python, runs [1, 2):", repr(anchored("py.docx")), "expected 'TARGET'")

subprocess.run(
    ["rdocx", "comment", "add", "in.docx", "--start-paragraph", "0", "--start-run", "1", "--end-paragraph", "0",
     "--end-run", "2", "--author", "A", "--text", "x", "-o", "cli.docx"],
    check=True, capture_output=True,
)
print("cli, --start-run 1 --end-run 2:", repr(anchored("cli.docx")), "expected 'TARGET'")

try:
    rdocx.Document.open("in.docx").add_comment(
        rdocx.RunRange(start=rdocx.RunPosition(body_index=0, run_index=2), end=rdocx.RunPosition(body_index=0, run_index=3)),
        author="A",
        text="x",
    )
    print("python, runs [2, 3): accepted")
except Exception as exc:
    print("python, runs [2, 3):", type(exc).__name__, exc)
