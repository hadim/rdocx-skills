import io, zipfile
import rdocx
from docx import Document
from docx.oxml.ns import qn
from PIL import Image


def build(path, default_ns):
    d = Document()
    d.add_paragraph("Before the control.")
    body = d.element.body
    sdt = body.makeelement(qn("w:sdt"), {})
    pr = sdt.makeelement(qn("w:sdtPr"), {})
    tag = pr.makeelement(qn("w:tag"), {}); tag.set(qn("w:val"), "goog_rdk_0"); pr.append(tag); sdt.append(pr)
    content = sdt.makeelement(qn("w:sdtContent"), {})
    p = d.add_paragraph("Inside the control.")._p; body.remove(p); content.append(p); sdt.append(content)
    body.insert(1, sdt)
    d.save(path)
    if default_ns:  # as Google Docs exports write the root
        items = [(i, zipfile.ZipFile(path).read(i.filename)) for i in zipfile.ZipFile(path).infolist()]
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for info, data in items:
                if info.filename == "word/document.xml":
                    data = data.replace(b"<w:document ", b'<w:document xmlns="http://schemas.microsoft.com/office/tasks/2019/documenttasks" ', 1)
                z.writestr(info, data)
    return path


png = io.BytesIO(); Image.new("RGB", (60, 40)).save(png, "PNG")
for default_ns in (False, True):
    doc = rdocx.Document.open(build("sdt.docx", default_ns))
    try:
        doc.add_picture(png.getvalue(), "x.png", width=rdocx.Inches(1), height=rdocx.Inches(1))
        result = "ok"
    except rdocx.RdocxError as e:
        result = f"RdocxError: {e}"
    print(f"content control, default namespace on the root: {default_ns!s:<5} -> add_picture {result}")
