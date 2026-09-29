"""Every python and bash block of the skills' recipes.md runs on the pinned build, in order, in one folder
per file that starts with the fixtures as `report.docx` and `deck.pptx`. A recipe that stops working after
a pin change fails here, before an agent follows it."""
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

from conftest import BIN

ROOT = Path(__file__).resolve().parent.parent
BLOCKS = []
for recipes in sorted(ROOT.glob("skills/*/references/recipes.md")):
    for k, (lang, code) in enumerate(re.findall(r"```(python|bash)\n(.*?)```", recipes.read_text(), re.S)):
        BLOCKS.append(pytest.param(recipes, lang, code, id=f"{recipes.parent.parent.name}-{k:02d}-{lang}"))


@pytest.fixture(scope="module")
def workdirs(tmp_path_factory, report_docx, deck_pptx):
    dirs = {}
    for recipes in ROOT.glob("skills/*/references/recipes.md"):
        d = tmp_path_factory.mktemp(recipes.parent.parent.name)
        shutil.copy(report_docx, d / "report.docx")
        shutil.copy(deck_pptx, d / "deck.pptx")
        dirs[recipes] = d
    return dirs


@pytest.mark.parametrize("recipes,lang,code", BLOCKS)
def test_recipe_block(recipes, lang, code, workdirs):
    skill = recipes.parent.parent
    # the runtime Python an agent gets (BIN/python), not the test environment: no pytest, no Pillow there
    env = dict(os.environ, R=str(BIN), SKILL=str(skill), RDOCX_BIN_DIR=str(BIN), PYTHONPATH=str(skill / "scripts"))
    cmd = [str(BIN / "python"), "-c", code] if lang == "python" else ["bash", "-euo", "pipefail", "-c", code]
    res = subprocess.run(cmd, cwd=workdirs[recipes], env=env, capture_output=True, text=True, timeout=300)
    assert res.returncode == 0, res.stdout[-2000:] + res.stderr[-2000:]


def test_recipe_outputs(workdirs):
    """The recipes' outputs hold what the recipes say (runs after every block, in the same folders)."""
    import io
    import zipfile

    import pptx
    import rdocx
    sys.path.insert(0, str(ROOT / "skills" / "docx" / "scripts"))
    import docx_ops
    d = {r.parent.parent.name: w for r, w in workdirs.items()}["docx"]
    texts = [p.text for p in rdocx.Document.open(d / "edited.docx").paragraphs]
    assert any("Issue C, for approval" in t for t in texts) and any("two points lower" in t for t in texts)
    doc = rdocx.Document.open(d / "commented.docx")
    first = next(c for c in doc.comments if c.parent_id is None)
    assert docx_ops.anchored_text(d / "commented.docx", first.id) == "about 12 mm" and first.resolved and first.date
    assert any(c.text == "Rename this column?" and c.date for c in doc.comments)
    red = rdocx.Document.open(d / "redline.docx")
    assert [(r.kind, r.story.kind) for r in red.revisions] == [("deletion", "body"), ("insertion", "body")]
    bold = rdocx.Document.open(d / "bold.docx")
    assert any(r.text == "safety of path users" and r.font.bold for p in bold.paragraphs for r in p.runs)
    assert not any(t.startswith("Access equipment, traffic management")
                   for t in (p.text for p in rdocx.Document.open(d / "restructured.docx").paragraphs))
    assert any("Replace the handrail caps" == c.text for t in rdocx.Document.open(d / "table.docx").tables
               for row in t.rows for c in row.cells)
    with zipfile.ZipFile(d / "from-template.docx") as z:
        assert b"wordprocessingml.document.main+xml" in z.read("[Content_Types].xml")
    assert [p.text for p in rdocx.Document.open(d / "note.docx").paragraphs] == ["Site visit note", "Visited on 27 September 2026."]
    p = {r.parent.parent.name: w for r, w in workdirs.items()}["pptx"]
    texts = "\n".join(sh.text_frame.text for s in pptx.Presentation(p / "edited.pptx").slides for sh in s.shapes if sh.has_text_frame)
    assert "EUR 236,000" in texts and "eight weeks" in texts
    shapes = pptx.Presentation(p / "shapes.pptx").slides[3].shapes
    badge = next(sh for sh in shapes if sh.has_text_frame and sh.text_frame.text == "Decision")
    assert str(badge.fill.fore_color.rgb) == "7B1E3A"
    re_deck = pptx.Presentation(p / "reordered.pptx")
    assert re_deck.slides[2].notes_slide.notes_text_frame.text == "Mention the photo record for each defect."
    assert re_deck.slides[-1].shapes.title.text == "Questions"
    with zipfile.ZipFile(p / "from-template.pptx") as z:
        assert b"presentationml.presentation.main+xml" in z.read("[Content_Types].xml")
