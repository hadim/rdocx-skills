"""docx: what the pinned build lacks of the everyday surface the skills used to cite, each as a strict xfail
(tests/gaps.py): python-docx's names and signatures, the refusal of bare-int lengths, list restarts, content
controls set as Word does, note references and removal, raw XML checks, charts read back, display equations, and the
values Word ignores that are accepted silently. Each test builds its own input."""
import datetime
import re

import docx
import pytest
import rdocx

from conftest import part, run

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def saved(doc, path):
    doc.save(path)
    return part(path, "word/document.xml").decode()


# ---------------------------------------------------------------- python-docx's names (what tensorbee/rdocx#379 left)
@pytest.mark.gap("python-docx-names")
def test_list_styles_of_python_docx_template(tmp_path):
    """python-docx's default template has List Bullet and List Number; a new rdocx document lacks them (KeyError)."""
    doc = rdocx.Document()
    doc.add_paragraph("First point", style="List Bullet")
    doc.add_paragraph("Step one", style="List Number")
    doc.save(tmp_path / "l.docx")
    assert [p.style.name for p in docx.Document(tmp_path / "l.docx").paragraphs] == ["List Bullet", "List Number"]


@pytest.mark.gap("python-docx-names")
def test_page_margins_the_python_docx_way(tmp_path):
    """doc.sections[i] is a snapshot: its header refuses and names doc.section(i), but setting a margin on it raises
    a bare AttributeError. Either set it as python-docx does or name the way."""
    doc = rdocx.Document()
    doc.add_paragraph("Body")
    try:
        doc.sections[0].left_margin = rdocx.Inches(1.5)
    except (AttributeError, NotImplementedError) as e:
        assert "section(" in str(e)
        return
    doc.save(tmp_path / "h.docx")
    assert docx.Document(tmp_path / "h.docx").sections[0].left_margin == docx.shared.Inches(1.5)


@pytest.mark.gap("python-docx-names")
def test_colour_type_enum_the_python_docx_way():
    from rdocx.enum.dml import MSO_COLOR_TYPE                      # python-docx: docx.enum.dml
    doc = rdocx.Document()
    doc.add_paragraph("Line")
    doc.paragraphs[0].runs[0].font.color.rgb = rdocx.RGBColor(0x7B, 0x1E, 0x3A)
    assert doc.paragraphs[0].runs[0].font.color.type == MSO_COLOR_TYPE.RGB


@pytest.mark.gap("python-docx-names")
def test_tab_stops_the_python_docx_way():
    doc = rdocx.Document()
    stops = doc.add_paragraph("Name\tValue").paragraph_format.tab_stops
    stops.add_tab_stop(rdocx.Inches(3))
    stops[0].alignment = rdocx.WD_TAB_ALIGNMENT.RIGHT                 # a live TabStop, as python-docx
    assert doc.paragraphs[0].paragraph_format.tab_stops[0].alignment == rdocx.WD_TAB_ALIGNMENT.RIGHT


@pytest.mark.gap("python-docx-names")
def test_a_python_docx_private_points_at_raw_xml():
    """Code written for python-docx reaches for `paragraph._p`: the error names the way to the XML."""
    doc = rdocx.Document()
    doc.add_paragraph("x")
    with pytest.raises(AttributeError, match="replace_xml|xml"):
        doc.paragraphs[0]._p


@pytest.mark.gap("python-docx-names")
def test_custom_properties_as_a_mapping(tmp_path):
    doc = rdocx.Document()
    doc.custom_properties["Client"] = "Acme"
    doc.custom_properties["Revision"] = 3
    doc.custom_properties["Approved"] = True
    doc.custom_properties["Due"] = datetime.datetime(2026, 10, 31, tzinfo=datetime.timezone.utc)
    del doc.custom_properties["Approved"]
    doc.save(tmp_path / "c.docx")
    back = rdocx.Document.open(tmp_path / "c.docx")
    assert dict(back.custom_properties.items()) == {"Client": "Acme", "Revision": 3,
                                                    "Due": datetime.datetime(2026, 10, 31, tzinfo=datetime.timezone.utc)}
    assert "Client" in back.custom_properties.keys() and back.custom_properties.get("Nope") is None


# ---------------------------------------------------------------- bare-int lengths
@pytest.mark.gap("docx-bare-int-length-saved-as-zero")
def test_a_bare_int_length_raises():
    """A bare integer is read as EMU: space_after = 6 would save as w:after="0". It raises and names Pt (or Twips for
    page lengths), as font.size = 12 does."""
    doc = rdocx.Document()
    doc.add_paragraph("x")
    with pytest.raises(ValueError, match="Pt"):
        doc.paragraphs[0].paragraph_format.space_after = 6
    with pytest.raises(ValueError, match="Pt"):
        doc.paragraphs[0].runs[0].font.character_spacing = 2
    with pytest.raises(ValueError, match="Twips"):
        doc.update_section(0, margin_top=1440)


@pytest.mark.gap("docx-bare-int-length-saved-as-zero")
def test_a_bare_number_line_spacing_is_a_multiple(tmp_path):
    """As python-docx: `line_spacing = 2` is double spacing, a length is exact."""
    doc = rdocx.Document()
    doc.add_paragraph("x").paragraph_format.line_spacing = 2
    assert doc.paragraphs[0].paragraph_format.line_spacing == 2.0
    doc.save(tmp_path / "p.docx")
    assert docx.Document(tmp_path / "p.docx").paragraphs[0].paragraph_format.line_spacing == 2.0


# ---------------------------------------------------------------- lists
@pytest.mark.gap("docx-list-numbering-options")
def test_a_numbered_list_restarts_and_a_bullet_level_takes_a_font(tmp_path):
    doc = rdocx.Document()
    doc.add_numbered_list_item("one")
    doc.add_numbered_list_item("two")
    doc.add_numbered_list_item("again one", restart=True)          # Word's "Restart at 1"
    doc.add_numbered_list_item("again two")
    assert doc.paragraphs[2].numbering[0] != doc.paragraphs[1].numbering[0]
    assert doc.paragraphs[3].numbering == doc.paragraphs[2].numbering
    new_id = doc.restart_numbering(doc.paragraphs[3], start=5)
    assert doc.paragraphs[3].numbering == (new_id, 0)
    plain = doc.add_numbering_definition([rdocx.ListLevel(format="bullet", text="-", font="Arial")])
    doc.add_paragraph("dash").numbering = (doc.add_numbering_instance(plain), 0)
    doc.save(tmp_path / "l.docx")
    numbering = part(tmp_path / "l.docx", "word/numbering.xml").decode()
    assert '<w:startOverride w:val="5"/>' in numbering and re.search(r'<w:rFonts [^>]*w:ascii="Arial"', numbering)


# ---------------------------------------------------------------- content controls
def controls_docx(path):
    """Body content controls: a plain text one (tag client), a drop-down (alias Status), a date and a check box."""
    d = docx.Document()
    body = d.element.body
    xml = (f'<w:p xmlns:w="{W}" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">'
           '<w:sdt><w:sdtPr><w:tag w:val="client"/><w:id w:val="11"/><w:showingPlcHdr/><w:text/></w:sdtPr>'
           '<w:sdtContent><w:r><w:t>Click here</w:t></w:r></w:sdtContent></w:sdt>'
           '<w:sdt><w:sdtPr><w:alias w:val="Status"/><w:tag w:val="status"/><w:id w:val="12"/><w:dropDownList>'
           '<w:listItem w:displayText="Draft" w:value="draft"/><w:listItem w:displayText="Final" w:value="final"/>'
           '</w:dropDownList></w:sdtPr><w:sdtContent><w:r><w:t>Draft</w:t></w:r></w:sdtContent></w:sdt>'
           '<w:sdt><w:sdtPr><w:tag w:val="due"/><w:id w:val="13"/><w:date><w:dateFormat w:val="yyyy-MM-dd"/></w:date>'
           '</w:sdtPr><w:sdtContent><w:r><w:t>date</w:t></w:r></w:sdtContent></w:sdt>'
           '<w:sdt><w:sdtPr><w:tag w:val="ok"/><w:id w:val="14"/><w14:checkbox><w14:checked w14:val="0"/></w14:checkbox>'
           '</w:sdtPr><w:sdtContent><w:r><w:t>☐</w:t></w:r></w:sdtContent></w:sdt></w:p>')
    from docx.oxml import parse_xml
    body.insert(0, parse_xml(xml))
    d.save(path)
    return path


@pytest.mark.gap("docx-content-control-value-like-word")
def test_a_content_control_value_is_set_as_word_does(rdocx_cli, tmp_path):
    """The placeholder goes, a date is written as a date, a drop-down takes one of its items (by value or display
    text, showing the display text), and an unknown tag is an error."""
    doc = rdocx.Document.open(controls_docx(tmp_path / "c.docx"))
    doc.set_content_control_value_by_tag("client", "Acme")
    doc.set_content_control_value_by_tag("due", "2026-10-31")
    xml = saved(doc, tmp_path / "f.docx")
    assert "<w:showingPlcHdr/>" not in xml and 'w:fullDate="2026-10-31' in xml
    doc.set_content_control_value_by_tag("status", "final")
    assert doc.content_controls()[1].text == "Final"
    with pytest.raises(ValueError, match="Draft|Final"):
        doc.set_content_control_value_by_alias("Status", "Archived")
    with pytest.raises(KeyError):
        doc.set_content_control_value_by_tag("nope", "x")
    res = run([rdocx_cli, "fill", tmp_path / "c.docx", "--tag", "status=final", "-o", tmp_path / "g.docx"])
    assert res.returncode == 0 and rdocx.Document.open(tmp_path / "g.docx").content_controls()[1].text == "Final"


# ---------------------------------------------------------------- notes
@pytest.mark.gap("docx-note-reference-and-removal")
def test_a_note_reference_is_styled_as_word_writes_it(tmp_path):
    """Word marks a note reference with the FootnoteReference / EndnoteReference character style, which raises it."""
    doc = rdocx.Document()
    doc.add_paragraph("A claim.")
    doc.add_footnote(doc.paragraphs[0], "The source.")
    doc.add_endnote(doc.paragraphs[0], "A closing note.")
    xml = saved(doc, tmp_path / "n.docx")
    assert 'w:val="FootnoteReference"' in xml and 'w:val="EndnoteReference"' in xml


@pytest.mark.gap("docx-note-reference-and-removal")
def test_a_note_is_removed_after_other_edits(rdocx_cli, tmp_path):
    """A note handle retires at the next structural edit, and nothing lists the notes again: remove one by its id."""
    doc = rdocx.Document()
    doc.add_paragraph("A claim.")
    doc.add_paragraph("Another claim.")
    first = doc.add_footnote(doc.paragraphs[0], "The source.").id
    doc.add_footnote(doc.paragraphs[1].runs[-1], "A second source.")
    doc.remove_footnote(first)
    doc.save(tmp_path / "n.docx")
    text = run([rdocx_cli, "text", tmp_path / "n.docx"], check=True).stdout
    assert "A second source." in text and "The source." not in text


# ---------------------------------------------------------------- raw XML
@pytest.mark.gap("docx-raw-xml-unchecked")
def test_raw_xml_that_word_refuses_is_refused():
    """An unknown Word element, and a foreign element that mc:Ignorable does not cover, make Word report the file as
    damaged: replace_xml refuses them, the document unchanged."""
    doc = rdocx.Document()
    doc.add_paragraph("Keep me")
    before = doc.to_bytes()
    xml = doc.paragraphs[0].xml
    for bad in (xml.replace("</w:p>", "<w:bogus/></w:p>"), xml.replace("</w:p>", '<x:foo xmlns:x="urn:example"/></w:p>')):
        with pytest.raises((ValueError, rdocx.RdocxError)):
            doc.paragraphs[0].replace_xml(bad)
    assert doc.to_bytes() == before


# ---------------------------------------------------------------- charts and equations
@pytest.mark.gap("docx-chart-read-back")
def test_charts_are_read_back(tmp_path):
    doc = rdocx.Document()
    doc.add_paragraph("Costs")
    data = rdocx.CategoryChartData()
    data.categories = ["2024", "2025"]
    data.add_series("Inspections", (12, 15))
    doc.add_chart(rdocx.XL_CHART_TYPE.COLUMN_CLUSTERED, data, rdocx.Inches(5), rdocx.Inches(3))
    doc.save(tmp_path / "c.docx")
    charts = rdocx.Document.open(tmp_path / "c.docx").charts
    assert [(c.chart_type, c.categories, [(s.name, s.values) for s in c.series]) for c in charts] == [
        (rdocx.XL_CHART_TYPE.COLUMN_CLUSTERED, ("2024", "2025"), [("Inspections", (12.0, 15.0))])]


@pytest.mark.gap("docx-display-equation")
def test_a_display_equation(tmp_path):
    doc = rdocx.Document()
    doc.add_paragraph("").add_equation(r"\frac{a}{b}", format="latex", display=True)
    assert doc.paragraphs[0].equations()[0].is_display is True
    assert "<m:oMathPara" in saved(doc, tmp_path / "e.docx")


# ---------------------------------------------------------------- values Word ignores, accepted silently
@pytest.mark.gap("docx-invalid-values-accepted")
def test_saving_under_the_extension_of_another_format_raises(tmp_path):
    doc = rdocx.Document()
    doc.add_paragraph("x")
    with pytest.raises(ValueError, match="to_pdf"):
        doc.save(tmp_path / "x.pdf")                             # a DOCX named .pdf
    assert not (tmp_path / "x.pdf").exists()


@pytest.mark.gap("docx-invalid-values-accepted")
def test_a_second_tab_stop_at_one_position_raises():
    pf = rdocx.Document().add_paragraph("Name\tValue").paragraph_format
    pf.tab_stops.add_tab_stop(rdocx.Inches(1))
    with pytest.raises(ValueError):
        pf.tab_stops.add_tab_stop(rdocx.Inches(1))              # one tab per position
    assert len(pf.tab_stops) == 1


@pytest.mark.gap("docx-invalid-values-accepted")
def test_a_page_break_in_a_header_raises():
    doc = rdocx.Document()
    doc.add_paragraph("Body")
    doc.section(0).header.add_paragraph("Head")
    with pytest.raises((ValueError, rdocx.RdocxError)):
        doc.section(0).header.paragraphs[0].runs[0].add_break("page")   # Word ignores it there


@pytest.mark.gap("docx-invalid-values-accepted")
def test_a_malformed_language_tag_raises():
    doc = rdocx.Document()
    doc.add_paragraph("x")
    with pytest.raises(ValueError):
        doc.paragraphs[0].runs[0].font.language = "not a tag!"


@pytest.mark.gap("docx-invalid-values-accepted")
def test_a_font_face_without_its_name_table_is_refused(free_font):
    doc = rdocx.Document()
    doc.add_paragraph("x")
    licence = rdocx.FontEmbeddingLicense(authorized=True, identity="CC0-1.0: Aileron")
    with pytest.raises(rdocx.RdocxError):
        doc.embed_font("Aileron", free_font.read_bytes().replace(b"name", b"nam_", 1), license=licence)
    assert doc.fonts == () or all(not f.embedded_fonts for f in doc.fonts)
