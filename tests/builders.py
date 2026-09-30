"""Small synthetic inputs for the tests, built with python-docx plus targeted XML where python-docx has no API.
Each builder writes a file and returns its path."""
import zipfile

import docx
from docx.oxml.ns import qn

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
FOOTNOTES_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
FOOTNOTES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
ENDNOTES_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/endnotes"
ENDNOTES_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.endnotes+xml"


def notes_part(kind, text):
    """footnotes.xml or endnotes.xml with the two separators and note 1 holding `text`."""
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:{kind}s xmlns:w="{W}">'
            f'<w:{kind} w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:{kind}>'
            f'<w:{kind} w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:{kind}>'
            f'<w:{kind} w:id="1"><w:p><w:r><w:{kind}Ref/></w:r><w:r><w:t xml:space="preserve"> {text}</w:t></w:r></w:p></w:{kind}>'
            f'</w:{kind}s>')


def wrap_run(paragraph, text, wrapper):
    """Append a run holding `text` inside a wrapper element: "sdt" (inline content control) or "ins" (tracked
    insertion)."""
    p = paragraph._p
    r = paragraph.add_run(text)._r
    p.remove(r)
    if wrapper == "sdt":
        outer = p.makeelement(qn("w:sdt"), {})
        pr = outer.makeelement(qn("w:sdtPr"), {})
        tag = pr.makeelement(qn("w:tag"), {qn("w:val"): "goog_rdk_9"})
        pr.append(tag)
        outer.append(pr)
        content = outer.makeelement(qn("w:sdtContent"), {})
        content.append(r)
        outer.append(content)
    elif wrapper == "ins":
        outer = p.makeelement(qn("w:ins"), {qn("w:id"): "901", qn("w:author"): "Editor",
                                            qn("w:date"): "2026-01-01T00:00:00Z"})
        outer.append(r)
    elif wrapper == "del":
        outer = p.makeelement(qn("w:del"), {qn("w:id"): "902", qn("w:author"): "Editor",
                                            qn("w:date"): "2026-01-01T00:00:00Z"})
        r.remove(r.find(qn("w:t")))
        dt = r.makeelement(qn("w:delText"), {})
        dt.text = text
        dt.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r.append(dt)
        outer.append(r)
    else:
        raise ValueError(wrapper)
    p.append(outer)
    return paragraph


def cell_text_docx(path):
    """"beta" in a body paragraph, in a cell paragraph whose single bold run is "alpha beta gamma" (table 0,
    cell 0,0), in a nested table (cell 1,1), and in a last body paragraph: body, cell, nested cell, body."""
    d = docx.Document()
    d.add_paragraph("beta in the body")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).paragraphs[0].add_run("alpha beta gamma").bold = True
    t.cell(0, 1).text = "delta"
    t.cell(1, 1).add_table(rows=1, cols=1).cell(0, 0).text = "nested beta"
    d.add_paragraph("after beta")
    d.save(path)
    return path


def wrapped_run_docx(path, wrapper, target="TARGET"):
    """One paragraph: "before " | <wrapper>target</wrapper> | " after"."""
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("before ")
    wrap_run(p, target, wrapper)
    p.add_run(" after")
    d.save(path)
    return path


def block_sdt(document, text):
    """A body-level content control holding one paragraph (as Google Docs and Word forms write them)."""
    body = document.element.body
    sdt = body.makeelement(qn("w:sdt"), {})
    pr = sdt.makeelement(qn("w:sdtPr"), {})
    pr.append(pr.makeelement(qn("w:tag"), {qn("w:val"): "goog_rdk_8"}))
    sdt.append(pr)
    content = sdt.makeelement(qn("w:sdtContent"), {})
    p = document.add_paragraph(text)._p
    body.remove(p)
    content.append(p)
    sdt.append(content)
    body.insert(len(body) - 1, sdt)  # before the final sectPr


def every_story_docx(path, token="NEEDLE"):
    """`token` once in each place a replacement may or may not reach. Returns (path, places) where places maps
    each place to its number of occurrences."""
    d = docx.Document()
    d.add_paragraph(f"Body {token} one.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 1).text = f"Cell {token}"
    p = d.add_paragraph("Inline control: ")
    wrap_run(p, f"sdt {token}", "sdt")
    block_sdt(d, f"Block control {token}")
    p = d.add_paragraph("Tracked: ")
    wrap_run(p, f"ins {token}", "ins")
    wrap_run(p, f"del {token}", "del")
    p = d.add_paragraph("Footnote here")
    s = d.sections[0]
    s.different_first_page_header_footer = True
    s.header.paragraphs[0].text = f"Header {token}"
    s.footer.paragraphs[0].text = f"Footer {token}"
    s.first_page_footer.paragraphs[0].text = f"First footer {token}"
    s.footer.add_table(1, 1, docx.shared.Inches(3)).cell(0, 0).text = f"Footer cell {token}"
    d.save(path)
    textbox = ('<w:r><w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml" style="width:100pt;height:50pt">'
               f'<v:textbox><w:txbxContent><w:p><w:r><w:t>Text box {token}</w:t></w:r></w:p></w:txbxContent>'
               '</v:textbox></v:shape></w:pict></w:r>')
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "word/document.xml":
                data = data.replace(b"<w:t>Footnote here</w:t></w:r>",
                                    b'<w:t>Footnote here</w:t></w:r><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
                                    b'<w:footnoteReference w:id="1"/></w:r><w:r><w:endnoteReference w:id="1"/></w:r>'
                                    + textbox.encode(), 1)
            elif info.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", f'<Override PartName="/word/footnotes.xml" ContentType="{FOOTNOTES_CT}"/>'
                                    f'<Override PartName="/word/endnotes.xml" ContentType="{ENDNOTES_CT}"/></Types>'.encode())
            elif info.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", f'<Relationship Id="rId99" Type="{FOOTNOTES_TYPE}" Target="footnotes.xml"/>'
                                    f'<Relationship Id="rId98" Type="{ENDNOTES_TYPE}" Target="endnotes.xml"/></Relationships>'.encode())
            z.writestr(info, data)
        z.writestr("word/footnotes.xml", notes_part("footnote", f"Footnote {token}."))
        z.writestr("word/endnotes.xml", notes_part("endnote", f"Endnote {token}."))
    places = {"body": 1, "table_cell": 1, "inline_sdt": 1, "block_sdt": 1, "tracked_insertion": 1, "header": 1,
              "footer": 2, "footer_table_cell": 1, "footnote": 1, "endnote": 1, "text_box": 1}
    return path, places


WPS_BOX = ('<w:drawing><wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
           'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" behindDoc="0" locked="0" '
           'layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column">'
           '<wp:posOffset>0</wp:posOffset></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0'
           '</wp:posOffset></wp:positionV><wp:extent cx="1270000" cy="635000"/><wp:effectExtent l="0" t="0" r="0" b="0"/>'
           '<wp:wrapNone/><wp:docPr id="10" name="Text Box 10"/><wp:cNvGraphicFramePr/>'
           '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
           '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
           '<wps:wsp xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:cNvSpPr txBox="1"/>'
           '<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1270000" cy="635000"/></a:xfrm><a:prstGeom prst="rect">'
           '<a:avLst/></a:prstGeom></wps:spPr><wps:txbx><w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p>'
           '</w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing>')
VML_BOX = ('<w:pict><v:shape xmlns:v="urn:schemas-microsoft-com:vml" style="width:100pt;height:50pt"><v:textbox>'
           '<w:txbxContent><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict>')
MC_BOX = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
          '<mc:Choice xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" Requires="wps">'
          '{choice}</mc:Choice><mc:Fallback>{fallback}</mc:Fallback></mc:AlternateContent>')


def word_textbox_docx(path, token="NEEDLE", footnote=True):
    """A text box as Word writes it (mc:AlternateContent: DrawingML in the Choice, a VML copy in the
    Fallback) holding `token` once, and, with footnote=True, a footnote holding it once: 2 visible occurrences."""
    d = docx.Document()
    d.add_paragraph("Anchor paragraph")
    d.save(path)
    box = MC_BOX.format(choice=WPS_BOX.format(text=f"Box {token}"), fallback=VML_BOX.format(text=f"Box {token}"))
    with zipfile.ZipFile(path) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "word/document.xml":
                extra = b'<w:r><w:footnoteReference w:id="1"/></w:r>' if footnote else b""
                data = data.replace(b"<w:t>Anchor paragraph</w:t></w:r>",
                                    b"<w:t>Anchor paragraph</w:t></w:r><w:r>" + box.encode() + b"</w:r>" + extra, 1)
            elif footnote and info.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", f'<Override PartName="/word/footnotes.xml" ContentType="{FOOTNOTES_CT}"/></Types>'.encode())
            elif footnote and info.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", f'<Relationship Id="rId99" Type="{FOOTNOTES_TYPE}" Target="footnotes.xml"/></Relationships>'.encode())
            z.writestr(info, data)
        if footnote:
            z.writestr("word/footnotes.xml", notes_part("footnote", f"Footnote {token}."))
    return path


def wrapped_text_docx(path, wrapper, target="MID"):
    """One paragraph: "before " | target inside `wrapper` | " after", for wrappers that hold runs:
    "fldSimple" (a simple field's cached result), "smartTag", "customXml"."""
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("before ")
    r = p.add_run(target)._r
    p._p.remove(r)
    attrs = {"fldSimple": {qn("w:instr"): " DOCPROPERTY Title "}, "smartTag": {qn("w:element"): "place"},
             "customXml": {qn("w:element"): "item"}}[wrapper]
    outer = p._p.makeelement(qn(f"w:{wrapper}"), attrs)
    outer.append(r)
    p._p.append(outer)
    p.add_run(" after")
    d.save(path)
    return path
