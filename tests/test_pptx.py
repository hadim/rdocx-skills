"""pptx: what the pinned build (tensorbee/rdocx main) lacks, as strict xfails of tests/gaps.py. Everything else rpptx
does is upstream's to test; the skills' own contract is checked by test_docs_snippets.py and test_skill_docs.py."""
import io

import pptx
import pytest
import rpptx
from PIL import Image
from rpptx.util import Inches

from conftest import digest, run


def deck(*layouts):
    prs = rpptx.Presentation()
    for k in layouts or (6,):
        prs.slides.add_slide(prs.slide_layouts[k])
    return prs


def text_box(prs, text="x"):
    prs.slides[0].shapes.add_textbox(0, 0, Inches(2), Inches(1)).text_frame.text = text
    return prs.slides[0].shapes[len(prs.slides[0].shapes) - 1]


# ---------------------------------------------------------------- python-pptx names
@pytest.mark.gap("python-pptx-names")
def test_slide_id_as_in_python_pptx():
    prs = deck(6, 6)
    assert prs.slides[0].slide_id == 256 and prs.slides[1].slide_id == 257


@pytest.mark.gap("python-pptx-names")
def test_colour_type_and_theme_colour_read_the_enums_as_in_python_pptx():
    """rpptx.enum.dml has both enums, but color.type and color.theme_color read strings ("rgb", "accent2")."""
    from rpptx.enum.dml import MSO_COLOR_TYPE, MSO_THEME_COLOR_INDEX
    font = text_box(deck()).text_frame.paragraphs[0].runs[0].font
    font.color.rgb = "123456"
    assert font.color.type == MSO_COLOR_TYPE.RGB
    font.color.theme_color = MSO_THEME_COLOR_INDEX.ACCENT_2
    assert (font.color.type, font.color.theme_color) == (MSO_COLOR_TYPE.SCHEME, MSO_THEME_COLOR_INDEX.ACCENT_2)


@pytest.mark.gap("python-pptx-names")
def test_language_id_as_in_python_pptx():
    from rpptx.enum.lang import MSO_LANGUAGE_ID
    font = text_box(deck(), "Bonjour").text_frame.paragraphs[0].runs[0].font
    font.language = "fr-FR"
    assert font.language_id == MSO_LANGUAGE_ID.FRENCH
    font.language_id = MSO_LANGUAGE_ID.GERMAN
    assert font.language == "de-DE"


@pytest.mark.gap("python-pptx-names")
@pytest.mark.parametrize("name", ["_element", "_txBody"])
def test_a_python_pptx_private_attribute_names_the_rpptx_way(name):
    shape = text_box(deck())
    owner, way = {"_element": (shape, r"\belement\b"), "_txBody": (shape.text_frame, r"\belement\b")}[name]
    with pytest.raises(AttributeError, match=way):
        getattr(owner, name)


# ---------------------------------------------------------------- values accepted silently
def save_as_pdf(prs, tmp_path):
    prs.save(tmp_path / "x.pdf")


INVALID = {   # call -> what PowerPoint cannot honour
    "save_to_a_pdf_path": save_as_pdf,                                    # writes a PPTX package named x.pdf
    "slide_width_over_56_inches": lambda prs, _: setattr(prs, "slide_width", Inches(60)),   # out of the schema range
    "slide_height_under_1_inch": lambda prs, _: setattr(prs, "slide_height", Inches(0.5)),
    "flip_a_table": lambda prs, _: setattr(prs.slides[0].shapes[0], "flip_h", True),
    "unknown_table_style": lambda prs, _: setattr(prs.slides[0].shapes[0].table, "style_id",
                                                  "{00000000-0000-0000-0000-000000000000}"),
}


@pytest.mark.gap("pptx-invalid-values-accepted")
@pytest.mark.parametrize("case", sorted(INVALID))
def test_a_value_powerpoint_cannot_honour_is_refused(case, tmp_path):
    prs = deck()
    prs.slides[0].shapes.add_table(1, 1, Inches(1), Inches(1), Inches(2), Inches(1))
    before = prs.to_bytes()
    with pytest.raises((ValueError, rpptx.RpptxError)):
        INVALID[case](prs, tmp_path)
    assert digest(prs.to_bytes()) == digest(before) and not (tmp_path / "x.pdf").exists()


RAW = {   # owner -> (old, new) of a replace_xml that PowerPoint cannot open
    "misplaced_in_p_sp": ("shape", "</p:sp>", "<a:p/></p:sp>"),
    "unknown_in_a_prstGeom": ("shape", "<a:avLst/>", "<a:avLst/><a:bogus/>"),
    "unknown_in_a_p": ("text_frame", "<a:r>", "<a:bogus/><a:r>"),
    "unknown_in_p_sld": ("slide", "<p:clrMapOvr>", "<p:bogus/><p:clrMapOvr>"),
}


@pytest.mark.gap("pptx-raw-xml-unchecked")
@pytest.mark.parametrize("case", sorted(RAW))
def test_raw_xml_with_an_unknown_or_misplaced_element_is_refused(case):
    prs = deck()
    prs.slides[0].shapes.add_shape(rpptx.MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(2), Inches(1)).text = "Box"
    kind, old, new = RAW[case]
    owner = {"shape": prs.slides[0].shapes[0], "text_frame": prs.slides[0].shapes[0].text_frame,
             "slide": prs.slides[0]}[kind].element
    xml = owner.xml
    assert old in xml
    before = prs.to_bytes()
    with pytest.raises((ValueError, rpptx.RpptxError)):
        owner.replace_xml(xml.replace(old, new, 1))
    assert digest(prs.to_bytes()) == digest(before)


# ---------------------------------------------------------------- wrong results
@pytest.mark.gap("pptx-chart-type-family-only")
@pytest.mark.parametrize("kind", ["BAR_STACKED", "COLUMN_STACKED", "LINE", "PIE_EXPLODED"])
def test_chart_type_of_a_chart_rpptx_does_not_author(kind, tmp_path):
    """python-pptx writes these; rpptx reads only the plot family (a stacked bar reads as COLUMN_CLUSTERED, a line
    without markers as LINE, which is python-pptx's LINE_MARKERS). The type, or None when rpptx cannot tell."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    p = pptx.Presentation()
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Cost", (1, 2))
    p.slides.add_slide(p.slide_layouts[6]).shapes.add_chart(getattr(XL_CHART_TYPE, kind), 0, 0, Inches(4), Inches(3), data)
    p.save(tmp_path / "c.pptx")
    chart = next(sh for sh in rpptx.Presentation(tmp_path / "c.pptx").slides[0].shapes if sh.has_chart).chart
    assert chart.chart_type in (None, int(getattr(XL_CHART_TYPE, kind)))


@pytest.mark.gap("pptx-replace-map-zero-match")
def test_a_map_pair_without_expect_that_finds_nothing_is_refused(rpptx_cli, tmp_path):
    """As `rpptx replace -p` does: a pair without "expect" must find at least one, so that a mistyped placeholder
    does not pass unnoticed."""
    deck().save(tmp_path / "in.pptx")
    (tmp_path / "map.json").write_text('[{"placeholder": "nothere", "value": "y"}]')
    res = run([rpptx_cli, "replace", tmp_path / "in.pptx", "--map", tmp_path / "map.json", "-o", tmp_path / "out.pptx"])
    assert res.returncode == 1 and not (tmp_path / "out.pptx").exists()


@pytest.mark.gap("pptx-render-ignores-slide-show-master-shapes")
def test_a_slide_that_hides_the_master_shapes_renders_without_them():
    prs = deck(5, 5)
    plain = prs.render_slide_to_png(1, 40)
    logo = io.BytesIO()
    Image.new("RGB", (20, 20), (0, 0, 0)).save(logo, "PNG")
    logo.seek(0)
    prs.slide_master.shapes.add_picture(logo, Inches(9), Inches(6), width=Inches(1))   # a logo on every slide
    assert digest(prs.render_slide_to_png(1, 40)) != digest(plain)
    prs.slides[1].show_master_shapes = False                              # written as showMasterSp="0"
    assert digest(prs.render_slide_to_png(1, 40)) == digest(plain)
