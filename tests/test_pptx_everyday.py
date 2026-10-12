"""pptx: the everyday API of the 2026-10-10 pin that the references document (text and fonts, bullets, fills,
arrangement, alt text, connectors, placeholders, autofit, slide numbers and footers, themes, masters and layouts,
transitions, tables, metadata, sections, media, notes, exports, raw XML, the errors that replace silent results)
and the CLI extras (footer, fit, meta, notes, slide, replace --map, --font-dir); and from the 2026-10-12 pin, charts,
slide import, Markdown and HTML export, embedded fonts and the nesting limit."""
import datetime
import io
import json
import re
import zipfile

import pptx
import pytest
import rpptx
from PIL import Image
from rpptx.dml.color import RGBColor
from rpptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL, MSO_PATTERN_TYPE, MSO_THEME_COLOR
from rpptx.enum.lang import MSO_LANGUAGE_ID
from rpptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE, MSO_SHAPE_TYPE, PP_PLACEHOLDER
from rpptx.enum.text import MSO_AUTO_SIZE
from rpptx.util import Inches, Pt

from conftest import (FREE_FONT_LICENSE, MONOSPACE_FONTS, digest, font_dir, libreoffice_pdf, part, parts, pdf_uses_font,
                      run)

LONG = "This sentence is long enough to need several lines in a small box. " * 4


def png_bytes(size=(40, 20), color=(200, 60, 60)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, "PNG")
    return buf.getvalue()


def deck(*layouts):
    prs = rpptx.Presentation()
    for k in layouts or (6,):
        prs.slides.add_slide(prs.slide_layouts[k])
    return prs


def slide_xml(prs, k=1):
    with zipfile.ZipFile(io.BytesIO(prs.to_bytes())) as z:
        return z.read(f"ppt/slides/slide{k}.xml").decode()


def valid(rpptx_cli, prs, tmp_path, name="v.pptx"):
    prs.save(tmp_path / name)
    res = run([rpptx_cli, "validate", tmp_path / name])
    assert res.returncode == 0, res.stdout + res.stderr
    return tmp_path / name


# ---------------------------------------------------------------- handles
BRAND = rpptx.Presentation()
EDITS = {  # edit -> what it retires, as python-api.md "Handles" says
    "alt_text_flip": (lambda p: (setattr(p.slides[1].shapes[0], "alt_text", "x"), setattr(p.slides[1].shapes[0], "flip_h", True)), ""),
    "align_distribute": (lambda p: (p.slides[1].shapes.align([p.slides[1].shapes[0], p.slides[1].shapes[1]], "left"),
                                    p.slides[1].shapes.distribute([p.slides[1].shapes[0], p.slides[1].shapes[1]], "vertical")), ""),
    "add_field": (lambda p: p.slides[1].shapes[0].text_frame.paragraphs[0].add_field("slidenum"), ""),
    "fit_text_refresh_autofit": (lambda p: (p.slides[1].shapes[0].text_frame.fit_text(), p.refresh_autofit()), ""),
    "header_footer_adds": (lambda p: (p.set_header_footer(footer="x"), setattr(p.slides[1].header_footer, "footer", "y")), ""),
    "transition_theme_styles": (lambda p: (setattr(p.slides[1].transition, "type", "fade"),
                                           p.slide_master.theme.colors.__setitem__("accent1", "000000"),
                                           setattr(p.slide_master.text_styles.body[0], "bullet", "-")), ""),
    "core_properties_sections": (lambda p: (setattr(p.core_properties, "title", "t"), p.set_sections([("a", [0, 1])])), ""),
    "ungroup": (lambda p: (p.slides[1].shapes.group([p.slides[1].shapes[0], p.slides[1].shapes[1]]),
                           p.slides[1].shapes[0].ungroup()), "shapes paragraph run"),
    "shape_replace_xml": (lambda p: p.slides[1].shapes[0].replace_xml(p.slides[1].shapes[0].xml), "shapes paragraph run"),
    "text_frame_replace_xml": (lambda p: p.slides[1].shapes[0].text_frame.replace_xml(p.slides[1].shapes[0].text_frame.xml),
                               "shapes paragraph run"),
    "gradient_stops": (lambda p: (p.slides[1].shapes[0].fill.gradient(), p.slides[1].shapes[0].fill.gradient_stops.append(0.5)),
                       "shapes paragraph run"),
    "slide_replace_xml": (lambda p: p.slides[1].replace_xml(p.slides[1].xml), "slide shapes paragraph run"),
    "layout_duplicate": (lambda p: p.slide_layouts.duplicate(p.slide_layouts[1]), "layout shapes paragraph run"),
    "layout_remove": (lambda p: p.slide_layouts.remove(p.slide_layouts[3]), "layout shapes paragraph run"),
    "layout_replace_xml": (lambda p: p.slide_layouts[6].replace_xml(p.slide_layouts[6].xml), "layout shapes paragraph run"),
    "apply_theme": (lambda p: p.apply_theme(BRAND), "layout shapes paragraph run"),
}


@pytest.mark.parametrize("op", sorted(EDITS))
def test_handles_after_the_new_edits(op):
    prs = deck(6, 6)
    for k in range(2):
        prs.slides[k].shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Draft a"
        prs.slides[k].shapes.add_textbox(Inches(1), Inches(3), Inches(4), Inches(1)).text_frame.text = "Draft b"
    slide, shape, layout = prs.slides[0], prs.slides[0].shapes[1], prs.slide_layouts[6]
    paragraph = prs.slides[0].shapes[1].text_frame.paragraphs[0]
    run_ = paragraph.runs[0]
    held = {"slide": lambda: slide.notes_text, "shapes": lambda: shape.left, "paragraph": lambda: paragraph.text,
            "run": lambda: run_.text, "layout": lambda: layout.name}
    edit, retired = EDITS[op]
    edit(prs)
    for kind, read in held.items():
        if kind in retired.split():
            with pytest.raises(rpptx.StaleElementError):
                read()
        else:
            read()


# ---------------------------------------------------------------- text, fonts, bullets
def test_font_language_spacing_highlight_small_caps_scripts(tmp_path, rpptx_cli):
    prs = deck()
    prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = "Bonjour"
    font = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font
    font.language = "fr-FR"
    assert (font.language, font.language_id) == ("fr-FR", MSO_LANGUAGE_ID.FRENCH)
    font.language_id = MSO_LANGUAGE_ID.GERMAN
    assert font.language == "de-DE"
    font.spacing, font.baseline, font.small_caps = Pt(2), -0.25, True
    font.highlight_color, font.east_asian_name, font.complex_script_name = "FFFF00", "MS Mincho", "Arial"
    assert (font.spacing, font.small_caps, str(font.highlight_color)) == (Pt(2), True, "FFFF00")
    x = slide_xml(prs)
    assert re.search(r'<a:rPr cap="small" spc="200" baseline="-25000" lang="de-DE"><a:highlight><a:srgbClr val="FFFF00"/>'
                     r'</a:highlight><a:ea typeface="MS Mincho"/><a:cs typeface="Arial"/></a:rPr>', x)
    font.baseline = 0.3
    assert 'baseline="30000"' in slide_xml(prs)
    valid(rpptx_cli, prs, tmp_path)


def test_bullets_and_auto_numbers(tmp_path):
    prs = deck()
    tf = prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2)).text_frame
    tf.text = "First"
    p = tf.add_paragraph()
    p.text = "Second"
    paragraphs = prs.slides[0].shapes[0].text_frame.paragraphs
    paragraphs[0].bullet, paragraphs[0].bullet_color = "\u2022", "#7B1E3A"
    paragraphs[0].bullet_size, paragraphs[0].bullet_font = 1.2, "Arial"
    paragraphs[1].auto_number, paragraphs[1].auto_number_start = "arabicPeriod", 3
    assert (paragraphs[0].bullet, str(paragraphs[0].bullet_color), paragraphs[0].bullet_size) == ("\u2022", "7B1E3A", 1.2)
    assert (paragraphs[1].auto_number, paragraphs[1].auto_number_start) == ("arabicPeriod", 3)
    x = slide_xml(prs)
    assert ('<a:buClr><a:srgbClr val="7B1E3A"/></a:buClr><a:buSzPct val="120000"/><a:buFont typeface="Arial"/>'
            '<a:buChar char="\u2022"/>') in x
    assert '<a:buAutoNum type="arabicPeriod" startAt="3"/>' in x
    paragraphs[0].bullet = False
    assert paragraphs[0].bullet is False and "<a:buNone/>" in slide_xml(prs)
    for scheme in ("alphaLcParenR", "romanUcPeriod"):
        paragraphs[1].auto_number = scheme
        assert f'<a:buAutoNum type="{scheme}" startAt="3"/>' in slide_xml(prs)


def test_add_field_slide_number_and_date(tmp_path, rpptx_cli):
    prs = deck()
    tf = prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame
    tf.text = "Page "
    prs.slides[0].shapes[0].text_frame.paragraphs[0].add_field("slidenum")
    prs.slides[0].shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    prs.slides[0].shapes[1].text_frame.paragraphs[0].add_field("datetime4", "10 October 2026")
    x = slide_xml(prs)
    assert re.search(r'<a:fld id="[^"]+" type="slidenum">', x)
    assert re.search(r'<a:fld id="[^"]+" type="datetime4"><a:t>10 October 2026</a:t></a:fld>', x)
    with pytest.raises(rpptx.RpptxError, match="slidenum or datetime1 to datetime13"):
        prs.slides[0].shapes[1].text_frame.paragraphs[0].add_field("page")
    valid(rpptx_cli, prs, tmp_path)


# ---------------------------------------------------------------- colours and fills
def test_colour_theme_alpha_gradient_pattern_picture_fills(tmp_path, rpptx_cli):
    prs = deck()
    for k in range(4):
        prs.slides[0].shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1 + 2.5 * k), Inches(1), Inches(2), Inches(1))
    solid = prs.slides[0].shapes[0].fill
    solid.solid()
    solid.fore_color.rgb = "336699"
    solid.fore_color.alpha = 0.5
    assert solid.fore_color.alpha == 0.5 and '<a:srgbClr val="336699"><a:alpha val="50000"/>' in slide_xml(prs)
    solid.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_2
    assert (solid.fore_color.type, solid.fore_color.theme_color) == (MSO_COLOR_TYPE.SCHEME, MSO_THEME_COLOR.ACCENT_2)
    grad = prs.slides[0].shapes[1].fill
    grad.gradient()
    assert grad.type == MSO_FILL.GRADIENT and len(grad.gradient_stops) == 2 and grad.gradient_path is None
    assert grad.gradient_stops[0].color.theme_color == MSO_THEME_COLOR.ACCENT_1
    grad.gradient_angle = 90.0
    grad.gradient_stops[0].color.rgb, grad.gradient_stops[1].color.rgb = "FF0000", "0000FF"
    grad.gradient_stops.append(0.5).color.rgb = "00FF00"
    with pytest.raises(rpptx.StaleElementError):             # append retires the shape handles
        grad.type
    stops = prs.slides[0].shapes[1].fill.gradient_stops
    assert [(s.position, str(s.color.rgb)) for s in stops] == [(0.0, "FF0000"), (0.5, "00FF00"), (1.0, "0000FF")]
    del prs.slides[0].shapes[1].fill.gradient_stops[1]
    assert len(prs.slides[0].shapes[1].fill.gradient_stops) == 2
    patt = prs.slides[0].shapes[2].fill
    patt.patterned()
    patt.pattern, patt.fore_color.rgb, patt.back_color.rgb = MSO_PATTERN_TYPE.CROSS, "000000", "FFFFFF"
    assert (patt.type, patt.pattern, str(patt.back_color.rgb)) == (MSO_FILL.PATTERNED, MSO_PATTERN_TYPE.CROSS, "FFFFFF")
    prs.slides[0].shapes[3].fill.picture(io.BytesIO(png_bytes()))
    assert prs.slides[0].shapes[3].fill.type == MSO_FILL.PICTURE
    prs.slides[0].shapes.add_table(1, 1, Inches(1), Inches(3), Inches(2), Inches(1))
    prs.slides[0].shapes[4].table.cell(0, 0).fill.picture(io.BytesIO(png_bytes()))
    prs.slides[0].background.fill.picture(io.BytesIO(png_bytes()))
    path = valid(rpptx_cli, prs, tmp_path)
    shapes = pptx.Presentation(path).slides[0].shapes
    assert shapes[1].fill.type == pptx.enum.dml.MSO_FILL.GRADIENT and shapes[1].fill.gradient_angle == 90.0
    assert shapes[2].fill.pattern == pptx.enum.dml.MSO_PATTERN.CROSS
    assert shapes[0].fill.fore_color.theme_color == pptx.enum.dml.MSO_THEME_COLOR.ACCENT_2
    assert "<a:blipFill" in re.search(r"<p:bg>.*?</p:bg>", parts(path)["ppt/slides/slide1.xml"].decode()).group()


# ---------------------------------------------------------------- shapes
def test_alt_text_decorative_and_flips(tmp_path, rpptx_cli):
    prs = deck()
    prs.slides[0].shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(1), Inches(1), Inches(2), Inches(1))
    prs.slides[0].shapes.add_picture(io.BytesIO(png_bytes()), Inches(4), Inches(1))
    prs.slides[0].shapes.add_table(1, 1, Inches(1), Inches(3), Inches(2), Inches(1))
    arrow, picture = prs.slides[0].shapes[0], prs.slides[0].shapes[1]
    arrow.alt_text, arrow.alt_title, arrow.flip_h = "Next step", "Arrow", True
    picture.decorative = True
    assert (arrow.alt_text, arrow.alt_title, arrow.decorative, arrow.flip_h, arrow.flip_v) == (
        "Next step", "Arrow", False, True, False)
    assert picture.decorative and picture.alt_text is None
    with pytest.raises(rpptx.RpptxError):
        prs.slides[0].shapes[2].flip_h = True
    path = valid(rpptx_cli, prs, tmp_path)
    x = parts(path)["ppt/slides/slide1.xml"].decode()
    assert 'descr="Next step" title="Arrow"' in x and '<a:xfrm flipH="1">' in x
    assert '<adec:decorative xmlns:adec="http://schemas.microsoft.com/office/drawing/2017/decorative" val="1"/>' in x
    assert pptx.Presentation(path).slides[0].shapes[0]._element.xpath("./p:nvSpPr/p:cNvPr/@descr") == ["Next step"]


def test_group_ungroup_align_distribute():
    prs = deck(6, 1)
    for k in range(3):
        prs.slides[0].shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1 + 2 * k + 0.3 * k * k), Inches(1 + 0.5 * k), Inches(1), Inches(1))
    s = prs.slides[0].shapes
    s.align([s[0], s[1], s[2]], "top")
    assert [sh.top for sh in prs.slides[0].shapes] == [Inches(1)] * 3
    s = prs.slides[0].shapes
    s.distribute([s[0], s[1], s[2]], "horizontal")
    left = [sh.left for sh in prs.slides[0].shapes]
    assert left[1] - left[0] == left[2] - left[1] and left[0] == Inches(1)
    s = prs.slides[0].shapes
    s.align([s[0]], "center", relative_to="slide")
    assert prs.slides[0].shapes[0].left == (prs.slide_width - Inches(1)) // 2
    s = prs.slides[0].shapes
    ids, before = [sh.shape_id for sh in s][:2], [(sh.left, sh.top) for sh in s][:2]
    group = s.group([s[0], s[1]])
    assert group.shape_type == MSO_SHAPE_TYPE.GROUP and [m.shape_id for m in group.shapes] == ids
    assert [(m.left, m.top) for m in group.shapes] == before and len(prs.slides[0].shapes) == 2
    k = next(i for i, sh in enumerate(prs.slides[0].shapes) if sh.shape_type == MSO_SHAPE_TYPE.GROUP)
    members = prs.slides[0].shapes[k].ungroup()
    assert [m.shape_id for m in members] == ids and [(m.left, m.top) for m in members] == before
    assert len(prs.slides[0].shapes) == 3
    prs.slides[0].shapes.add_group_shape([prs.slides[0].shapes[0], prs.slides[0].shapes[1]])
    assert len(prs.slides[0].shapes) == 2
    with pytest.raises(rpptx.RpptxError, match="placeholder"):
        prs.slides[1].shapes.group([prs.slides[1].shapes[0], prs.slides[1].shapes[1]])


def test_glued_connector_sites_and_endpoints(tmp_path, rpptx_cli):
    prs = deck()
    prs.slides[0].shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    prs.slides[0].shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5), Inches(3), Inches(2), Inches(1))
    c = prs.slides[0].shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, 0, 0)
    sites = []
    for site in range(4):
        c.begin_connect(prs.slides[0].shapes[0], site)
        sites.append((c.begin_x / 914400, c.begin_y / 914400))
    assert sites == [(2.0, 1.0), (1.0, 1.5), (2.0, 2.0), (3.0, 1.5)]       # top, left, bottom, right
    c.end_connect(prs.slides[0].shapes[1], 1)
    assert (c.end_x, c.end_y) == (Inches(5), Inches(3.5))
    x = slide_xml(prs)
    assert '<a:stCxn id="2" idx="3"/><a:endCxn id="3" idx="1"/>' in x
    c.end_x = Inches(9)                                      # moving an end releases its glue
    assert c.end_x == Inches(9)
    x = slide_xml(prs)
    assert '<a:stCxn id="2" idx="3"/>' in x and "endCxn" not in x
    valid(rpptx_cli, prs, tmp_path)


def test_placeholders_and_insert_picture(tmp_path, rpptx_cli):
    prs = deck(8, 6)                                         # Picture with Caption
    kinds = [(ph.is_placeholder, ph.placeholder_format.idx, ph.placeholder_format.type) for ph in prs.slides[0].placeholders]
    assert kinds == [(True, 0, PP_PLACEHOLDER.TITLE), (True, 1, PP_PLACEHOLDER.PICTURE), (True, 2, PP_PLACEHOLDER.BODY)]
    prs.slides[1].shapes.add_textbox(0, 0, Inches(1), Inches(1))
    assert prs.slides[1].shapes[0].is_placeholder is False
    picture = prs.slides[0].placeholders[1].insert_picture(io.BytesIO(png_bytes((40, 10))))
    assert picture.shape_type == MSO_SHAPE_TYPE.PICTURE and picture.is_placeholder
    assert picture.crop_left == picture.crop_right > 0 and picture.crop_top == picture.crop_bottom == 0
    assert prs.slides[0].placeholders[1].shape_type == MSO_SHAPE_TYPE.PICTURE
    valid(rpptx_cli, prs, tmp_path)


# ---------------------------------------------------------------- autofit
def three_boxes():
    prs = deck()
    for k in range(3):
        box = prs.slides[0].shapes.add_textbox(Inches(1 + 3.5 * k), Inches(1), Inches(3), Inches(1))
        box.text_frame.word_wrap = True
        box.text_frame.text = LONG
        prs.slides[0].shapes[k].text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    return prs


def test_auto_size_is_stored_as_powerpoint_does_whatever_the_order(tmp_path, rpptx_cli):
    prs = three_boxes()
    prs.slides[0].shapes[0].text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    prs.slides[0].shapes[1].text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    assert [(f.autofit, f.font_scale) for f in prs.text_layout()][:2] == [("normal", 0.625), ("shape", 1.0)]  # refreshed
    path = valid(rpptx_cli, prs, tmp_path)
    x = parts(path)["ppt/slides/slide1.xml"].decode()
    assert '<a:normAutofit fontScale="62500" lnSpcReduction="20000"/>' in x
    heights = [sh.height for sh in rpptx.Presentation(path).slides[0].shapes]
    assert heights[0] == heights[2] == Inches(1) and heights[1] > Inches(2)
    # the python-pptx order: auto_size first, then the text, then save
    other = deck()
    other.slides[0].shapes.add_textbox(Inches(4.5), Inches(1), Inches(3), Inches(1))
    other.slides[0].shapes[0].text_frame.word_wrap = True
    other.slides[0].shapes[0].text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    other.slides[0].shapes[0].text_frame.text = LONG
    other.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    assert rpptx.Presentation.from_bytes(other.to_bytes()).slides[0].shapes[0].height == heights[1]


def test_refresh_autofit_and_fit_text():
    prs = three_boxes()
    prs.slides[0].shapes[0].text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    prs.slides[0].shapes[1].text_frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
    results = prs.refresh_autofit()
    assert [(r.shape_id, r.autofit, r.font_scale, r.line_spacing_reduction, r.fits) for r in results] == [
        (2, "normal", 0.625, 0.2, True), (3, "shape", 1.0, 0.0, True)]
    assert results[1].height == prs.slides[0].shapes[1].height and results[0].height is None
    assert all(isinstance(font, str) and isinstance(stand_in, str) for font, stand_in in results[0].font_substitutions)
    one = prs.slides[0].shapes[0].text_frame.refresh_autofit()
    assert (one.shape_id, one.font_scale) == (2, 0.625)
    box = prs.slides[0].shapes.add_textbox(Inches(1), Inches(4), Inches(3), Inches(1))
    box.text_frame.text = LONG
    fit = prs.slides[0].shapes[3].text_frame.fit_text(max_size=18)
    assert (fit.autofit, fit.font_size, fit.fits) == ("none", Pt(9), True)
    tf = prs.slides[0].shapes[3].text_frame
    assert tf.auto_size == MSO_AUTO_SIZE.NONE and tf.word_wrap is True and tf.paragraphs[0].runs[0].font.size == Pt(9)
    assert prs.slides[0].shapes.add_textbox(0, 0, Inches(1), Inches(1)).text_frame.fit_text() is None   # empty frame
    tiny = prs.slides[0].shapes.add_textbox(0, 0, Inches(1), Inches(0.3))
    tiny.text_frame.text = "word " * 3000
    with pytest.raises(rpptx.RpptxError, match="even at 1 point"):
        prs.slides[0].shapes[5].text_frame.fit_text(max_size=10)


def test_cli_fit_reports_overflow_with_exit_1(rpptx_cli, deck_pptx, tmp_path):
    res = run([rpptx_cli, "fit", "--json", deck_pptx])
    assert res.returncode == 1
    report = json.loads(res.stdout)
    assert report["fits"] is False and [(f["slide"], f["shape_id"]) for f in report["overflowing"]] == [(5, 19)]
    assert report["overflowing"][0]["needed_font_scale"] == 0.825
    import sys
    sys.path.insert(0, str(__import__("pathlib").Path(__file__).parent.parent / "skills" / "pptx" / "scripts"))
    import pptx_ops
    assert [(s + 1, sid) for s, sid, _ in pptx_ops.overflowing(deck_pptx)] == [(5, 19)]   # the same frames
    prs = deck()
    prs.slides[0].shapes.add_table(1, 1, Inches(1), Inches(1), Inches(1), Inches(0.3))
    prs.slides[0].shapes[0].table.cell(0, 0).text = LONG * 3                 # tables are not checked
    prs.save(tmp_path / "t.pptx")
    res = run([rpptx_cli, "fit", tmp_path / "t.pptx"])
    assert res.returncode == 0 and "0 of" in res.stdout
    assert run([rpptx_cli, "fit", tmp_path / "missing.pptx"]).returncode == 2


# ---------------------------------------------------------------- slide numbers, footers, dates
def test_set_header_footer_and_slide_header_footer(tmp_path, rpptx_cli):
    prs = deck(0, 1, 5)
    prs.set_header_footer(footer="Board 2026", date="auto")
    rows = [(s.header_footer.slide_number, s.header_footer.footer, s.header_footer.date, s.header_footer.date_format)
            for s in prs.slides]
    assert rows == [(False, None, None, None)] + [(True, "Board 2026", "auto", "datetime1")] * 2   # hide_on_title
    types = [ph.placeholder_format.type for ph in prs.slides[1].placeholders]
    assert types[-3:] == [PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER]
    x = slide_xml(prs, 2)
    assert re.search(r'<a:fld id="[^"]+" type="slidenum">', x) and re.search(r'<a:fld id="[^"]+" type="datetime1">', x)
    prs.slides[2].header_footer.footer = "Only here"
    prs.slides[2].header_footer.date = "1 January 2026"
    prs.slides[2].header_footer.slide_number = False
    assert (prs.slides[2].header_footer.footer, prs.slides[2].header_footer.date) == ("Only here", "1 January 2026")
    assert PP_PLACEHOLDER.SLIDE_NUMBER not in [ph.placeholder_format.type for ph in prs.slides[2].placeholders]
    path = valid(rpptx_cli, prs, tmp_path)
    texts = [sh.text_frame.text for sh in pptx.Presentation(path).slides[2].shapes if sh.has_text_frame]
    assert "Only here" in texts and "1 January 2026" in texts


def test_a_header_footer_call_that_removes_a_placeholder_retires_shapes():
    prs = deck(1)
    prs.set_header_footer(footer="F", date="fixed")
    held = prs.slides[0].shapes[0]
    prs.slides[0].header_footer.footer = None
    with pytest.raises(rpptx.StaleElementError):
        held.name
    held = prs.slides[0].shapes[0]
    prs.set_header_footer(footer=None, date=None)
    with pytest.raises(rpptx.StaleElementError):
        held.name


def test_cli_footer(rpptx_cli, tmp_path):
    prs = deck(0, 1, 1)
    prs.save(tmp_path / "in.pptx")
    res = run([rpptx_cli, "footer", tmp_path / "in.pptx", "--slide-number", "--footer", "Board", "--date", "Q3 2026",
               "--skip-title", "-o", tmp_path / "out.pptx", "--json"], check=True)
    record = json.loads(res.stdout)["slides"]
    assert record[0] == {"date": None, "footer": None, "slide_number": False}
    assert record[1] == {"date": "Q3 2026", "footer": "Board", "slide_number": True}
    out = rpptx.Presentation(tmp_path / "out.pptx")
    assert (out.slides[2].header_footer.footer, out.slides[2].header_footer.slide_number) == ("Board", True)
    run([rpptx_cli, "footer", tmp_path / "in.pptx", "--date", "auto", "--date-format", "datetime4", "-o",
         tmp_path / "auto.pptx"], check=True)
    auto = rpptx.Presentation(tmp_path / "auto.pptx").slides[1].header_footer
    assert (auto.date, auto.date_format, auto.footer, auto.slide_number) == ("auto", "datetime4", None, False)   # opt-in


# ---------------------------------------------------------------- masters, layouts, themes, transitions
def test_theme_colours_and_fonts(tmp_path, rpptx_cli):
    prs = deck(1)
    theme = prs.slide_master.theme
    assert prs.slide_masters[0] == prs.slide_master and len(prs.slide_masters) == 1
    assert theme.colors.keys() == ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5",
                                   "accent6", "hlink", "folHlink"]
    assert dict(theme.colors.items())["accent1"] == theme.colors["accent1"] and isinstance(theme.colors["accent1"], RGBColor)
    theme.colors["accent1"] = "#7B1E3A"
    theme.fonts.major.latin, theme.fonts.minor.east_asian, theme.fonts.minor.complex_script = "Georgia", "MS Mincho", "Arial"
    t = prs.slide_master.theme
    assert (str(t.colors["accent1"]), t.fonts.major.latin, t.fonts.minor.east_asian, t.fonts.minor.complex_script) == (
        "7B1E3A", "Georgia", "MS Mincho", "Arial")
    with pytest.raises(KeyError, match="accent1"):
        t.colors["brand"] = "000000"
    path = valid(rpptx_cli, prs, tmp_path)
    theme_xml = parts(path)["ppt/theme/theme1.xml"].decode()
    assert '<a:accent1><a:srgbClr val="7B1E3A"/></a:accent1>' in theme_xml and '<a:latin typeface="Georgia"' in theme_xml


def test_master_text_styles_logo_and_show_master_shapes(tmp_path, rpptx_cli):
    prs = deck(5, 5)
    styles = prs.slide_master.text_styles
    assert (len(styles.title), len(styles.body), len(styles.other)) == (9, 9, 9)
    level = styles.body[0]
    level.bullet, level.bullet_color = "\u2013", "7B1E3A"
    level.left_indent, level.first_line_indent = Inches(0.4), -Inches(0.4)
    styles.title[0].font.name, styles.title[0].font.size = "Georgia", Pt(40)
    prs.slide_master.text_styles.body[1].bullet = False
    assert prs.slide_master.text_styles.body[1].bullet is False
    prs.slide_master.background.fill.gradient()
    assert prs.slide_master.background.fill.type == MSO_FILL.GRADIENT and len(prs.slide_master.placeholders) == 5
    level = prs.slide_master.text_styles.body[0]
    assert (level.bullet, str(level.bullet_color), level.left_indent, level.first_line_indent) == (
        "\u2013", "7B1E3A", Inches(0.4), -Inches(0.4))
    assert prs.slide_master.text_styles.title[0].font.size == Pt(40)
    plain = prs.render_slide_to_png(1, 40)
    prs.slide_master.shapes.add_picture(io.BytesIO(png_bytes((20, 20), (0, 0, 0))), Inches(9), Inches(6), width=Inches(1))
    assert digest(prs.render_slide_to_png(1, 40)) != digest(plain)         # the logo shows on every slide
    prs.slides[1].show_master_shapes = False
    assert prs.slides[1].show_master_shapes is False and digest(prs.render_slide_to_png(1, 40)) == digest(plain)
    assert digest(prs.render_slide_to_png(0, 40)) != digest(plain)
    path = valid(rpptx_cli, prs, tmp_path)
    master = parts(path)["ppt/slideMasters/slideMaster1.xml"].decode()
    assert '<a:buChar char="\u2013"/>' in master and "<p:pic>" in master


def test_layouts_lookup_duplicate_remove_and_edit(tmp_path, rpptx_cli):
    prs = deck(0, 1)
    layouts = prs.slide_layouts
    title = layouts.get_by_name("Title Slide")
    assert layouts.get_by_name("Nope") is None and layouts.get_by_name("Nope", layouts[6]).name == "Blank"
    assert title.used_by_slides == (prs.slides[0],) and title.slide_master == prs.slide_master
    assert len(prs.slide_master.slide_layouts) == 11
    with pytest.raises(ValueError, match="in use"):
        layouts.remove(title)
    copy = layouts.duplicate(layouts[1])
    assert copy.name == "1_Title and Content" and len(prs.slide_layouts) == 12
    copy.name = "Agenda"
    assert prs.slide_layouts.get_by_name("Agenda") is not None
    prs.slide_layouts.remove(prs.slide_layouts.get_by_name("Agenda"))
    assert len(prs.slide_layouts) == 11
    blank = prs.slide_layouts[6]
    blank.show_master_shapes = False
    blank.background.fill.solid()
    blank.background.fill.fore_color.rgb = "F2F2F2"
    blank.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(1), Inches(1))
    blank = prs.slide_layouts[6]
    assert (blank.show_master_shapes, blank.follow_master_background, len(blank.shapes)) == (False, False, 4)
    assert len(blank.placeholders) == 3
    prs.slides[1].slide_layout = prs.slide_layouts[5]
    assert prs.slides[1].slide_layout.name == "Title Only"
    valid(rpptx_cli, prs, tmp_path)


def test_apply_theme(tmp_path):
    brand = rpptx.Presentation()
    brand.slide_master.theme.colors["accent1"] = "112233"
    brand.slide_master.theme.fonts.major.latin = "Georgia"
    brand.save(tmp_path / "brand.pptx")
    prs = deck(1, 5)
    prs.slides[0].shapes.title.text = "Kept"
    prs.apply_theme(tmp_path / "brand.pptx")
    assert (str(prs.slide_master.theme.colors["accent1"]), prs.slide_master.theme.fonts.major.latin) == ("112233", "Georgia")
    prs.apply_theme(brand, import_master=True)
    assert prs.slides[0].shapes.title.text == "Kept" and prs.slides[0].slide_layout.name == "Title and Content"
    held = prs.slide_layouts[0]
    prs.apply_theme((tmp_path / "brand.pptx").read_bytes())
    with pytest.raises(rpptx.StaleElementError):
        held.name


def test_transitions(tmp_path, rpptx_cli):
    prs = deck(6, 6)
    tr = prs.slides[1].transition
    assert (tr.type, tr.direction, tr.duration, tr.advance_on_click, tr.advance_after) == (None, None, None, True, None)
    tr.type, tr.direction, tr.duration, tr.advance_after = "push", "left", 0.75, 5.0
    assert (tr.type, tr.direction, tr.duration, tr.advance_after) == ("push", "left", 0.75, 5.0)
    with pytest.raises(ValueError, match="push does not take direction in"):
        tr.direction = "in"
    prs.slides[1].transition.apply_to_all()
    assert (prs.slides[0].transition.type, prs.slides[0].transition.duration) == ("push", 0.75)
    path = valid(rpptx_cli, prs, tmp_path)
    assert '<p:transition spd="med" p14:dur="750" advTm="5000"><p:push dir="l"/>' in parts(path)["ppt/slides/slide1.xml"].decode()


# ---------------------------------------------------------------- tables, metadata, sections, notes, media
def test_table_style_options_and_style_id(tmp_path, rpptx_cli):
    prs = deck()
    prs.slides[0].shapes.add_table(3, 3, Inches(1), Inches(1), Inches(6), Inches(2))
    t = prs.slides[0].shapes[0].table
    assert (t.first_row, t.last_row, t.first_col, t.last_col, t.horz_banding, t.vert_banding, t.style_id) == (
        True, False, False, False, True, False, None)
    t.first_row, t.vert_banding = False, True
    t.style_id = "{5c22544a-7ee6-4342-b048-85bdc9fd1c3a}"
    assert prs.slides[0].shapes[0].table.style_id == "{5C22544A-7EE6-4342-B048-85BDC9FD1C3A}"
    with pytest.raises(ValueError, match="built-in table styles"):
        t.style_id = "{00000000-0000-0000-0000-000000000000}"
    prs.slides[0].shapes[0].table.cell(0, 1).text = "b"
    assert [c.text for c in prs.slides[0].shapes[0].table.rows[0].cells] == ["", "b", ""]
    path = valid(rpptx_cli, prs, tmp_path)
    table = next(sh.table for sh in pptx.Presentation(path).slides[0].shapes if sh.has_table)
    assert (table.first_row, table.vert_banding, table.horz_banding) == (False, True, True)


def test_core_properties_are_writable(tmp_path, rpptx_cli):
    prs = deck()
    cp = prs.core_properties
    assert (cp.title, cp.author, cp.created, cp.revision) == ("", "", None, 0)        # nothing stamped
    prs.save(tmp_path / "fresh.pptx")
    assert b"<dcterms:created" not in parts(tmp_path / "fresh.pptx")["docProps/core.xml"]
    cp.title, cp.author, cp.subject, cp.keywords, cp.comments, cp.category = "Deck", "Claude", "Q3", "a; b", "desc", "Report"
    cp.last_modified_by, cp.content_status, cp.identifier, cp.language, cp.version = "C2", "Draft", "ID-1", "en-GB", "1.0"
    cp.revision = 3
    cp.modified = datetime.datetime(2026, 10, 10, 12, 0, tzinfo=datetime.timezone.utc)
    cp.created = datetime.datetime(2026, 10, 1, 9, 0, tzinfo=datetime.timezone.utc)
    cp.last_printed = datetime.datetime(2026, 1, 1)                                    # naive: read as UTC
    assert prs.core_properties.last_printed == datetime.datetime(2026, 1, 1, tzinfo=datetime.timezone.utc)
    path = valid(rpptx_cli, prs, tmp_path)
    c = pptx.Presentation(path).core_properties
    assert (c.title, c.author, c.subject, c.keywords, c.comments, c.category, c.last_modified_by, c.revision) == (
        "Deck", "Claude", "Q3", "a; b", "desc", "Report", "C2", 3)
    assert (c.content_status, c.identifier, c.language, c.version) == ("Draft", "ID-1", "en-GB", "1.0")
    assert c.modified == datetime.datetime(2026, 10, 10, 12, 0) and c.created == datetime.datetime(2026, 10, 1, 9, 0)
    prs.core_properties.title = None
    assert prs.core_properties.title == ""


def test_sections(tmp_path, rpptx_cli):
    prs = deck(6, 6, 6)
    assert prs.sections == ()
    prs.set_sections([("Intro", [0]), ("Body", [1, 2])])
    assert [(s.name, s.slide_indices) for s in prs.sections] == [("Intro", [0]), ("Body", [1, 2])]
    assert all(re.fullmatch(r"\{[0-9A-F-]{36}\}", s.id) for s in prs.sections)
    with pytest.raises(ValueError, match="every slide once, in slide order"):
        prs.set_sections([("A", [1]), ("B", [0, 2])])
    path = valid(rpptx_cli, prs, tmp_path)
    assert '<p14:section xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" name="Body"' in parts(path)[
        "ppt/presentation.xml"].decode()


def test_notes_slide_as_in_python_pptx():
    prs = deck(6, 6)
    assert (prs.slides[0].has_notes_slide, prs.slides[0].notes_text) == (False, None)
    notes = prs.slides[0].notes_slide
    assert prs.slides[0].has_notes_slide and notes.notes_text_frame.text == ""
    notes.notes_text_frame.text = "Say hello"
    assert prs.slides[0].notes_text == "Say hello" and prs.slides[0].slide_id != prs.slides[1].slide_id


def test_movie_media(tmp_path, rpptx_cli):
    mp4 = b"\x00\x00\x00\x18ftypmp42\x00\x00\x00\x00mp42isom" + b"\x00" * 64
    prs = deck()
    with pytest.raises(ValueError, match="mime_type"):
        prs.slides[0].shapes.add_movie(io.BytesIO(mp4), Inches(1), Inches(1), Inches(4), Inches(3))
    movie = prs.slides[0].shapes.add_movie(io.BytesIO(mp4), Inches(1), Inches(1), Inches(4), Inches(3),
                                           poster_frame_image=io.BytesIO(png_bytes()), mime_type="video/mp4")
    info = prs.slides[0].media
    assert [(m.shape_id, m.kind, m.content_type, m.linked, m.target) for m in info] == [
        (movie.shape_id, "video", "video/mp4", False, "/ppt/media/media1.mp4")]
    assert prs.slides[0].extract_media(movie.shape_id) == mp4
    prs.slides[0].replace_media(movie.shape_id, io.BytesIO(mp4 + b"x"), "video/mp4")
    assert prs.slides[0].extract_media(movie.shape_id) == mp4 + b"x"
    path = valid(rpptx_cli, prs, tmp_path)
    assert pptx.Presentation(path).slides[0].shapes[0].shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.MEDIA
    prs.slides[0].remove_media(movie.shape_id)
    assert prs.slides[0].media == () and len(prs.slides[0].shapes) == 0
    assert not any("media" in n for n in parts(valid(rpptx_cli, prs, tmp_path, "r.pptx")))


# ---------------------------------------------------------------- exports
def with_handout_master(prs):
    """The deck's bytes with a minimal handout master, as PowerPoint adds one (rpptx's default template has none)."""
    rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    ns = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="' + rel[:-1] +
          '" xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')
    master = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:handoutMaster {ns}><p:cSld><p:spTree>'
              '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/></p:spTree></p:cSld>'
              '<p:clrMap bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" accent3="accent3" '
              'accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/></p:handoutMaster>')
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(prs.to_bytes())) as zin, zipfile.ZipFile(out, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/ppt/handoutMasters/handoutMaster1.xml" ContentType='
                                    b'"application/vnd.openxmlformats-officedocument.presentationml.handoutMaster+xml"/></Types>')
            elif info.filename == "ppt/_rels/presentation.xml.rels":
                data = data.replace(b"</Relationships>", f'<Relationship Id="rIdHm1" Type="{rel}handoutMaster" '
                                    'Target="handoutMasters/handoutMaster1.xml"/></Relationships>'.encode())
            elif info.filename == "ppt/presentation.xml":
                data = re.sub(rb"(</p:sldMasterIdLst>)",
                              rb'\1<p:handoutMasterIdLst><p:handoutMasterId r:id="rIdHm1"/></p:handoutMasterIdLst>', data)
            zout.writestr(info, data)
        zout.writestr("ppt/handoutMasters/handoutMaster1.xml", master)
        zout.writestr("ppt/handoutMasters/_rels/handoutMaster1.xml.rels",
                      '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.'
                      f'openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{rel}theme" '
                      'Target="../theme/theme1.xml"/></Relationships>')
    return out.getvalue()


def test_render_slides_pdfa_handouts(tmp_path, rpptx_cli):
    prs = deck(5, 5, 5)
    for k in range(3):
        prs.slides[k].shapes.title.text = f"Slide {k + 1}"
    pngs = prs.render_slides(dpi=40)
    assert len(pngs) == 3 and all(p[:4] == b"\x89PNG" for p in pngs)
    jpeg = prs.render_slides(dpi=40, format="jpeg", quality=80, slides=[1])
    assert len(jpeg) == 1 and jpeg[0][:3] == b"\xff\xd8\xff"
    tiff = prs.render_slides(dpi=40, format="tiff")
    assert isinstance(tiff, bytes) and Image.open(io.BytesIO(tiff)).n_frames == 3
    with pytest.raises(ValueError, match="quality applies to JPEG only"):
        prs.render_slides(format="png", quality=80)
    with pytest.raises(ValueError, match="transparent applies to PNG only"):
        prs.render_slides(format="jpeg", transparent=True)
    assert b"<pdfaid:part>2</pdfaid:part>" in prs.to_pdfa() and b"<pdfaid:part>3</pdfaid:part>" in prs.to_pdfa("pdfa-3b")
    with pytest.raises(rpptx.RpptxError, match="no handout master"):
        prs.to_handout_pdf()
    with_master = rpptx.Presentation.from_bytes(with_handout_master(prs))
    pdf = with_master.to_handout_pdf(2)
    assert pdf[:5] == b"%PDF-" and len(re.findall(rb"/Type\s*/Page[^s]", pdf)) == 2
    assert len(with_master.render_all_handouts(2, dpi=40)) == 2 and len(with_master.render_all_handouts(4, dpi=40)) == 1


def test_odp_round_trip_reports_what_it_drops(tmp_path):
    prs = deck(5, 6)
    prs.slides[0].shapes.title.text = "Inherited geometry"
    prs.slides[1].shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "Kept"
    odp, diagnostics = prs.to_odp()
    assert odp[:2] == b"PK" and diagnostics == [("slides/0/shapes/0", "shape without bounds was dropped")]
    assert prs.save_odp(tmp_path / "d.odp") == diagnostics
    back, read = rpptx.Presentation.from_odp(tmp_path / "d.odp")
    assert len(back.slides) == 2 and back.slides[1].shapes[0].text == "Kept" and read == []
    assert len(rpptx.Presentation.from_odp(odp)[0].slides) == 2


def test_streams(tmp_path):
    prs = deck()
    buf = io.BytesIO()
    prs.save(buf)
    assert len(rpptx.Presentation(io.BytesIO(buf.getvalue())).slides) == 1


# ---------------------------------------------------------------- raw XML
def test_raw_xml_replace_and_refusals():
    prs = deck(5)
    prs.slides[0].shapes.title.text = "Hello"
    prs.slides[0].shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(2), Inches(1)).text = "Box"
    xml = prs.slides[0].shapes[1].xml
    assert isinstance(xml, bytes) and xml.startswith(b"<p:sp ")
    held = prs.slides[0].shapes[0]
    prs.slides[0].shapes[1].replace_xml(xml.replace(b'prst="rect"', b'prst="ellipse"'))
    assert prs.slides[0].shapes[1].auto_shape_type == MSO_SHAPE.OVAL
    with pytest.raises(rpptx.StaleElementError):
        held.name
    tx = prs.slides[0].shapes[1].text_frame.xml
    assert tx.startswith(b"<p:txBody ")
    prs.slides[0].shapes[1].text_frame.replace_xml(tx.replace(b"Box", b"Boxed").decode())   # str works too
    assert prs.slides[0].shapes[1].text == "Boxed"
    prs.slides[0].replace_xml(prs.slides[0].xml.replace(b"Hello", b"Hi"))
    assert prs.slides[0].shapes.title.text == "Hi"
    layout = prs.slide_layouts[5]
    layout.replace_xml(layout.xml.replace(b'name="Title Only"', b'name="Heading Only"'))
    assert prs.slide_layouts[5].name == "Heading Only"
    sp = prs.slides[0].shapes[1].xml
    refused = {
        "malformed": (lambda: prs.slides[0].shapes[1].replace_xml(sp[:-10]), "malformed"),
        "doctype": (lambda: prs.slides[0].shapes[1].replace_xml(b"<!DOCTYPE x>" + sp), "DOCTYPE"),
        "other root": (lambda: prs.slides[0].shapes[1].text_frame.replace_xml(
            b'<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'), "expected one txBody"),
        "missing r:id": (lambda: prs.slides[0].shapes[1].replace_xml(sp.replace(
            b"</p:spPr>", b'<a:blipFill><a:blip r:embed="rId99"/></a:blipFill></p:spPr>', 1)), "rId99"),
        "misplaced in p:sp": (lambda: prs.slides[0].shapes[1].replace_xml(sp.replace(b"</p:sp>", b"<a:p/></p:sp>")),
                              "cannot sit directly in p:sp"),
        "unknown in a:p": (lambda: prs.slides[0].shapes[1].text_frame.replace_xml(
            prs.slides[0].shapes[1].text_frame.xml.replace(b"<a:r>", b"<a:bogus/><a:r>", 1)), "cannot sit directly in a:p"),
        "unknown in p:sld": (lambda: prs.slides[0].replace_xml(prs.slides[0].xml.replace(
            b"<p:clrMapOvr>", b"<p:bogus/><p:clrMapOvr>")), "cannot sit directly in p:sld"),
        "unknown deeper": (lambda: prs.slides[0].shapes[1].replace_xml(sp.replace(b"<a:avLst/>", b"<a:avLst/><a:bogus/>", 1)),
                           "cannot sit directly in a:prstGeom"),
    }
    for name, (call, message) in refused.items():
        before = prs.to_bytes()
        with pytest.raises(ValueError, match=message):
            call()
        assert digest(prs.to_bytes()) == digest(before), name
    with pytest.raises(AttributeError, match=r"shape\.replace_xml"):
        prs.slides[0].shapes[1]._element


# ---------------------------------------------------------------- errors instead of silent results
def test_errors_instead_of_silent_results(tmp_path):
    prs = deck()
    prs.slides[0].shapes.add_textbox(0, 0, Inches(2), Inches(1)).text_frame.text = "x"
    p = prs.slides[0].shapes[0].text_frame.paragraphs[0]
    with pytest.raises(ValueError, match="1 to 4000 points"):
        p.runs[0].font.size = 12
    for attr in ("space_before", "space_after"):
        with pytest.raises(ValueError, match=r"Pt\(6\)"):
            setattr(p, attr, 6)
    with pytest.raises(ValueError, match="must not be empty"):
        p.runs[0].font.name = ""
    with pytest.raises(ValueError, match=r"to_pdf\(\)"):
        prs.save(tmp_path / "x.pdf")
    assert not (tmp_path / "x.pdf").exists()
    with pytest.raises(ValueError, match="56 inches"):
        prs.slide_width = Inches(60)
    with pytest.raises(AttributeError, match="immutable"):
        RGBColor(1, 2, 3).rgb = 5
    with pytest.raises(AttributeError, match=r"prs\.slides\.remove"):
        prs.slides._sldIdLst
    with pytest.raises(AttributeError, match=r"text_frame\.replace_xml"):
        prs.slides[0].shapes[0].text_frame._txBody


# ---------------------------------------------------------------- CLI extras
def cli_deck(tmp_path):
    prs = deck(0, 1, 1, 5)
    for k in range(4):
        prs.slides[k].shapes.title.text = f"Slide {k + 1} {{name}}"
    prs.save(tmp_path / "in.pptx")
    return tmp_path / "in.pptx"


def titles(path):
    return [s.shapes.title.text if s.shapes.title else "" for s in pptx.Presentation(path).slides]


def test_cli_slide_commands(rpptx_cli, tmp_path):
    src, o = cli_deck(tmp_path), lambda n: tmp_path / f"{n}.pptx"
    rec = json.loads(run([rpptx_cli, "slide", "add", src, "--layout", "Title Only", "--at", "2", "-o", o(1), "--json"],
                         check=True).stdout)
    assert (rec["slide"], rec["slides"], rec["layout_name"]) == (2, 5, "Title Only")
    run([rpptx_cli, "slide", "add", src, "--layout", "7", "-o", o("blank")], check=True)
    assert pptx.Presentation(o("blank")).slides[4].slide_layout.name == "Blank"
    res = run([rpptx_cli, "slide", "add", src, "--layout", "Nope", "-o", o("x")])
    assert res.returncode == 1 and '6 "Title Only"' in res.stderr and not o("x").exists()
    run([rpptx_cli, "slide", "duplicate", o(1), "3", "-o", o(2)], check=True)
    run([rpptx_cli, "slide", "move", o(2), "1", "--to", "4", "-o", o(3)], check=True)
    run([rpptx_cli, "slide", "hide", o(3), "2", "-o", o(4)], check=True)
    again = json.loads(run([rpptx_cli, "slide", "hide", o(4), "2", "-o", o(5), "--json"], check=True).stdout)
    assert again["changed"] is False and o(5).exists()
    run([rpptx_cli, "slide", "remove", o(4), "5", "-o", o(6)], check=True)
    assert titles(o(6)) == ["", "Slide 2 {name}", "Slide 2 {name}", "Slide 1 {name}", "Slide 4 {name}"]
    assert pptx.Presentation(o(6)).slides[1]._element.get("show") == "0"
    run([rpptx_cli, "slide", "show", o(6), "2", "-o", o(7)], check=True)
    assert rpptx.Presentation(o(7)).slides[1].hidden is False
    for args in (["slide", "hide", o(7), "1"], ["notes", "set", o(7), "1", "--text", "t"], ["meta", "set", o(7), "--title", "t"],
                 ["footer", o(7), "--slide-number"]):
        res = run([rpptx_cli, *args, "-o", o(6)])                       # an existing output is refused
        assert res.returncode == 1 and "already exists" in res.stderr


def test_cli_notes_set_meta_and_replace_map(rpptx_cli, tmp_path):
    src = cli_deck(tmp_path)
    (tmp_path / "notes.txt").write_text("Line one\nLine two\n")
    rec = json.loads(run([rpptx_cli, "notes", "set", src, "1", "--text", "First note", "-o", tmp_path / "n1.pptx", "--json"],
                         check=True).stdout)
    assert (rec["action"], rec["slide"]) == ("set", 1)
    run([rpptx_cli, "notes", "set", tmp_path / "n1.pptx", "2", "--from-file", tmp_path / "notes.txt", "-o", tmp_path / "n2.pptx"],
        check=True)
    notes = [s.notes_slide.notes_text_frame.text for s in list(pptx.Presentation(tmp_path / "n2.pptx").slides)[:2]]
    assert notes == ["First note", "Line one\nLine two"]
    (tmp_path / "map.json").write_text(json.dumps([{"placeholder": "{name}", "value": "Ada", "expect": 4},
                                                   {"placeholder": "Slide", "value": "Page"},
                                                   {"placeholder": "zzz", "value": "y", "expect": 0}]))
    res = json.loads(run([rpptx_cli, "replace", src, "--map", tmp_path / "map.json", "-o", tmp_path / "r.pptx", "--json"],
                         check=True).stdout)
    assert [p["count"] for p in res["pairs"]] == [4, 4, 0] and res["total"] == 8
    assert titles(tmp_path / "r.pptx")[0] == "Page 1 Ada"
    for pairs, message in (([{"placeholder": "{name}", "value": "Ada", "expect": 2}], "expected 2"),
                           ([{"placeholder": "nothere", "value": "y"}], "no replacements found")):
        (tmp_path / "bad.json").write_text(json.dumps(pairs))
        res = run([rpptx_cli, "replace", src, "--map", tmp_path / "bad.json", "-o", tmp_path / "bad.pptx"])
        assert res.returncode == 1 and message in res.stderr and not (tmp_path / "bad.pptx").exists()
    rec = json.loads(run([rpptx_cli, "meta", "set", src, "--title", "Deck", "--author", "Claude", "--subject", "S",
                          "--keywords", "k1, k2", "--description", "D", "--category", "C", "-o", tmp_path / "m.pptx",
                          "--json"], check=True).stdout)
    assert rec["core"]["title"] == "Deck" and rec["core"]["author"] == "Claude"
    cp = pptx.Presentation(tmp_path / "m.pptx").core_properties
    assert (cp.title, cp.author, cp.subject, cp.keywords, cp.comments, cp.category) == ("Deck", "Claude", "S", "k1, k2", "D", "C")
    got = json.loads(run([rpptx_cli, "meta", "get", "--json", tmp_path / "m.pptx"], check=True).stdout)["core"]
    assert (got["description"], got["category"], got["created"]) == ("D", "C", None)
    assert "title: Deck" in run([rpptx_cli, "meta", "get", tmp_path / "m.pptx"], check=True).stdout


def test_cli_font_dir_on_thumbnail(rpptx_cli, tmp_path):
    fonts = font_dir(tmp_path, MONOSPACE_FONTS, "Monofamily")
    prs = deck()
    prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2)).text_frame.text = "Sample text " * 4
    prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.name = "Monofamily"
    prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    prs.save(tmp_path / "f.pptx")
    run([rpptx_cli, "thumbnail", tmp_path / "f.pptx", "-o", tmp_path / "a.png"], check=True)
    run([rpptx_cli, "thumbnail", tmp_path / "f.pptx", "--font-dir", fonts, "-o", tmp_path / "b.png"], check=True)
    assert digest((tmp_path / "a.png").read_bytes()) != digest((tmp_path / "b.png").read_bytes())
    res = run([rpptx_cli, "thumbnail", tmp_path / "f.pptx", "--font-dir", tmp_path / "none", "-o", tmp_path / "c.png"])
    assert res.returncode == 1 and "does not exist" in res.stderr


# ---------------------------------------------------------------- silent failures the 2026-10-10 pins closed
def test_a_bare_int_character_spacing_raises():
    """A bare int is EMU: 2 would round to spc="0", nothing visible; the other length setters raise naming Pt."""
    prs = rpptx.Presentation()
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, 914400, 914400).text_frame.text = "x"
    with pytest.raises(ValueError, match="Pt"):
        prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0].font.spacing = 2


@pytest.mark.parametrize("value", [0, -1, Pt(0)])
def test_a_line_spacing_of_zero_or_less_raises(value):
    prs = rpptx.Presentation()
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, 914400, 914400).text_frame.text = "x"
    with pytest.raises(ValueError, match="line_spacing must be a positive"):
        prs.slides[0].shapes[0].text_frame.paragraphs[0].line_spacing = value


def test_removing_the_last_footer_placeholder_retires_its_handle():
    prs = rpptx.Presentation()
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, 914400, 914400)
    prs.set_header_footer(slide_number=True, hide_on_title=False)
    number = prs.slides[0].shapes[-1]                        # the slide-number placeholder, the last shape
    prs.slides[0].header_footer.slide_number = False
    with pytest.raises(rpptx.StaleElementError):
        number.left


# ---------------------------------------------------------------- the 2026-10-12 pin: charts, import, export, fonts
def chart_data(categories, *series):
    from rpptx.chart.data import CategoryChartData
    data = CategoryChartData()
    data.categories = categories
    for name, values in series:
        data.add_series(name, values)
    return data


def test_charts_are_added_read_back_and_their_data_replaced(rpptx_cli, tmp_path):
    from rpptx.enum.chart import XL_CHART_TYPE
    prs = deck(5)
    frame = prs.slides[0].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5),
                                           chart_data(["A", "B"], ("Cost", (236, 230)), ("Risk", (3, 2))))
    assert frame.has_chart and not prs.slides[0].shapes.title.has_chart
    prs.save(tmp_path / "c.pptx")
    p = pptx.Presentation(tmp_path / "c.pptx")                            # python-pptx reads an editable chart
    chart = next(sh for sh in p.slides[0].shapes if sh.has_chart).chart
    assert chart.chart_type == XL_CHART_TYPE.COLUMN_CLUSTERED and list(chart.plots[0].categories) == ["A", "B"]
    chart.plots[0].series[0].format.fill.solid()                          # formatting the replacement must keep
    chart.plots[0].series[0].format.fill.fore_color.rgb = pptx.dml.color.RGBColor(0x7B, 0x1E, 0x3A)
    p.save(tmp_path / "styled.pptx")
    prs = rpptx.Presentation(tmp_path / "styled.pptx")
    chart = next(sh for sh in prs.slides[0].shapes if sh.has_chart).chart
    assert (chart.chart_type, chart.categories, [(s.name, s.values) for s in chart.series]) == (
        XL_CHART_TYPE.COLUMN_CLUSTERED, ("A", "B"), [("Cost", (236.0, 230.0)), ("Risk", (3.0, 2.0))])
    chart.replace_data(chart_data(["A", "B", "C"], ("Cost", (236, 230, 251))))
    assert (chart.categories, [(s.name, s.values) for s in chart.series]) == (("A", "B", "C"), [("Cost", (236.0, 230.0, 251.0))])
    prs.save(tmp_path / "r.pptx")
    after = next(sh for sh in pptx.Presentation(tmp_path / "r.pptx").slides[0].shapes if sh.has_chart).chart
    assert [list(s.values) for s in after.plots[0].series] == [[236.0, 230.0, 251.0]]
    assert str(after.plots[0].series[0].format.fill.fore_color.rgb) == "7B1E3A"
    with zipfile.ZipFile(tmp_path / "r.pptx") as z:
        assert any(n.startswith("ppt/embeddings/") and n.endswith(".xlsx") for n in z.namelist())
    assert run([rpptx_cli, "validate", tmp_path / "r.pptx"]).returncode == 0


def test_a_chart_rpptx_does_not_author_is_read_with_its_missing_points(tmp_path):
    """A stacked bar chart from python-pptx: rpptx authors neither horizontal nor stacked bars."""
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE as PPTX_CHART
    from rpptx.enum.chart import XL_CHART_TYPE
    assert sorted(m.name for m in XL_CHART_TYPE) == sorted(
        ["AREA", "COLUMN_CLUSTERED", "DOUGHNUT", "LINE_MARKERS", "PIE", "RADAR_MARKERS", "XY_SCATTER_LINES"])
    p = pptx.Presentation()
    data = CategoryChartData()
    data.categories = ["A", "B"]
    data.add_series("Cost", (1, None))
    p.slides.add_slide(p.slide_layouts[6]).shapes.add_chart(PPTX_CHART.BAR_STACKED, 0, 0, Inches(4), Inches(3), data)
    p.save(tmp_path / "bar.pptx")
    chart = next(sh for sh in rpptx.Presentation(tmp_path / "bar.pptx").slides[0].shapes if sh.has_chart).chart
    assert (chart.chart_type, chart.categories, [(s.name, s.values) for s in chart.series]) == (
        None, ("A", "B"), [("Cost", (1.0, None))])


def jump_deck():
    """Three titled slides from python-pptx's default template; slide 1 jumps to slide 3, slide 2 has a chart."""
    from rpptx.enum.chart import XL_CHART_TYPE
    src = deck(1, 1, 1)
    for k, s in enumerate(src.slides):
        s.shapes.title.text = f"Source {k + 1}"
    src.slides[0].shapes.title.click_action.target_slide = src.slides[2]
    src.slides[1].shapes.add_chart(XL_CHART_TYPE.PIE, Inches(1), Inches(2), Inches(4), Inches(3),
                                   chart_data(["x", "y"], ("Share", (1, 3))))
    src.add_comment_author(id="{5B5B5B5B-0000-4000-8000-000000000001}", name="Claude", user_id="Claude", provider_id="None")
    src.slides[1].notes_text = "Notes travel too"
    src.slides[1].add_comment(id="{5B5B5B5B-0000-4000-8000-000000000002}", author_id="{5B5B5B5B-0000-4000-8000-000000000001}",
                              created="2026-10-12T10:00:00Z", text="Check the split")
    return src


def test_import_slides_with_charts_comments_jumps_and_source_formatting(rpptx_cli, tmp_path):
    src = jump_deck()
    prs = deck(1)
    prs.slides[0].shapes.title.text = "Own"
    copies = prs.slides.import_slides(src, index=1)                       # every slide: the jump between copies holds
    assert [s.shapes.title.text for s in prs.slides] == ["Own", "Source 1", "Source 2", "Source 3"]
    assert copies[0].shapes.title.click_action.target_slide == prs.slides[3]
    assert prs.slides[2].comments[0].text == "Check the split" and prs.slides[2].notes_text == "Notes travel too"
    assert [sh.chart.categories for sh in prs.slides[2].shapes if sh.has_chart] == [("x", "y")]
    with pytest.warns(UserWarning):                                       # its target is left behind: dropped
        prs.slides.import_slide(src.slides[0])
    assert prs.slides[len(prs.slides) - 1].shapes.title.click_action.target_slide is None
    prs.save(tmp_path / "m.pptx")
    assert run([rpptx_cli, "validate", tmp_path / "m.pptx"]).returncode == 0
    other = rpptx.Presentation()
    other.slide_layouts[1].name = "Agenda"                                # no layout of that name in the destination
    other.slide_master.theme.colors["accent1"] = "7B1E3A"
    other.slides.add_slide(other.slide_layouts[1]).shapes.title.text = "Kept"
    other = rpptx.Presentation.from_bytes(other.to_bytes())              # gap import-stale-source-formatting
    dest = deck(1)
    dest.slides.import_slides(rpptx.Presentation.from_bytes(jump_deck().to_bytes()), keep_source_formatting=True)
    assert len(dest.slide_masters) == 1                                   # an equal master and layouts are reused
    with pytest.raises(rpptx.RpptxError):
        dest.slides.import_slide(other.slides[0])
    kept = dest.slides.import_slide(other.slides[0], keep_source_formatting=True)
    assert kept.slide_layout.name == "Agenda" and kept.shapes.title.text == "Kept"
    assert [str(m.theme.colors["accent1"]) for m in dest.slide_masters] == ["4F81BD", "7B1E3A"]
    with pytest.raises(rpptx.RpptxError, match="exclude each other"):
        dest.slides.import_slide(other.slides[0], layout=dest.slide_layouts[1], keep_source_formatting=True)
    sectioned = deck(1, 1)
    sectioned.set_sections([("Intro", [0]), ("Body", [1])])
    dest = deck(1)
    dest.slides.import_slides(sectioned, keep_sections=True)
    assert [(s.name, s.slide_indices) for s in dest.sections][1:] == [("Intro", [1]), ("Body", [2])]


@pytest.mark.gap("import-stale-source-formatting")
def test_import_keeps_the_source_formatting_edited_in_memory():
    other = rpptx.Presentation()
    other.slide_layouts[1].name = "Agenda"
    other.slide_master.theme.colors["accent1"] = "7B1E3A"
    other.slides.add_slide(other.slide_layouts[1]).shapes.title.text = "Kept"
    dest = deck(1)
    dest.slides.import_slides(other, keep_source_formatting=True)
    assert dest.slides[1].slide_layout.name == "Agenda"
    assert [str(m.theme.colors["accent1"]) for m in dest.slide_masters] == ["4F81BD", "7B1E3A"]


def test_cli_slide_import(rpptx_cli, tmp_path):
    jump_deck().save(tmp_path / "src.pptx")
    deck(1).save(tmp_path / "dst.pptx")
    res = run([rpptx_cli, "slide", "import", tmp_path / "dst.pptx", tmp_path / "src.pptx", "--slides", "1-2", "--at", "1",
               "--json", "-o", tmp_path / "out.pptx"])
    assert res.returncode == 0, res.stderr
    record = json.loads(res.stdout)
    assert (record["source_slides"], record["slides"]) == ([1, 2], [1, 2])
    assert record["warnings"] and "Warning:" in res.stderr                # slide 1's jump to slide 3 is dropped
    assert titles(tmp_path / "out.pptx")[:2] == ["Source 1", "Source 2"]
    res = run([rpptx_cli, "slide", "import", tmp_path / "dst.pptx", tmp_path / "src.pptx", "--keep-source-formatting",
               "--keep-sections", "-o", tmp_path / "all.pptx"])
    assert res.returncode == 0 and len(pptx.Presentation(tmp_path / "all.pptx").slides) == 4
    res = run([rpptx_cli, "slide", "import", tmp_path / "dst.pptx", tmp_path / "src.pptx", "-o", tmp_path / "out.pptx"])
    assert res.returncode == 1 and "already exists" in res.stderr


def test_markdown_and_html_export(rpptx_cli, tmp_path):
    prs = deck(1, 1, 1)
    prs.slides[0].shapes.title.text = "Options"
    body = prs.slides[0].placeholders[1].text_frame
    body.text = "Option A"
    body.paragraphs[0].runs[0].font.bold = True
    sub = body.add_paragraph()
    sub.text, sub.level = "cheaper", 1
    body.add_paragraph().add_run().text = "Survey"
    body.paragraphs[2].runs[0].hyperlink.address = "https://example.com/s"
    prs.slides[0].notes_text = "Say the cost."
    prs.slides[1].shapes.title.text = "Hidden one"
    prs.slides[1].hidden = True
    table = prs.slides[2].shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for (r, c), t in {(0, 0): "Item", (0, 1): "Cost", (1, 0): "A", (1, 1): "236"}.items():
        table.cell(r, c).text = t
    prs.slides[2].shapes.add_picture(io.BytesIO(png_bytes()), Inches(6), Inches(2)).alt_text = "Deck photo"
    md = prs.to_markdown(notes=True, skip_hidden=True)
    assert md.startswith("# Options") and "Hidden one" not in md
    assert "- **Option A**\n  - cheaper" in md and "[Survey](https://example.com/s)" in md
    assert "> Say the cost." in md and "| Item | Cost |" in md and "![Deck photo]" in md
    assert "Hidden one" in prs.to_markdown() and "Say the cost." not in prs.to_markdown()
    html = prs.to_html(notes=True)
    assert '<section id="slide-1">' in html and "<h1>Options</h1>" in html and "<aside" in html
    assert "<strong>Option A</strong>" in html and 'alt="Deck photo"' in html and "data:image/png;base64," in html
    prs.save(tmp_path / "d.pptx")
    run([rpptx_cli, "convert", tmp_path / "d.pptx", "--to", "md", "--notes", "--skip-hidden", "-o", tmp_path / "d.md"], check=True)
    assert (tmp_path / "d.md").read_text() == md
    run([rpptx_cli, "convert", tmp_path / "d.pptx", "--to", "html", "-o", tmp_path / "d.html"], check=True)
    assert (tmp_path / "d.html").read_text() == prs.to_html()
    for fmt, flag in (("pdf", "--notes"), ("png", "--skip-hidden")):
        res = run([rpptx_cli, "convert", tmp_path / "d.pptx", "--to", fmt, flag, "-o", tmp_path / f"d.{fmt}"])
        assert res.returncode != 0 and not (tmp_path / f"d.{fmt}").exists()


def test_embed_list_and_remove_fonts(rpptx_cli, tmp_path, free_font):
    prs = deck()
    licence = rpptx.FontEmbeddingLicense(authorized=True, identity=FREE_FONT_LICENSE)
    assert prs.embed_true_type_fonts is False and prs.embedded_fonts == ()
    prs.embed_font("Aileron", free_font, "regular", license=licence)
    prs.embed_font("Aileron", free_font.read_bytes(), "bold", license=licence)
    assert prs.embed_true_type_fonts is True                  # PowerPoint uses embedded fonts only with this flag
    prs.save(tmp_path / "f.pptx")
    assert re.search(r'<p:presentation [^>]*embedTrueTypeFonts="1"', part(tmp_path / "f.pptx", "ppt/presentation.xml").decode())
    fonts = rpptx.Presentation(tmp_path / "f.pptx").embedded_fonts
    assert [(t.typeface, [(f.kind, f.format, f.data == free_font.read_bytes()) for f in t.faces]) for t in fonts] == [
        ("Aileron", [("regular", "eot", True), ("bold", "eot", True)])]
    assert all(f.part_name.endswith(".fntdata") and f.stored_len > len(f.data) for f in fonts[0].faces)
    with zipfile.ZipFile(tmp_path / "f.pptx") as z:
        eot = z.read(fonts[0].faces[0].part_name.lstrip("/"))
    os2 = free_font.read_bytes().index(b"OS/2")                            # the table record: tag, checksum, offset
    offset = int.from_bytes(free_font.read_bytes()[os2 + 8:os2 + 12], "big")
    fs_type = free_font.read_bytes()[offset + 8:offset + 10]
    assert eot[32:34] == fs_type[::-1] and int.from_bytes(eot[12:16], "little") == 0   # its fsType; not compressed
    prs = rpptx.Presentation(tmp_path / "f.pptx")
    assert prs.remove_embedded_font("Aileron", "bold").kind == "bold"
    assert prs.remove_embedded_font("Aileron", "italic") is None
    assert prs.remove_font("Aileron").typeface == "Aileron" and prs.embedded_fonts == ()
    prs.save_subset_fonts = False
    prs.embed_true_type_fonts = False
    assert (prs.save_subset_fonts, prs.embed_true_type_fonts) == (False, False)
    prs.save(tmp_path / "none.pptx")
    assert run([rpptx_cli, "validate", tmp_path / "f.pptx"]).returncode == 0
    with zipfile.ZipFile(tmp_path / "none.pptx") as z:
        assert not [n for n in z.namelist() if n.endswith(".fntdata")]      # the parts go with the last face


def test_a_deck_font_is_embedded_only_with_an_explicit_licence(tmp_path, free_font):
    prs = deck()
    before = prs.to_bytes()
    for licence in (rpptx.FontEmbeddingLicense(authorized=False, identity="CC0-1.0: Aileron"),
                    rpptx.FontEmbeddingLicense(authorized=True, identity=" ")):
        with pytest.raises(rpptx.RpptxError):
            prs.embed_font("Aileron", free_font, license=licence)
    with pytest.raises(TypeError):
        prs.embed_font("Aileron", free_font)
    assert digest(prs.to_bytes()) == digest(before)


def test_cli_font_commands(rpptx_cli, tmp_path, free_font):
    deck().save(tmp_path / "d.pptx")
    res = run([rpptx_cli, "font", "embed", tmp_path / "d.pptx", free_font, "--name", "Aileron",
               "--license-identity", FREE_FONT_LICENSE, "-o", tmp_path / "no.pptx"])
    assert res.returncode == 1 and "authorization" in res.stderr and not (tmp_path / "no.pptx").exists()
    res = run([rpptx_cli, "font", "embed", tmp_path / "d.pptx", free_font, "--name", "Aileron", "--license-authorized",
               "--license-identity", FREE_FONT_LICENSE, "-o", tmp_path / "e.pptx", "--json"])
    assert res.returncode == 0, res.stderr
    record = json.loads(res.stdout)
    assert (record["name"], record["face"], record["byte_count"]) == ("Aileron", "regular", len(free_font.read_bytes()))
    listed = json.loads(run([rpptx_cli, "font", "list", "--json", tmp_path / "e.pptx"], check=True).stdout)
    face = listed["fonts"][0]["embedded"][0]
    assert (listed["fonts"][0]["name"], face["face"], face["format"], face["byte_count"]) == (
        "Aileron", "regular", "eot", len(free_font.read_bytes()))
    assert listed["settings"] == {"embed_true_type_fonts": True, "save_subset_fonts": True}   # the default deck's
    res = run([rpptx_cli, "font", "settings", tmp_path / "e.pptx", "--save-subset-fonts", "false", "-o", tmp_path / "s.pptx",
               "--json"])
    assert json.loads(res.stdout)["settings"] == {"embed_true_type_fonts": True, "save_subset_fonts": False}
    run([rpptx_cli, "font", "remove", tmp_path / "s.pptx", "--name", "Aileron", "--face", "regular", "-o", tmp_path / "r.pptx"],
        check=True)
    assert rpptx.Presentation(tmp_path / "r.pptx").embedded_fonts == ()
    res = run([rpptx_cli, "font", "remove", tmp_path / "r.pptx", "--name", "Aileron", "-o", tmp_path / "g.pptx"])
    assert res.returncode == 1 and "not embedded" in res.stderr
    res = run([rpptx_cli, "font", "remove", tmp_path / "s.pptx", "--name", "Aileron", "-o", tmp_path / "r.pptx"])
    assert res.returncode == 1 and "already exists" in res.stderr


def nested_groups(depth):
    """The spTree XML of `depth` groups nested in one another."""
    group = '<p:grpSp><p:nvGrpSpPr><p:cNvPr id="{id}" name="G{id}"/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr/>'
    return "".join(group.format(id=k + 10) for k in range(depth)) + "</p:grpSp>" * depth


def test_slides_nested_past_16_levels_are_refused(tmp_path):
    for depth, opens in ((16, True), (17, False)):
        src = io.BytesIO(deck().to_bytes())
        with zipfile.ZipFile(src) as z, zipfile.ZipFile(tmp_path / f"{depth}.pptx", "w") as out:
            for info in z.infolist():
                data = z.read(info.filename)
                if info.filename == "ppt/slides/slide1.xml":
                    data = data.replace(b"</p:spTree>", nested_groups(depth).encode() + b"</p:spTree>")
                out.writestr(info, data)
        if opens:
            assert rpptx.Presentation(tmp_path / f"{depth}.pptx").slides[0].shapes is not None
        else:
            with pytest.raises(rpptx.XmlError, match="nesting exceeds 16 levels"):
                rpptx.Presentation(tmp_path / f"{depth}.pptx")
    prs = deck()
    shapes = prs.slides[0].shapes
    for _ in range(16):
        group = shapes.add_group_shape()
        shapes = group.shapes
    with pytest.raises(rpptx.RpptxError, match="nesting exceeds 16 levels"):
        shapes.add_group_shape()


def test_what_the_new_edits_retire(free_font):
    """Charts and appended imports retire nothing; a font embedded or removed, or an import at an index, retires the
    slides and everything below; keeping the source formatting retires the layouts."""
    from rpptx.enum.chart import XL_CHART_TYPE
    licence = rpptx.FontEmbeddingLicense(authorized=True, identity=FREE_FONT_LICENSE)

    def source(keep=False):
        other = rpptx.Presentation()
        if keep:
            other.slide_layouts[1].name = "Agenda"
            other = rpptx.Presentation.from_bytes(other.to_bytes())
        other.slides.add_slide(other.slide_layouts[1])
        return other

    def replace(prs):
        frame = next(sh for sh in prs.slides[0].shapes if sh.has_chart)
        frame.chart.replace_data(chart_data(["a", "b"], ("s", (1, 2))))

    edits = {
        "add_chart": lambda p: p.slides[0].shapes.add_chart(XL_CHART_TYPE.PIE, 0, 0, Inches(1), Inches(1),
                                                            chart_data(["a"], ("s", (1,)))),
        "replace_data": replace,
        "import at the end": lambda p: p.slides.import_slides(source()),
        "font flags": lambda p: setattr(p, "save_subset_fonts", False),
        "embed_font": lambda p: p.embed_font("Aileron", free_font, license=licence),
        "remove_font": lambda p: (p.embed_font("Aileron", free_font, license=licence),
                                  p.slides[0].shapes.title.text_frame.paragraphs[0].runs[0].text,  # re-fetched: valid
                                  p.remove_font("Aileron")),
        "import at an index": lambda p: p.slides.import_slides(source(), index=0),
        "keep_source_formatting": lambda p: p.slides.import_slides(source(keep=True), keep_source_formatting=True),
    }
    retired = {}
    for name, edit in edits.items():
        prs = deck(1)
        prs.slides[0].shapes.add_chart(XL_CHART_TYPE.PIE, 0, 0, Inches(1), Inches(1), chart_data(["a"], ("s", (1,))))
        prs.slides[0].shapes.title.text = "t"
        slide = prs.slides[0]
        held = {"slide": slide, "shape": slide.shapes.title, "run": slide.shapes.title.text_frame.paragraphs[0].runs[0],
                "layout": prs.slide_layouts[1]}
        edit(prs)
        retired[name] = []
        for kind, handle in held.items():
            try:
                handle.name if kind in ("shape", "layout") else handle.slide_id if kind == "slide" else handle.text
            except rpptx.StaleElementError:
                retired[name].append(kind)
    everything = ["slide", "shape", "run"]
    assert retired == {"add_chart": [], "replace_data": [], "import at the end": [], "font flags": [],
                       "embed_font": everything, "remove_font": everything, "import at an index": everything,
                       "keep_source_formatting": ["layout"]}


def embedded_font_deck(path, free_font):
    prs = deck()
    box = prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "Embedded face text"
    box.text_frame.paragraphs[0].runs[0].font.name = "Aileron"
    prs.embed_font("Aileron", free_font, license=rpptx.FontEmbeddingLicense(authorized=True, identity=FREE_FONT_LICENSE))
    prs.save(path)
    return path


def test_an_embedded_font_is_drawn_by_libreoffice_and_by_rpptx_given_the_file(rpptx_cli, tmp_path, free_font):
    src = embedded_font_deck(tmp_path / "f.pptx", free_font)
    (tmp_path / "fonts").mkdir()
    (tmp_path / "fonts" / free_font.name).write_bytes(free_font.read_bytes())
    assert pdf_uses_font(rpptx.Presentation(src).to_pdf(font_dir=tmp_path / "fonts"), b"Aileron")   # the workaround
    run([rpptx_cli, "convert", src, "--to", "pdf", "--font-dir", tmp_path / "fonts", "-o", tmp_path / "cli.pdf"], check=True)
    assert pdf_uses_font((tmp_path / "cli.pdf").read_bytes(), b"Aileron")
    assert pdf_uses_font(libreoffice_pdf(src, tmp_path), b"Aileron")


@pytest.mark.gap("render-embedded-font")
def test_rpptx_draws_an_embedded_font(tmp_path, free_font):
    src = embedded_font_deck(tmp_path / "f.pptx", free_font)
    assert pdf_uses_font(rpptx.Presentation(src).to_pdf(), b"Aileron")
