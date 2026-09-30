"""docx, reading: text, structure, JSON views, metadata, hyperlinks, sections, styles."""
import json
import zipfile

import docx
import pytest
import rdocx
from docx.oxml.ns import qn
from docx.shared import Twips

from builders import every_story_docx
from conftest import run


def test_cli_text_plain(rdocx_cli, report_docx):
    out = run([rdocx_cli, "text", report_docx], check=True).stdout
    assert "This report records the principal inspection" in out
    assert "Figure 1. Cross-section through pier 2" in out


def test_cli_text_json_schema(rdocx_cli, report_docx):
    data = json.loads(run([rdocx_cli, "text", "--json", report_docx], check=True).stdout)
    assert data["schema"] == 1
    paras = data["paragraphs"]
    first = paras[0]
    assert {"body_index", "path", "style", "numbering", "runs", "text"} <= set(first)
    assert {"index", "text", "formatting"} <= set(first["runs"][0])
    # table cell paragraphs are listed too, with a path through the table
    assert any(len(p["path"]) > 1 for p in paras)
    headings = [p["text"] for p in paras if p["style"] == "Heading1"]
    assert headings[:2] == ["Summary", "Scope and method"]


def test_cli_text_includes_block_content_control(rdocx_cli, report_docx):
    out = run([rdocx_cli, "text", report_docx], check=True).stdout
    assert "Access to the river was agreed with the harbour office" in out


def test_cli_inspect_json(rdocx_cli, report_docx):
    data = json.loads(run([rdocx_cli, "inspect", "--json", report_docx], check=True).stdout)
    assert data["tables"] == 3
    assert data["metadata"]["title"].startswith("Riverton Footbridge")
    assert "Heading1" in data["styles_used"]


def test_cli_layout_json(rdocx_cli, report_docx):
    data = json.loads(run([rdocx_cli, "layout", "--json", report_docx], check=True).stdout)
    items = data["body_items"]
    assert items[0]["fragments"][0]["physical_page"] == 1
    pages = {f["physical_page"] for it in items for f in it["fragments"]}
    assert max(pages) >= 10


def test_python_structure(report_docx):
    doc = rdocx.Document.open(report_docx)
    assert len(doc.tables) == 3
    assert len(doc.tables[0].rows) == 57
    assert doc.tables[0].cell(0, 0).text == "Ref"
    sec = doc.sections[0]
    assert sec.orientation == "portrait" and sec.page_width == 7772400 and sec.different_first_page
    ids = {s.style_id for s in doc.styles}
    assert {"Normal", "Heading1", "Caption", "TOC1", "Hyperlink"} <= ids
    kinds = {(v.kind, v.variant) for v in doc.header_footer_variants}
    assert ("footer", "default") in kinds and ("header", "first") in kinds


def test_python_paragraphs_and_runs(report_docx):
    doc = rdocx.Document.open(report_docx)
    heads = [p.text for p in doc.paragraphs if p.style == "Heading1"]
    assert "Inspection findings" in heads
    p = next(p for p in doc.paragraphs if p.text.startswith("The main findings are"))
    assert p.runs[0].text == "The main findings are:"
    bullets = [p for p in doc.paragraphs if p.numbering]
    assert bullets and bullets[0].numbering[0] == 1


def test_python_hyperlinks(report_docx):
    links = rdocx.Document.open(report_docx).hyperlinks
    assert len(links) == 6
    assert all(h.url.startswith("https://example.org/") for h in links)


def test_find_content_indices_includes_toc_block(report_docx):
    """A heading's text is found twice: in the body and in the table of contents block. Never delete by it."""
    doc = rdocx.Document.open(report_docx)
    hits = doc.find_content_indices("Recommendations")
    assert len(hits) == 2 and 7 in hits


def test_python_story_items_cover_headers_and_footers(report_docx):
    doc = rdocx.Document.open(report_docx)
    kinds = {it.story.kind for it in doc.story_items}
    assert {"body", "header", "footer"} <= kinds


def test_python_comments_and_revisions_empty(report_docx):
    doc = rdocx.Document.open(report_docx)
    assert list(doc.comments) == [] and list(doc.revisions) == []


def test_convert_markdown_and_html(rdocx_cli, report_docx, tmp_path):
    md, html = tmp_path / "r.md", tmp_path / "r.html"
    run([rdocx_cli, "convert", report_docx, "--to", "md", "-o", md], check=True)
    run([rdocx_cli, "convert", report_docx, "--to", "html", "-o", html], check=True)
    assert "# Summary" in md.read_text() or "Summary" in md.read_text()
    assert "<table" in html.read_text()


# ---------------------------------------------------------------- stories other than the body, tracked changes
def test_story_items_read_every_story(tmp_path):
    """Headers, footers (each variant), footnotes, text boxes and table cells are read through story_items."""
    path, _ = every_story_docx(tmp_path / "s.docx")
    texts = {(it.story.kind, it.text) for it in rdocx.Document.open(path).story_items if it.text and "NEEDLE" in it.text}
    assert {("table_cell", "Cell NEEDLE"), ("text_box", "Text box NEEDLE"), ("header", "Header NEEDLE"),
            ("footer", "Footer NEEDLE"), ("footer", "First footer NEEDLE"), ("footnote", " Footnote NEEDLE."),
            ("endnote", " Endnote NEEDLE.")} <= texts


def test_text_json_is_the_accepted_view(rdocx_cli, tmp_path):
    """`text --json` and Paragraph.text show tracked insertions and hide deletions."""
    path, _ = every_story_docx(tmp_path / "s.docx")
    data = json.loads(run([rdocx_cli, "text", "--json", path], check=True).stdout)
    assert data["revision_view"] == "accepted"
    assert "Tracked: ins NEEDLE" in [p["text"] for p in data["paragraphs"]]
    assert "Tracked: ins NEEDLE" in [p.text for p in rdocx.Document.open(path).paragraphs]


def test_cli_plain_text_keeps_tracked_insertions(rdocx_cli, tmp_path):
    path, _ = every_story_docx(tmp_path / "s.docx")
    assert "ins NEEDLE" in run([rdocx_cli, "text", path], check=True).stdout


@pytest.mark.parametrize("view", ["plain", "json", "md"])
def test_cli_text_includes_headers_footers_and_notes(rdocx_cli, tmp_path, view):
    path, _ = every_story_docx(tmp_path / "s.docx")
    if view == "plain":
        out = run([rdocx_cli, "text", path], check=True).stdout
        assert "--- header (/word/header1.xml) ---\nHeader NEEDLE\n" in out
    elif view == "json":
        out = run([rdocx_cli, "text", "--json", path], check=True).stdout
        data = json.loads(out)
        assert data["scope"] == "all-supported-stories"
        assert {"header", "footer", "footnote", "endnote", "text_box"} <= {s["kind"] for s in data["stories"]}
    else:
        run([rdocx_cli, "convert", path, "--to", "md", "-o", tmp_path / "s.md"], check=True)
        out = (tmp_path / "s.md").read_text()
    assert "Header NEEDLE" in out and "Footnote NEEDLE" in out and "Text box NEEDLE" in out


# ---------------------------------------------------------------- validation
def broken(src, dst, name, change):
    with zipfile.ZipFile(src) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(dst, "w") as z:
        for info, data in items:
            z.writestr(info, change(data) if info.filename == name else data)
    return dst


def small_with_header(path):
    d = docx.Document()
    d.add_paragraph("x")
    d.sections[0].header.paragraphs[0].text = "Head"
    d.save(path)
    return path


def test_validate_catches_a_truncated_document_part(rdocx_cli, tmp_path):
    src = small_with_header(tmp_path / "a.docx")
    bad = broken(src, tmp_path / "b.docx", "word/document.xml", lambda b: b[: len(b) // 2])
    assert run([rdocx_cli, "validate", bad]).returncode != 0


def test_validate_catches_a_truncated_header(rdocx_cli, tmp_path):
    src = small_with_header(tmp_path / "a.docx")
    with zipfile.ZipFile(src) as z:
        header = next(n for n in z.namelist() if n.startswith("word/header"))
    bad = broken(src, tmp_path / "b.docx", header, lambda b: b[: len(b) // 2])
    assert run([rdocx_cli, "validate", bad]).returncode != 0


def test_validate_catches_a_dangling_style_id(rdocx_cli, tmp_path):
    src = small_with_header(tmp_path / "a.docx")
    bad = broken(src, tmp_path / "b.docx", "word/document.xml",
                 lambda b: b.replace(b"<w:p>", b'<w:p><w:pPr><w:pStyle w:val="NoSuchStyle"/></w:pPr>', 1))
    assert b"NoSuchStyle" in zipfile.ZipFile(bad).read("word/document.xml")
    assert run([rdocx_cli, "validate", bad]).returncode != 0


def test_cli_text_keeps_the_body_when_a_header_is_truncated(rdocx_cli, tmp_path):
    src = small_with_header(tmp_path / "a.docx")
    with zipfile.ZipFile(src) as z:
        header = next(n for n in z.namelist() if n.startswith("word/header"))
    bad = broken(src, tmp_path / "b.docx", header, lambda b: b[: len(b) // 2])
    res = run([rdocx_cli, "text", bad])
    assert res.returncode == 0 and res.stdout.startswith("x") and "Head" not in res.stdout and res.stderr
    data = json.loads(run([rdocx_cli, "text", "--json", bad], check=True).stdout)
    assert (data["scope"], data["stories"]) == ("main", [])


def test_reopen_with_story_items_catches_a_truncated_header(tmp_path):
    """A second check beside `validate`: a story read fails on a truncated header."""
    src = small_with_header(tmp_path / "a.docx")
    with zipfile.ZipFile(src) as z:
        header = next(n for n in z.namelist() if n.startswith("word/header"))
    bad = broken(src, tmp_path / "b.docx", header, lambda b: b[: len(b) // 2])
    with pytest.raises(rdocx.RdocxError):
        rdocx.Document.open(bad).story_items


# ---------------------------------------------------------------- measurements with a decimal part
def measured(path):
    """An indented paragraph with spacing, a 1x2 table with a row height and a width, a header with an indent: the
    measurements Google Docs writes with a floating-point tail."""
    d = docx.Document()
    p = d.add_paragraph("Indented paragraph")
    p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Twips(720), Twips(-227)
    p.paragraph_format.space_before, p.paragraph_format.line_spacing = Twips(120), 1.15
    t = d.add_table(rows=1, cols=2)
    t.rows[0].height = Twips(535)
    t.cell(0, 0).text = "Cell A"
    width = t._tbl.tblPr.find(qn("w:tblW"))
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), "8640")
    d.sections[0].header.paragraphs[0].text = "Header"
    d.sections[0].header.paragraphs[0].paragraph_format.left_indent = Twips(360)
    d.save(path)
    return path


def twips(emu):
    return emu / 635


# part, attribute as python-docx writes it with {} for the value, a decimal value, the integer it stands for, and
# where the binding reads it back in twips (None: not exposed)
DECIMAL = [
    ("document", 'w:gridCol w:w="{}"', "4320.0", 4320, lambda d: twips(d.tables[0].grid_widths[0])),
    ("document", 'w:gridCol w:w="{}"', "4319.999999999999", 4320, lambda d: twips(d.tables[0].grid_widths[0])),
    ("document", 'w:left="{}"', "720.0", 720, lambda d: twips(d.paragraphs[0].paragraph_format.left_indent)),
    ("document", 'w:hanging="{}"', "226.99999999999977", 227, lambda d: -twips(d.paragraphs[0].paragraph_format.first_line_indent)),
    ("document", 'w:before="{}"', "120.0", 120, lambda d: twips(d.paragraphs[0].paragraph_format.space_before)),
    ("document", 'w:trHeight w:val="{}"', "535.0000000000182", 535, lambda d: twips(d.tables[0].rows[0].height)),
    ("document", 'w:pgSz w:w="{}"', "12240.0", 12240, lambda d: twips(d.sections[0].page_width)),
    ("document", 'w:top="{}"', "1440.0", 1440, lambda d: twips(d.sections[0].margin_top)),
    ("document", 'w:tcW w:type="dxa" w:w="{}"', "4319.999999999999", 4320, lambda d: twips(d.tables[0].cell(0, 0).width)),
    ("document", 'w:tblW w:type="dxa" w:w="{}"', "8639.999999999999", 8640, lambda d: twips(d.tables[0].width)),
    ("styles", 'w:after="{}"', "200.0", 200, None),
    ("styles", '<w:sz w:val="{}"/>', "22.0", 22, None),
    ("numbering", 'w:left="{}"', "360.0", 360, None),
]
# the same cases that open today
DECIMAL_ACCEPTED = [
    ("document", 'w:line="{}"', "275.99999999999994", 276, lambda d: round(d.paragraphs[0].paragraph_format.line_spacing * 240)),
    ("document", 'w:tcW w:type="dxa" w:w="{}"', "4320.0", 4320, lambda d: twips(d.tables[0].cell(0, 0).width)),
    ("header1", 'w:left="{}"', "360.0", 360, None),
]


def with_values(src, dst, rows):
    """`src` with, for each row, the first occurrence of the attribute set to the row's decimal value."""
    with zipfile.ZipFile(src) as z:
        items = [(i, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            for part, template, value, integer, _ in rows:
                if info.filename == f"word/{part}.xml":
                    before = template.format(integer).encode()
                    assert before in data, (part, before)
                    data = data.replace(before, template.format(value).encode(), 1)
            z.writestr(info, data)
    return dst


def cases(rows, gap=None):
    return [pytest.param(*row, marks=[pytest.mark.gap(gap)] if gap else [], id=f"{row[0]} {row[1].format(row[2])}")
            for row in rows]


@pytest.mark.parametrize("part,template,value,integer,read", cases(DECIMAL) + cases(DECIMAL_ACCEPTED))
def test_decimal_measurement_is_read_as_the_nearest_integer(part, template, value, integer, read, rdocx_cli, tmp_path):
    """Google Docs writes measurements with a floating-point tail (`w:gridCol w:w="2210.0000000000005"`), which Word
    rounds. The file opens in the binding and in the CLI, and the value reads back as the nearest integer."""
    f = with_values(measured(tmp_path / "a.docx"), tmp_path / "b.docx", [(part, template, value, integer, read)])
    doc = rdocx.Document.open(f)
    if read:
        assert read(doc) == integer
    res = run([rdocx_cli, "text", f])
    assert res.returncode == 0 and "Indented paragraph" in res.stdout, res.stderr


def test_measurements_read_back_where_the_decimal_tests_look(tmp_path):
    """The control of the test above: each attribute is in the fixture, and its integer reads back there."""
    doc = rdocx.Document.open(measured(tmp_path / "a.docx"))
    for part, template, _, integer, read in DECIMAL + DECIMAL_ACCEPTED:
        assert template.format(integer).encode() in zipfile.ZipFile(tmp_path / "a.docx").read(f"word/{part}.xml")
        assert read is None or read(doc) == integer, template

