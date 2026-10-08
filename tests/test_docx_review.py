"""docx, review features: comments and threads, tracked changes, redline by comparison."""
import io
import json
import os
import re
import zipfile

import docx
import pytest
import rdocx
from docx.oxml.ns import qn
from PIL import Image

from builders import cell_text_docx, commented_thread, google_wrapped, rewrite_body, wrapped_run_docx
from conftest import STAMP, digest, part, run


def two_paragraphs(path, first="Alpha paragraph with some words here.", second="Beta paragraph."):
    d = docx.Document()
    d.add_paragraph(first)
    d.add_paragraph(second)
    d.save(path)
    return path


def anchored(path, cid):
    x = part(path, "word/document.xml").decode()
    m = re.search(r'<w:commentRangeStart w:id="%d"/>(.*?)<w:commentRangeEnd w:id="%d"/>' % (cid, cid), x, re.S)
    return "".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", m.group(1))) if m else None


def run_range(bi, first, last):
    return rdocx.RunRange(start=rdocx.RunPosition(body_index=bi, run_index=first),
                          end=rdocx.RunPosition(body_index=bi, run_index=last))


def noop_saved(src, dst):
    """The original passed through rdocx once (open, save): the workaround for empty-comments-reserialised."""
    rdocx.Document.open(src).save(dst)
    return dst


def compare(cli, a, b, out, *options):
    if os.path.exists(out):
        os.remove(out)
    return run([cli, "compare", a, b, "--author", "Reviewer", "--timestamp", STAMP, *options, "-o", out])


# ---------------------------------------------------------------- comments
def test_comment_thread_python(tmp_path):
    doc = rdocx.Document.open(two_paragraphs(tmp_path / "c.docx"))
    rng = rdocx.RunRange(start=rdocx.RunPosition(body_index=0, run_index=0),
                         end=rdocx.RunPosition(body_index=0, run_index=1))
    cid = doc.add_comment(rng, author="Reviewer", text="Check this.", initials="R", date=STAMP)
    rid = doc.reply_to(cid, author="Author", text="Done.")
    doc.resolve_comment(cid)
    got = [(c.id, c.author, c.text, c.parent_id, c.resolved) for c in doc.comments]
    assert got == [(cid, "Reviewer", "Check this.", None, True), (rid, "Author", "Done.", cid, False)]
    doc.save(tmp_path / "c2.docx")
    assert [(c.author, c.text) for c in docx.Document(tmp_path / "c2.docx").comments] == [("Reviewer", "Check this."), ("Author", "Done.")]
    again = rdocx.Document.open(tmp_path / "c2.docx")
    again.remove_comment(cid)
    assert len(again.comments) == 0


def test_comment_thread_cli(rdocx_cli, tmp_path):
    src = two_paragraphs(tmp_path / "c.docx")
    a, b, c = tmp_path / "a.docx", tmp_path / "b.docx", tmp_path / "d.docx"
    add = run([rdocx_cli, "comment", "add", src, "--start-paragraph", "0", "--start-run", "0", "--end-paragraph", "0",
               "--end-run", "1", "--author", "Reviewer", "--text", "Why?", "-o", a, "--json"], check=True)
    cid = json.loads(add.stdout)
    run([rdocx_cli, "comment", "reply", a, "--id", "0", "--author", "Author", "--text", "Because.", "-o", b], check=True)
    run([rdocx_cli, "comment", "resolve", b, "--id", "0", "-o", c], check=True)
    listed = json.loads(run([rdocx_cli, "comment", "list", "--json", c], check=True).stdout)["comments"]
    assert [x["text"] for x in listed] == ["Why?", "Because."] and listed[0]["resolved"] is True
    assert cid  # the operation record is JSON


def test_comment_ids_are_stable_across_save(tmp_path):
    doc = rdocx.Document.open(two_paragraphs(tmp_path / "c.docx"))
    a = doc.add_comment(run_range(1, 0, 1), author="A", text="on the second paragraph")
    b = doc.add_comment(run_range(0, 0, 1), author="A", text="on the first paragraph")
    doc.save(tmp_path / "d.docx")
    assert [(c.id, c.text) for c in rdocx.Document.open(tmp_path / "d.docx").comments] == [
        (a, "on the second paragraph"), (b, "on the first paragraph")]


def test_comment_date_is_written_only_when_given(tmp_path):
    doc = rdocx.Document.open(two_paragraphs(tmp_path / "c.docx"))
    doc.add_comment(run_range(0, 0, 1), author="A", text="undated")
    doc.add_comment(run_range(1, 0, 1), author="A", text="dated", date=STAMP)
    doc.save(tmp_path / "d.docx")
    assert [(c.text, c.date) for c in rdocx.Document.open(tmp_path / "d.docx").comments] == [("undated", None), ("dated", STAMP)]


def test_cli_comment_add_takes_a_date(rdocx_cli, tmp_path):
    src = two_paragraphs(tmp_path / "c.docx")
    res = run([rdocx_cli, "comment", "add", src, "--start-paragraph", "0", "--start-run", "0", "--end-paragraph", "0",
               "--end-run", "1", "--author", "R", "--text", "x", "--date", STAMP, "-o", tmp_path / "d.docx"])
    assert res.returncode == 0 and rdocx.Document.open(tmp_path / "d.docx").comments[0].date == STAMP


def test_comment_on_a_table_cell_through_a_story_range(tmp_path):
    d = docx.Document()
    d.add_paragraph("before")
    d.add_table(rows=1, cols=2).cell(0, 1).text = "cell text"
    d.save(tmp_path / "c.docx")
    doc = rdocx.Document.open(tmp_path / "c.docx")
    item = next(it for it in doc.story_items if it.kind == "paragraph" and it.text == "cell text")
    rng = rdocx.StoryRunRange(start=rdocx.StoryRunPosition(item=item, run_index=0),
                              end=rdocx.StoryRunPosition(item=item, run_index=1))
    cid = doc.add_comment(rng, author="A", text="on a cell", date=STAMP)
    doc.save(tmp_path / "d.docx")
    assert anchored(tmp_path / "d.docx", cid) == "cell text"


def test_comment_on_part_of_a_table_cell_paragraph_by_text(tmp_path):
    """add_comment_on_text numbers the anchor over the main story in document order, from 0, table cells
    included, and splits the cell's run with its format kept. split_run itself refuses a cell paragraph."""
    path = cell_text_docx(tmp_path / "c.docx")
    doc = rdocx.Document.open(path)
    with pytest.raises(ValueError, match="table cell"):
        doc.split_run(doc.tables[0].cell(0, 0).paragraphs[0], 0, 6)
    cid = doc.add_comment_on_text("beta", author="A", text="x", occurrence=1, date=STAMP)
    doc.save(tmp_path / "d.docx")
    assert anchored(tmp_path / "d.docx", cid) == "beta"
    runs = [r for r in rdocx.Document.open(tmp_path / "d.docx").tables[0].cell(0, 0).paragraphs[0].runs if r.text]
    assert [r.text for r in runs] == ["alpha ", "beta", " gamma"] and all(r.font.bold for r in runs)
    with pytest.raises(rdocx.RdocxError, match="occurs 4 times"):
        rdocx.Document.open(path).add_comment_on_text("beta", author="A", text="x", occurrence=4)


def test_cli_comment_on_part_of_a_table_cell_paragraph_by_text(rdocx_cli, tmp_path):
    path = cell_text_docx(tmp_path / "c.docx")
    run([rdocx_cli, "comment", "add", path, "--anchor", "beta", "--occurrence", "1", "--author", "A", "--text", "x",
         "--date", STAMP, "-o", tmp_path / "d.docx"], check=True)
    assert anchored(tmp_path / "d.docx", 0) == "beta"
    assert [r.text for r in rdocx.Document.open(tmp_path / "d.docx").tables[0].cell(0, 0).paragraphs[0].runs if r.text] \
        == ["alpha ", "beta", " gamma"]


@pytest.mark.parametrize("wrapper", ["sdt", "ins"])
def test_comment_run_index_counts_wrapped_runs(tmp_path, wrapper):
    path = wrapped_run_docx(tmp_path / "w.docx", wrapper)
    doc = rdocx.Document.open(path)
    assert [r.text for r in doc.paragraphs[0].runs] == ["before ", "TARGET", " after"]
    cid = doc.add_comment(run_range(0, 1, 2), author="A", text="x")
    doc.save(tmp_path / "c.docx")
    assert anchored(tmp_path / "c.docx", cid) == "TARGET"


def test_cli_comment_run_index_counts_wrapped_runs(rdocx_cli, tmp_path):
    path = wrapped_run_docx(tmp_path / "w.docx", "sdt")
    run([rdocx_cli, "comment", "add", path, "--start-paragraph", "0", "--start-run", "1", "--end-paragraph", "0",
         "--end-run", "2", "--author", "A", "--text", "x", "-o", tmp_path / "c.docx"], check=True)
    assert anchored(tmp_path / "c.docx", 0) == "TARGET"


def test_comment_cli_refuses_to_overwrite_its_input(rdocx_cli, tmp_path):
    src = two_paragraphs(tmp_path / "c.docx")
    before = src.read_bytes()
    res = run([rdocx_cli, "comment", "add", src, "--start-paragraph", "0", "--start-run", "0", "--end-paragraph", "0",
               "--end-run", "1", "--author", "R", "--text", "x", "-o", src])
    assert res.returncode != 0 and digest(src.read_bytes()) == digest(before)


# Comments of several paragraphs, as Word and Google Docs write them: each w15:commentEx row is keyed by the
# w14:paraId of the comment's LAST paragraph (and names its parent by the parent's last paragraph).
SEVERAL_PARAGRAPHS = [("Ada", ["1A000001"]), ("Ben", ["1B000001", "1B000002"]), ("Ada", ["2A000001", "2A000002"]),
                      ("Ben", ["2B000001"]), ("Ada", ["3A000001", "3A000002"])]
SEVERAL_PARAGRAPHS_EX = [("1A000001", None, 0), ("1B000002", "1A000001", 0), ("2A000002", None, 0),
                         ("2B000001", "2A000002", 0), ("3A000002", None, 1)]


def several_paragraph_comments(path):
    """Five anchored comments: a reply of two paragraphs, a reply to a parent of two paragraphs, a resolved
    comment of two paragraphs. Returns the path and the comment ids in part order."""
    doc = rdocx.Document()
    doc.add_paragraph("Anchor paragraph.")
    ids = [doc.add_comment(run_range(0, 0, 1), author=a, text="x", date=STAMP) for a, _ in SEVERAL_PARAGRAPHS]
    base = doc.to_bytes()
    comments = "".join(
        '<w:comment w:id="%d" w:author="%s" w:date="%s">%s</w:comment>'
        % (cid, author, STAMP, "".join('<w:p w14:paraId="%s"><w:r><w:t>%s</w:t></w:r></w:p>' % (p, p) for p in paras))
        for cid, (author, paras) in zip(ids, SEVERAL_PARAGRAPHS))
    replaced = {
        "word/comments.xml": '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                             'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml">%s</w:comments>' % comments,
        "word/commentsExtended.xml": '<w15:commentsEx xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml">%s'
                                     '</w15:commentsEx>' % "".join(
            '<w15:commentEx w15:paraId="%s"%s w15:done="%d"/>' % (p, ' w15:paraIdParent="%s"' % par if par else "", d)
            for p, par, d in SEVERAL_PARAGRAPHS_EX),
    }
    with zipfile.ZipFile(io.BytesIO(base)) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, replaced.get(item.filename, zin.read(item.filename)))
    return path, ids


def comment_ex_rows(path):
    """{paraId: (paraIdParent, done)} of word/commentsExtended.xml."""
    rows = re.findall(r"<w15:commentEx\b[^>]*/>", part(path, "word/commentsExtended.xml").decode())
    attr = lambda row, name: (re.search(r'w15:%s="([^"]*)"' % name, row) or [None, None])[1]
    return {attr(r, "paraId"): (attr(r, "paraIdParent"), attr(r, "done") in ("1", "true")) for r in rows}


def test_comment_of_several_paragraphs_threads_on_its_last_paragraph(tmp_path):
    src, ids = several_paragraph_comments(tmp_path / "m.docx")
    doc = rdocx.Document.open(src)
    assert [(c.parent_id, c.resolved) for c in doc.comments] == [
        (None, False), (ids[0], False), (None, False), (ids[2], False), (None, True)]
    reply = doc.reply_to(ids[2], author="Cy", text="first\nsecond", date=STAMP)
    doc.resolve_comment(ids[2])
    out = tmp_path / "out.docx"
    doc.save(out)
    rows = comment_ex_rows(out)
    assert rows["2A000002"] == (None, True) and "2A000001" not in rows
    body = re.search(r'<w:comment\b[^>]*w:id="%d"[^>]*>(.*?)</w:comment>' % reply,
                     part(out, "word/comments.xml").decode(), re.S).group(1)
    assert len(re.findall(r"<w:p[ >]", body)) == 2
    assert rows[re.findall(r'w14:paraId="([0-9A-Fa-f]+)"', body)[-1]][0] == "2A000002"
    reread = {c.id: (c.parent_id, c.text) for c in rdocx.Document.open(out).comments}
    assert reread[reply] == (ids[2], "first\nsecond")


# ---------------------------------------------------------------- comment anchors, moves, removals
def test_comment_anchor_text_through_google_wrappers(rdocx_cli, tmp_path):
    src, cid, rid = commented_thread(tmp_path / "c.docx")
    g = google_wrapped(src, tmp_path / "g.docx")
    assert part(g, "word/document.xml").count(b"goog_rdk_") == 3
    doc = rdocx.Document.open(g)
    assert [(c.id, c.anchor_text) for c in doc.comments] == [(cid, "paragraph"), (rid, None)]
    anchor = doc.comments[0].anchor
    assert isinstance(anchor, rdocx.StoryRunRange) and doc.comments[1].anchor is None
    item = anchor.start.item                                              # the paragraph inside the block control
    assert (item.story.kind, item.kind, item.index_path, item.direct_body_index) == ("body", "paragraph", (0, 0), 0)
    assert (anchor.start.run_index, anchor.end.run_index) == (1, 2)
    listed = json.loads(run([rdocx_cli, "comment", "list", "--json", g], check=True).stdout)["comments"]
    assert [(x["id"], x["anchor_text"]) for x in listed] == [(cid, "paragraph"), (rid, None)]
    assert listed[0]["anchor"]["story"]["kind"] == "body" and listed[0]["anchor"]["start"]["body_index"] == 0
    assert (listed[0]["anchor"]["start"]["run_index"], listed[0]["anchor"]["end"]["run_index"]) == (1, 2)
    assert listed[0]["reference"]["story"]["kind"] == "body" and listed[1]["anchor"] is listed[1]["reference"] is None


def test_comment_anchor_text_of_several_paragraphs_and_of_no_range(tmp_path):
    src, cid, _ = commented_thread(tmp_path / "c.docx")
    doc = rdocx.Document.open(src)
    doc.move_comment(cid, rdocx.RunRange(start=rdocx.RunPosition(body_index=0, run_index=1),
                                         end=rdocx.RunPosition(body_index=1, run_index=1)))
    assert doc.comments[0].anchor_text == "paragraph with some words here.\nBeta paragraph."
    bare = rewrite_body(src, tmp_path / "ref.docx", lambda x: re.sub(r'<w:commentRange(Start|End) w:id="0"/>', "", x))
    assert rdocx.Document.open(bare).comments[0].anchor_text == ""          # a reference, no range


def thread_state(doc):
    return [(c.id, c.parent_id, c.author, c.text, c.date, c.resolved, c.anchor_text) for c in doc.comments]


def test_move_comment_keeps_its_thread(tmp_path):
    src, cid, rid = commented_thread(tmp_path / "c.docx")
    doc = rdocx.Document.open(src)
    before = thread_state(doc)
    doc.move_comment_to_text(cid, "Beta")
    assert thread_state(doc) == [before[0][:-1] + ("Beta",), before[1]]
    doc.move_comment(cid, rdocx.RunRange(start=rdocx.RunPosition(body_index=2, run_index=0),
                                         end=rdocx.RunPosition(body_index=2, run_index=1)))
    assert doc.comments[0].anchor_text == "Gamma paragraph." and doc.comments[0].resolved
    for bad in (lambda: doc.move_comment_to_text(rid, "Beta"),       # a reply moves with its root
                lambda: doc.move_comment_to_text(99, "Beta"),        # unknown id
                lambda: doc.move_comment_to_text(cid, "Nowhere")):   # text not found
        with pytest.raises(rdocx.RdocxError):
            bad()
    doc.save(tmp_path / "moved.docx")
    xml = part(tmp_path / "moved.docx", "word/document.xml").decode()
    assert [xml.count(f'<w:{m} w:id="{cid}"/>') for m in ("commentRangeStart", "commentRangeEnd", "commentReference")] == [1, 1, 1]
    again = rdocx.Document.open(tmp_path / "moved.docx")
    assert thread_state(again) == thread_state(doc)
    anchor = rdocx.Document.open(src).comments[0].anchor                     # an anchor read back moves a comment
    again.move_comment(cid, anchor)
    assert again.comments[0].anchor_text == "paragraph"


def test_move_comment_out_of_google_wrappers_cli(rdocx_cli, tmp_path):
    src, cid, rid = commented_thread(tmp_path / "c.docx")
    g = google_wrapped(src, tmp_path / "g.docx")
    out = tmp_path / "moved.docx"
    rec = json.loads(run([rdocx_cli, "comment", "move", g, "--id", str(cid), "--anchor", "paragraph", "--occurrence", "1",
                          "-o", out, "--json"], check=True).stdout)
    assert rec["comment_id"] == cid and rec["action"] == "move"
    doc = rdocx.Document.open(out)
    assert [(c.id, c.parent_id, c.resolved, c.anchor_text) for c in doc.comments] == [
        (cid, None, True, "paragraph"), (rid, cid, False, None)]
    assert doc.comments[0].anchor.start.item.text == "Beta paragraph."
    xml = part(out, "word/document.xml").decode()
    assert "goog_rdk_0" not in xml and "goog_rdk_1" in xml               # the emptied wrapper went, the other stays
    assert run([rdocx_cli, "validate", out]).returncode == 0
    refused = run([rdocx_cli, "comment", "move", g, "--id", str(cid), "--anchor", "paragraph", "-o", out])
    assert refused.returncode == 1 and "already exists" in refused.stderr
    for bad in (["--id", str(cid), "--anchor", "Nowhere"], ["--id", str(rid), "--anchor", "Beta"]):
        res = run([rdocx_cli, "comment", "move", g, *bad, "-o", tmp_path / "bad.docx"])
        assert res.returncode == 1 and not (tmp_path / "bad.docx").exists()


def test_removing_the_commented_content_removes_its_thread(rdocx_cli, tmp_path):
    src, cid, _ = commented_thread(tmp_path / "c.docx")
    doc = rdocx.Document.open(src)
    assert doc.remove_content(0)
    assert len(doc.comments) == 0
    doc.save(tmp_path / "removed.docx")
    xml = part(tmp_path / "removed.docx", "word/document.xml").decode()
    assert "commentRangeStart" not in xml and "commentReference" not in xml
    assert run([rdocx_cli, "validate", tmp_path / "removed.docx"]).returncode == 0
    d = docx.Document()
    table = d.add_table(rows=2, cols=1)
    table.cell(0, 0).text, table.cell(1, 0).text = "kept row", "removed row"
    d.save(tmp_path / "t.docx")
    doc = rdocx.Document.open(tmp_path / "t.docx")
    doc.add_comment_on_text("removed", author="Reviewer", text="Drop this row?", date=STAMP)
    doc.tables[0].remove_row(1)
    assert len(doc.comments) == 0


def test_a_partial_cut_keeps_the_comment_on_what_is_left(tmp_path):
    src, cid, rid = commented_thread(tmp_path / "c.docx")
    doc = rdocx.Document.open(src)
    doc.move_comment(cid, rdocx.RunRange(start=rdocx.RunPosition(body_index=0, run_index=1),
                                         end=rdocx.RunPosition(body_index=1, run_index=1)))
    doc.remove_content(0)
    assert [(c.id, c.parent_id, c.anchor_text) for c in doc.comments] == [(cid, None, "Beta paragraph."), (rid, cid, None)]


def test_a_popped_paragraph_carries_its_thread_back(tmp_path):
    src, _, _ = commented_thread(tmp_path / "c.docx")
    doc = rdocx.Document.open(src)
    fragment = doc.pop_content(0)
    assert len(doc.comments) == 0
    doc.insert_content(2, fragment)
    assert [p.text for p in doc.paragraphs][-1].startswith("Alpha")
    root, reply = doc.comments
    assert (root.text, root.anchor_text, root.resolved, reply.text, reply.parent_id) == (
        "Check.", "paragraph", True, "Done.", root.id)


def test_validate_flags_a_comment_without_range_or_reference(rdocx_cli, tmp_path):
    src, cid, _ = commented_thread(tmp_path / "c.docx")
    bare = rewrite_body(src, tmp_path / "bare.docx", lambda x: re.sub(
        r'<w:commentRange(Start|End) w:id="0"/>|<w:r>(?:(?!</w:r>).)*?<w:commentReference w:id="0"/></w:r>', "", x,
        flags=re.S))
    assert b"commentReference" not in part(bare, "word/document.xml")
    assert rdocx.Document.open(bare).comments[0].anchor is None and rdocx.Document.open(bare).comments[0].anchor_text is None
    res = run([rdocx_cli, "validate", bare])
    assert res.returncode == 1 and f"comment {cid} has no range and no reference in any story" in res.stdout
    assert run([rdocx_cli, "validate", src]).returncode == 0


# ---------------------------------------------------------------- tracked changes
def test_compare_then_accept_or_reject(rdocx_cli, tmp_path):
    a = two_paragraphs(tmp_path / "a.docx")
    b = two_paragraphs(tmp_path / "b.docx", first="Alpha paragraph with other words here.")
    red = tmp_path / "red.docx"
    assert compare(rdocx_cli, a, b, red).returncode == 0
    doc = rdocx.Document.open(red)
    kinds = sorted({r.kind for r in doc.revisions})
    assert "insertion" in kinds and "deletion" in kinds
    acc = rdocx.Document.open(red)
    acc.accept_all()
    rej = rdocx.Document.open(red)
    rej.reject_all()
    assert acc.paragraphs[0].text == "Alpha paragraph with other words here."
    assert rej.paragraphs[0].text == "Alpha paragraph with some words here."


def test_revision_cli_list_and_accept_by_author(rdocx_cli, tmp_path):
    a = two_paragraphs(tmp_path / "a.docx")
    b = two_paragraphs(tmp_path / "b.docx", second="Beta paragraph, edited.")
    red, acc = tmp_path / "red.docx", tmp_path / "acc.docx"
    compare(rdocx_cli, a, b, red)
    listed = json.loads(run([rdocx_cli, "revision", "list", "--json", red], check=True).stdout)["revisions"]
    assert listed and all(r["author"] == "Reviewer" for r in listed)
    run([rdocx_cli, "revision", "accept", red, "--author", "Reviewer", "-o", acc], check=True)
    assert json.loads(run([rdocx_cli, "revision", "list", "--json", acc], check=True).stdout)["revisions"] == []


def test_accept_revisions_in_a_date_range(rdocx_cli, tmp_path):
    a = two_paragraphs(tmp_path / "a.docx")
    b = two_paragraphs(tmp_path / "b.docx", first="Alpha paragraph with other words here.")
    compare(rdocx_cli, a, b, tmp_path / "red.docx")
    doc = rdocx.Document.open(tmp_path / "red.docx")
    assert doc.accept_revisions_in_date_range(start="2026-09-27T00:00:00Z", end="2026-09-28T00:00:00Z") >= 2
    assert doc.paragraphs[0].text == "Alpha paragraph with other words here." and len(doc.revisions) == 0


def test_python_compare_method(tmp_path):
    a = rdocx.Document.open(two_paragraphs(tmp_path / "a.docx"))
    b = rdocx.Document.open(two_paragraphs(tmp_path / "b.docx", second="Gamma paragraph."))
    a.compare(b, "Reviewer", STAMP)
    assert len(a.revisions) >= 2


def body_doc(spec):
    """A new document from a spec: "T" is a one-cell table, any other string a paragraph."""
    d = rdocx.Document()
    for item in spec:
        if item == "T":
            d.add_table(1, 1)
            d.tables[len(d.tables) - 1].cell(0, 0).text = "cell"
        else:
            d.add_paragraph(item)
    return d


def body_shape(d):
    return ([p.text for p in d.paragraphs], len(d.tables))


@pytest.mark.parametrize("original, edited", [
    (["a"], ["a", "T", ""]),
    ([""], ["", "T", ""]),
    (["a", "T", ""], ["a", "T", "", "T", ""]),
    (["a"], ["a", "T", "b"]),
    (["a", "T", "b"], ["a"]),
])
def test_compare_a_table_added_or_removed_at_the_end_of_the_body(original, edited):
    """A table followed by a paragraph added (or removed) at the very end of the body compares, and accepting or
    rejecting the redline gives back the edited or the original body."""
    red = body_doc(original)
    red.compare(body_doc(edited), "Reviewer", STAMP, granularity="word")
    accepted, rejected = rdocx.Document.from_bytes(red.to_bytes()), rdocx.Document.from_bytes(red.to_bytes())
    accepted.accept_all()
    rejected.reject_all()
    assert body_shape(accepted) == body_shape(body_doc(edited))
    assert body_shape(rejected) == body_shape(body_doc(original))


def test_compare_a_table_inserted_between_two_paragraphs():
    red = body_doc(["a", "b"])
    red.compare(body_doc(["a", "T", "b"]), "Reviewer", STAMP, granularity="word")
    red.accept_all()
    assert body_shape(red) == (["a", "b"], 1)


def solid_png(rgb):
    buf = io.BytesIO()
    Image.new("RGB", (60, 40), rgb).save(buf, "PNG")
    return buf.getvalue()


def figure_doc(image, caption="Figure 1. Caption."):
    d = rdocx.Document()
    d.add_paragraph("Before the figure.")
    d.add_picture(image, "figure1.png", rdocx.Inches(1), rdocx.Inches(0.66))
    d.add_paragraph(caption)
    return d


def figure_images(d, tmp_path):
    """The bytes of each picture the body shows, in order (python-docx reads the saved file)."""
    d.save(tmp_path / "f.docx")
    doc = docx.Document(tmp_path / "f.docx")
    rids = [b.get(qn("r:embed")) for b in doc.element.body.iter(qn("a:blip"))]
    return [doc.part.related_parts[rid].blob for rid in rids]


@pytest.mark.parametrize("caption", ["Figure 1. Caption.", "Figure 1. New caption."])
def test_compare_tracks_a_picture_whose_image_changed(caption, tmp_path):
    """A picture whose image changed is marked deleted and inserted, as Word's Compare does: accepting gives the
    new image, rejecting the old one, with or without another change next to it."""
    old, new = solid_png((200, 30, 30)), solid_png((30, 30, 200))
    red = figure_doc(old)
    red.compare(figure_doc(new, caption), "Reviewer", STAMP, granularity="word")
    assert {"deletion", "insertion"} <= {r.kind for r in red.revisions}
    accepted, rejected = rdocx.Document.from_bytes(red.to_bytes()), rdocx.Document.from_bytes(red.to_bytes())
    accepted.accept_all()
    rejected.reject_all()
    assert figure_images(accepted, tmp_path) == [new] and accepted.paragraphs[-1].text == caption
    assert figure_images(rejected, tmp_path) == [old] and rejected.paragraphs[-1].text == "Figure 1. Caption."


def test_compare_marks_only_the_changed_word(rdocx_cli, tmp_path):
    text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor."
    a = two_paragraphs(tmp_path / "a.docx", first=text)
    b = two_paragraphs(tmp_path / "b.docx", first=text.replace("dolor", "DOLOR"))
    compare(rdocx_cli, a, b, tmp_path / "red.docx", "--granularity", "word")
    deleted = "".join(re.findall(r"<w:delText[^>]*>([^<]*)</w:delText>", part(tmp_path / "red.docx", "word/document.xml").decode()))
    assert deleted.strip() == "dolor"


def test_python_compare_marks_only_the_changed_word(tmp_path):
    text = "Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor."
    doc = rdocx.Document.open(two_paragraphs(tmp_path / "a.docx", first=text))
    edited = rdocx.Document.open(two_paragraphs(tmp_path / "b.docx", first=text.replace("dolor", "DOLOR")))
    doc.compare(edited, "Reviewer", "2026-09-29T12:00:00Z", granularity="word")
    doc.save(tmp_path / "red.docx")
    deleted = "".join(re.findall(r"<w:delText[^>]*>([^<]*)</w:delText>", part(tmp_path / "red.docx", "word/document.xml").decode()))
    assert deleted.strip() == "dolor"


def test_python_compare_ignores_comments_on_request(tmp_path):
    doc = rdocx.Document.open(two_paragraphs(tmp_path / "a.docx"))
    edited = rdocx.Document.open(two_paragraphs(tmp_path / "b.docx", first="Alpha paragraph with other words here."))
    edited.add_comment(rdocx.RunRange(start=rdocx.RunPosition(body_index=1, run_index=0),
                                      end=rdocx.RunPosition(body_index=1, run_index=1)), author="R", text="New comment")
    assert doc.compare(edited, "Reviewer", "2026-09-29T12:00:00Z", ignore_comments=True) == ()
    assert doc.comments == () and [r.kind for r in doc.revisions] == ["deletion", "insertion"]


def test_compare_accepts_a_pair_whose_comments_differ(rdocx_cli, tmp_path):
    a = two_paragraphs(tmp_path / "a.docx")
    b = tmp_path / "b.docx"
    run([rdocx_cli, "comment", "add", a, "--start-paragraph", "1", "--start-run", "0", "--end-paragraph", "1",
         "--end-run", "1", "--author", "R", "--text", "New comment", "-o", b], check=True)
    # The comment added on the edited side replaces its paragraph in the redline (deleted and reinserted with
    # the anchor): accepting gives the edited comments, rejecting the original ones, and the text is unchanged.
    assert compare(rdocx_cli, a, b, tmp_path / "red.docx").returncode == 0
    red = rdocx.Document.open(tmp_path / "red.docx")
    assert [c.text for c in red.comments] == ["New comment"]
    assert {r.kind for r in red.revisions} == {"deletion", "insertion"}
    texts = [p.text for p in rdocx.Document.open(a).paragraphs]
    for resolve, comments in (("accept_all", ["New comment"]), ("reject_all", [])):
        doc = rdocx.Document.open(tmp_path / "red.docx")
        getattr(doc, resolve)()
        doc = rdocx.Document.from_bytes(doc.to_bytes())
        assert [c.text for c in doc.comments] == comments
        assert [p.text for p in doc.paragraphs] == texts and doc.revisions == ()
    assert compare(rdocx_cli, a, b, tmp_path / "red.docx", "--ignore-comments").returncode == 0
    red = rdocx.Document.open(tmp_path / "red.docx")
    assert red.comments == () and red.revisions == ()


def test_compare_ignores_a_content_control_id(rdocx_cli, report_docx, tmp_path):
    b = tmp_path / "b.docx"
    with zipfile.ZipFile(report_docx) as zin, zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "word/document.xml":
                data = data.replace(b'w:val="1209466531"', b'w:val="1209466532"')
            zout.writestr(info, data)
    assert compare(rdocx_cli, report_docx, b, tmp_path / "red.docx").returncode == 0


def test_compare_against_its_own_rdocx_save(rdocx_cli, report_docx, tmp_path):
    doc = rdocx.Document.open(report_docx)
    doc.try_replace_text("described", "outlined")
    doc.save(tmp_path / "e.docx")
    assert compare(rdocx_cli, report_docx, tmp_path / "e.docx", tmp_path / "red.docx").returncode == 0


def test_compare_after_rdocx_edit_reports_only_the_edit(rdocx_cli, tmp_path):
    d = docx.Document()
    for t in ("Paragraph 1, lorem ipsum.", "WORD", "Paragraph 3, lorem ipsum."):
        d.add_paragraph(t)
    sect = d.sections[0]._sectPr
    pgsz = sect.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pgSz")
    pgsz.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}orient", "portrait")
    d.save(tmp_path / "a.docx")
    doc = rdocx.Document.open(tmp_path / "a.docx")
    doc.try_replace_text("WORD", "TERM")
    doc.save(tmp_path / "b.docx")
    compare(rdocx_cli, tmp_path / "a.docx", tmp_path / "b.docx", tmp_path / "red.docx")
    kinds = [r.kind for r in rdocx.Document.open(tmp_path / "red.docx").revisions]
    assert "section_property_change" not in kinds and len(kinds) == 2


def test_compare_after_a_plain_edit_of_an_rdocx_saved_original(rdocx_cli, report_docx, tmp_path):
    """The base of the two gap tests below: without the field refresh or the TOC rebuild, the pair compares."""
    base = noop_saved(report_docx, tmp_path / "base.docx")
    doc = rdocx.Document.open(base)
    doc.try_replace_text("described", "outlined")
    doc.save(tmp_path / "e.docx")
    assert compare(rdocx_cli, base, tmp_path / "e.docx", tmp_path / "red.docx").returncode == 0


def test_compare_after_refreshing_packed_page_fields(rdocx_cli, report_docx, tmp_path):
    base = noop_saved(report_docx, tmp_path / "base.docx")
    doc = rdocx.Document.open(base)
    doc.update_page_fields()
    doc.save(tmp_path / "f.docx")
    res = compare(rdocx_cli, base, tmp_path / "f.docx", tmp_path / "red.docx")
    assert res.returncode == 0, res.stderr


def test_compare_after_toc_rebuild(rdocx_cli, report_docx, tmp_path):
    """The redline also holds the rebuilt TOC; accepting it gives the edited text, rejecting it the original's."""
    base = noop_saved(report_docx, tmp_path / "base.docx")
    doc = rdocx.Document.open(base)
    doc.try_replace_text("described", "outlined")  # an edit rewrites the identity attributes, then the TOC rebuilds
    doc.rebuild_toc()
    doc.save(tmp_path / "t.docx")
    res = compare(rdocx_cli, base, tmp_path / "t.docx", tmp_path / "red.docx")
    assert res.returncode == 0, res.stderr
    texts = {name: [p.text for p in rdocx.Document.open(path).paragraphs]
             for name, path in (("base", base), ("edited", tmp_path / "t.docx"))}
    red = rdocx.Document.open(tmp_path / "red.docx")
    assert len(red.revisions) > 2  # the edit (a deletion and an insertion) and the TOC entries
    red.accept_all()
    assert [p.text for p in red.paragraphs] == texts["edited"]
    red = rdocx.Document.open(tmp_path / "red.docx")
    red.reject_all()
    assert [p.text for p in red.paragraphs] == texts["base"]


def test_revision_accept_record_counts_every_story(rdocx_cli, tmp_path):
    """The workaround for revisions-main-story: `revision accept --json` counts revisions in every story."""
    d = docx.Document()
    d.add_paragraph("Body.")
    d.sections[0].footer.paragraphs[0].text = "Footer lorem ipsum"
    d.save(tmp_path / "a.docx")
    d = docx.Document(tmp_path / "a.docx")
    d.sections[0].footer.paragraphs[0].text = "Footer lorem IPSUM"
    d.save(tmp_path / "b.docx")
    compare(rdocx_cli, tmp_path / "a.docx", tmp_path / "b.docx", tmp_path / "red.docx")
    rec = json.loads(run([rdocx_cli, "revision", "accept", tmp_path / "red.docx", "-o", tmp_path / "acc.docx", "--json"],
                         check=True).stdout)
    assert rec["resolved"] >= 1 and rec["scope"] == "all-supported-stories"


def test_revisions_listed_across_stories(rdocx_cli, tmp_path):
    d = docx.Document()
    d.add_paragraph("Body.")
    d.sections[0].footer.paragraphs[0].text = "Footer lorem ipsum"
    d.save(tmp_path / "a.docx")
    d = docx.Document(tmp_path / "a.docx")
    d.sections[0].footer.paragraphs[0].text = "Footer lorem IPSUM"
    d.save(tmp_path / "b.docx")
    compare(rdocx_cli, tmp_path / "a.docx", tmp_path / "b.docx", tmp_path / "red.docx")
    assert b"<w:ins " in part(tmp_path / "red.docx", "word/footer1.xml")
    assert {r.story.kind for r in rdocx.Document.open(tmp_path / "red.docx").revisions} == {"footer"}
    listed = json.loads(run([rdocx_cli, "revision", "list", "--json", tmp_path / "red.docx"], check=True).stdout)
    assert {r["story"]["kind"] for r in listed["revisions"]} == {"footer"}
