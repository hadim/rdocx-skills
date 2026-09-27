"""Fixture builders of the two acceptance matrices (identity attributes, producer traits).

`build(path, where, attr, word)` writes a python-docx document with what a real report has (a cached TOC field,
headings, a table, a block content control, PAGE / NUMPAGES in the footer) and, optionally, one identity
attribute on every element of one kind. `rewrite(path, fn_doc, fn_footer, empty_comments)` then applies one
producer trait to the whole file. Both are exercised by test_docx_matrices.py."""
import re
import zipfile

from docx import Document
from docx.oxml.ns import qn

XS = "{http://www.w3.org/XML/1998/namespace}space"
W14 = "{http://schemas.microsoft.com/office/word/2010/wordml}"
T = "2026-09-27T12:00:00Z"


def field(p, instr, cached=None):
    out = []
    for kind, text in (("begin", None), ("instr", instr), ("separate", None)):
        r = p.makeelement(qn("w:r"), {})
        if kind == "instr":
            el = r.makeelement(qn("w:instrText"), {}); el.text = text; el.set(XS, "preserve")
        else:
            el = r.makeelement(qn("w:fldChar"), {}); el.set(qn("w:fldCharType"), kind)
        r.append(el); p.append(r); out.append(r)
    if cached is not None:
        r = p.makeelement(qn("w:r"), {}); t = r.makeelement(qn("w:t"), {}); t.text = cached; r.append(t); p.append(r)
        r = p.makeelement(qn("w:r"), {}); el = r.makeelement(qn("w:fldChar"), {}); el.set(qn("w:fldCharType"), "end")
        r.append(el); p.append(r); out.append(r)
    return out


def build(path, where=None, attr=None, word="alpha"):
    doc = Document()
    body = doc.element.body
    field_runs, entry_runs = [], []
    entries = ["Chapter 1", "Section 1.1", "Chapter 2", "Section 2.1", "Chapter 3", "Section 3.1"]
    for k, title in enumerate(entries):
        p = doc.add_paragraph()
        if k == 0:
            field_runs += field(p._p, ' TOC \\o "1-3" \\h \\z \\u ')
        entry_runs.append(p.add_run(title + "\t1")._r)
        if k == len(entries) - 1:
            r = p._p.makeelement(qn("w:r"), {}); el = r.makeelement(qn("w:fldChar"), {})
            el.set(qn("w:fldCharType"), "end"); r.append(el); p._p.append(r); field_runs.append(r)
    for i in range(1, 4):
        doc.add_heading("Chapter %d" % i, level=1)
        doc.add_paragraph("Body text of chapter %d, lorem %s dolor." % (i, word if i == 2 else "ipsum"))
        doc.add_heading("Section %d.1" % i, level=2)
        doc.add_paragraph("Body text of section %d.1." % i)
    t = doc.add_table(rows=2, cols=2)
    for c in t._cells:
        c.text = "cell"
    sdt = body.makeelement(qn("w:sdt"), {}); pr = sdt.makeelement(qn("w:sdtPr"), {})
    tag = pr.makeelement(qn("w:tag"), {}); tag.set(qn("w:val"), "block"); pr.append(tag); sdt.append(pr)
    content = sdt.makeelement(qn("w:sdtContent"), {})
    for s in ("Inside the content control, one.", "Inside the content control, two."):
        pp = doc.add_paragraph(s)._p; body.remove(pp); content.append(pp)
    sdt.append(content); body.find(qn("w:tbl")).addnext(sdt)
    fp = doc.sections[0].footer.paragraphs[0]._p
    r = fp.makeelement(qn("w:r"), {}); tt = r.makeelement(qn("w:t"), {}); tt.text = "Page "; tt.set(XS, "preserve"); r.append(tt); fp.append(r)
    footer_runs = field(fp, " PAGE ", "1") + field(fp, " NUMPAGES ", "1")
    targets = {
        "paragraphs": [p for p in body.iter(qn("w:p"))],
        "runs": [r for r in body.iter(qn("w:r"))],
        "field runs": field_runs,
        "footer field runs": footer_runs,
        "content control": [pr],
        "table rows": [tr for tr in body.iter(qn("w:tr"))],
    }
    if where:
        for el in targets[where]:
            if attr == "w:tag":
                el.find(qn("w:tag")).set(qn("w:val"), "block-2")
            elif attr == "w:id":
                i = el.makeelement(qn("w:id"), {}); i.set(qn("w:val"), "-2000000001"); el.insert(0, i)
            elif attr.startswith("w14:"):
                el.set(W14 + attr[4:], "1A2B3C4D")
            else:
                el.set(qn(attr), "00A1B2C3")
    doc.save(path)
    return path



EMPTY_COMMENTS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:comments '
                  'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                  'xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml"/>')


def rewrite(path, fn_doc=None, fn_footer=None, empty_comments=False):
    zin = zipfile.ZipFile(path); items = [(i, zin.read(i.filename)) for i in zin.infolist()]; zin.close()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info, data in items:
            if info.filename == "word/document.xml" and fn_doc:
                data = fn_doc(data.decode()).encode()
            if info.filename == "word/footer1.xml" and fn_footer:
                data = fn_footer(data.decode()).encode()
            if empty_comments and info.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/word/comments.xml" ContentType="application/'
                                    b'vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/></Types>')
            if empty_comments and info.filename == "word/_rels/document.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/'
                                    b'officeDocument/2006/relationships/comments" Target="comments.xml"/></Relationships>')
            z.writestr(info, data)
        if empty_comments:
            z.writestr("word/comments.xml", EMPTY_COMMENTS)
    return path


def pack(s):
    # begin + instrText + separate of each field into one run, as Google Docs exports write them
    return re.sub(r'<w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r>(<w:instrText[^>]*>[^<]*</w:instrText>)</w:r>'
                  r'<w:r><w:fldChar w:fldCharType="separate"/></w:r>',
                  r'<w:r><w:fldChar w:fldCharType="begin"/>\1<w:fldChar w:fldCharType="separate"/></w:r>', s)


def pack_no_cache(s):
    s = pack(s)
    return re.sub(r'(<w:fldChar w:fldCharType="separate"/></w:r>)<w:r><w:t>[^<]*</w:t></w:r><w:r>(<w:fldChar w:fldCharType="end"/>)',
                  r'<w:fldChar w:fldCharType="separate"/>\2', s.replace('<w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>1</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r>',
                                                                     '<w:fldChar w:fldCharType="separate"/><w:fldChar w:fldCharType="end"/></w:r>'))


def inline_sdt(s):
    # wrap the first run of every body paragraph in an inline content control
    return re.sub(r'(<w:p>(?:<w:pPr>.*?</w:pPr>)?)(<w:r>.*?</w:r>)',
                  r'\1<w:sdt><w:sdtPr><w:tag w:val="goog_rdk_0"/></w:sdtPr><w:sdtContent>\2</w:sdtContent></w:sdt>', s)


TRAITS = {
    "control: none": (None, None),
    "xml:space=preserve on every w:t": (lambda s: re.sub(r"<w:t>", '<w:t xml:space="preserve">', s), None),
    "w:val=\"0\" toggles on every run": (lambda s: s.replace("<w:r>", '<w:r><w:rPr><w:b w:val="0"/><w:i w:val="0"/></w:rPr>'), None),
    "empty <w:pPr/> on plain paragraphs": (lambda s: s.replace("<w:p><w:r>", "<w:p><w:pPr/><w:r>"), None),
    "pageBreakBefore/keepNext w:val=\"0\"": (lambda s: s.replace("<w:p><w:r>", '<w:p><w:pPr><w:keepNext w:val="0"/><w:pageBreakBefore w:val="0"/></w:pPr><w:r>'), None),
    "w:orient=\"portrait\" on pgSz": (lambda s: s.replace("<w:pgSz ", '<w:pgSz w:orient="portrait" '), None),
    "default namespace on the root": (lambda s: s.replace("<w:document ", '<w:document xmlns="http://schemas.microsoft.com/office/tasks/2019/documenttasks" ', 1), None),
    "packed TOC field run": (pack, None),
    "packed footer fields, no cached result": (None, pack_no_cache),
    "inline content control on first runs": (inline_sdt, None),
    "empty comments part": (None, None, True),
}


