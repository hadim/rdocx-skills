"""docx, editing: counted replacement, structural edits, tables, pictures, handles, save semantics."""
import io
import re

import docx
import pytest
import rdocx
from PIL import Image

from builders import word_textbox_docx
from conftest import part, run


def simple(path, *texts, table_after=None):
    d = docx.Document()
    for i, t in enumerate(texts):
        d.add_paragraph(t)
        if table_after == i:
            d.add_table(rows=2, cols=2).cell(0, 0).text = "cell"
    d.save(path)
    return path


def test_replace_text_at_a_text_box_edits_both_copies_or_refuses(tmp_path):
    box = rdocx.Document.open(word_textbox_docx(tmp_path / "t.docx"))
    text_box = next(i for i in box.story_items if i.story.kind == "text_box" and i.kind == "paragraph")
    try:
        box.replace_text_at(text_box, "NEEDLE", "PIN")
    except rdocx.RdocxError:
        return                                                            # refused: nothing changed
    box.save(tmp_path / "o.docx")
    xml = part(tmp_path / "o.docx", "word/document.xml").decode()
    copies = [re.search(r"<mc:%s\b.*?</mc:%s>" % (tag, tag), xml, re.S).group(0) for tag in ("Choice", "Fallback")]
    assert ["PIN" in c and "NEEDLE" not in c for c in copies] == [True, True]   # Word's copy and the VML fallback


KEEP = {
    "font": lambda d: setattr(d.paragraphs[1].runs[0].font, "bold", True),
    "paragraph_format": lambda d: setattr(d.paragraphs[1].paragraph_format, "space_after", rdocx.Pt(6)),
    "style": lambda d: setattr(d.paragraphs[1], "style", "Heading1"),
    "alignment": lambda d: setattr(d.paragraphs[1], "alignment", rdocx.WD_ALIGN_PARAGRAPH.CENTER),
    "numbering": lambda d: setattr(d.paragraphs[1], "numbering", (1, 0)),
    "cell_width": lambda d: setattr(d.tables[0].cell(0, 0), "width", rdocx.Inches(2)),
    "run_text": lambda d: setattr(d.paragraphs[1].runs[0], "text", "X"),
}
# Appends, inserts and text setters keep the paragraph, run, table, row and cell handles they do not replace; story
# items still go stale.
KEEP_NODES = {
    "add_paragraph": lambda d: d.add_paragraph("more"),
    "add_heading": lambda d: d.add_heading("Head", 1),
    "add_table": lambda d: d.add_table(1, 1),
    "add_page_break": lambda d: d.add_page_break(),
    "insert_paragraph": lambda d: d.insert_paragraph(0, "first"),
    "insert_paragraph_before": lambda d: d.paragraphs[1].insert_paragraph_before("before"),
    "add_run": lambda d: d.paragraphs[1].add_run("more"),
    "paragraph_text": lambda d: setattr(d.paragraphs[1], "text", "Z"),
    "cell_text": lambda d: setattr(d.tables[0].cell(1, 1), "text", "Y"),
    "paragraph_replace_xml": lambda d: d.paragraphs[1].replace_xml(d.paragraphs[1].xml),
    "table_add_row": lambda d: d.tables[0].add_row(),
    "table_add_column": lambda d: d.tables[0].add_column(rdocx.Inches(1)),
}
# Row and grid edits keep the table handle and retire its rows and cells.
TABLE_EDITS = {
    "clone_row": lambda d: d.tables[0].clone_row(0),
    "remove_row": lambda d: d.tables[0].remove_row(1),
    "insert_column": lambda d: d.tables[0].insert_column(0),
    "remove_column": lambda d: d.tables[0].remove_column(1),
    "set_cell_grid_span": lambda d: d.tables[0].set_cell_grid_span(1, 0, 2),
}
INVALIDATE = {
    "try_replace_text": lambda d: d.try_replace_text("delta", "D"),
    "replace_all_regex": lambda d: d.replace_all_regex([("delta", "D")]),
    "paragraph_replace_text": lambda d: d.paragraphs[1].replace_text("delta", "D"),
    "cell_replace_text": lambda d: d.tables[0].cell(0, 0).replace_text("cell", "C"),
    "replace_text_at": lambda d: d.replace_text_at(d.story_items[1], "delta", "D"),
    "add_comment": lambda d: d.add_comment(rdocx.RunRange(start=rdocx.RunPosition(body_index=0, run_index=0),
                                                          end=rdocx.RunPosition(body_index=0, run_index=1)), author="A", text="c"),
    "remove_content": lambda d: d.remove_content(2),
    "split_run": lambda d: d.split_run(1, 0, 2),
    "set_story_text": lambda d: d.set_story_text(d.story_items[1], "Q"),
    "pop_content": lambda d: d.pop_content(1),
    "move_content": lambda d: d.move_content(d.paragraphs[1], 0),
    "clone_content": lambda d: d.clone_content(d.paragraphs[1], 0),
    "add_picture": lambda d: d.add_picture(png(), "red.png", rdocx.Inches(1), rdocx.Inches(0.5)),
}


@pytest.mark.parametrize("op", sorted(KEEP) + sorted(KEEP_NODES) + sorted(TABLE_EDITS) + sorted(INVALIDATE))
def test_which_calls_invalidate_handles(tmp_path, op):
    """The rule in python-api.md, "Handles": formatting setters and Run.text keep every handle valid; appends,
    inserts and text setters keep the other paragraph, run, table and cell handles; row edits retire the rows and
    cells of that table; the other edits invalidate every handle (paragraph, run, table, cell, story item)."""
    path = simple(tmp_path / "a.docx", "Alpha beta", "Gamma delta", table_after=1)
    doc = rdocx.Document.open(path)
    doc.tables[0].cell(1, 1).text = ""                 # an empty cell for set_cell_grid_span to consume
    nodes = [lambda p=doc.paragraphs[0]: p.text, lambda r=doc.paragraphs[0].runs[0]: r.text,
             lambda t=doc.tables[0]: len(t.rows)]
    cell = [lambda c=doc.tables[0].cell(0, 0): c.text]
    item = [lambda i=doc.story_items[0]: doc.set_story_text(i, "Alpha beta")]
    edited = [lambda r=doc.paragraphs[1].runs[0]: r.text] if op in ("paragraph_text", "paragraph_replace_xml") else []
    edited += [lambda p=doc.tables[0].cell(1, 1).paragraphs[0]: p.text] if op == "cell_text" else []
    (KEEP.get(op) or KEEP_NODES.get(op) or TABLE_EDITS.get(op) or INVALIDATE[op])(doc)
    valid = (nodes + cell + item if op in KEEP else nodes + cell if op in KEEP_NODES
             else nodes if op in TABLE_EDITS else [])
    for h in nodes + cell + item + edited:
        if h in valid:
            h()
        else:
            with pytest.raises(rdocx.StaleElementError):
                h()
    if op == "insert_paragraph":
        assert nodes[0]() == "Alpha beta"                 # a held handle follows its paragraph, not its index


def test_a_stale_handle_names_the_call_that_retired_it(tmp_path):
    doc = rdocx.Document.open(simple(tmp_path / "a.docx", "Alpha", "Beta"))
    held = doc.paragraphs[0]
    doc.remove_content(1)
    with pytest.raises(rdocx.StaleElementError, match="remove_content"):
        held.text


def test_table_merge_borders_shading_widths_and_row_height_from_python(tmp_path):
    doc = rdocx.Document.open(simple(tmp_path / "a.docx", "x", table_after=0))
    doc.tables[0].cell(1, 1).text = "full"
    with pytest.raises(rdocx.RdocxError):
        doc.tables[0].set_cell_grid_span(1, 0, 2)  # a span consumes empty cells only
    t = doc.tables[0]
    held = t.cell(1, 0)
    t.set_cell_grid_span(0, 0, 2)  # consumes the empty cell to its right; retires the table's rows and cells
    with pytest.raises(rdocx.StaleElementError):
        held.text
    assert t.cell(0, 0).grid_span == 2                 # the table handle stays valid
    t.cell(1, 1).shading = "FFEE00"
    t.set_borders("single", size=4, color="000000")
    t.set_column_width(0, rdocx.Inches(1))
    t.rows[1].height = rdocx.Inches(0.5)
    doc.save(tmp_path / "b.docx")
    x = part(tmp_path / "b.docx", "word/document.xml").decode()
    assert [len(r.cells) for r in docx.Document(tmp_path / "b.docx").tables[0].rows] == [2, 2]
    assert '<w:gridSpan w:val="2"/>' in x and 'w:fill="FFEE00"' in x and "<w:tblBorders>" in x
    assert '<w:gridCol w:w="1440"/>' in x and '<w:trHeight w:val="720"' in x


def png(size=(60, 40)):
    buf = io.BytesIO()
    Image.new("RGB", size).save(buf, "PNG")
    return buf.getvalue()


# ---------------------------------------------------------------- formatting values and units
def test_font_color_is_a_colour_format(tmp_path):
    """python-docx's shape: `font.color` is a ColorFormat (`rgb`, `type`, `theme_color`), never None."""
    doc = rdocx.Document()
    doc.add_paragraph("x")
    font = doc.paragraphs[0].runs[0].font
    assert (font.color.rgb, font.color.type) == (None, None)
    font.color = rdocx.RGBColor(0x7B, 0x1E, 0x3A)                      # shortcut for font.color.rgb = ...
    rgb = doc.paragraphs[0].runs[0].font.color.rgb
    assert tuple(rgb) == (0x7B, 0x1E, 0x3A) and str(rgb) == "7B1E3A" and rgb != "7B1E3A"
    assert doc.paragraphs[0].runs[0].font.color.type == rdocx.MSO_COLOR_TYPE.RGB
    doc.paragraphs[0].runs[0].font.color.theme_color = rdocx.MSO_THEME_COLOR.ACCENT_1
    color = doc.paragraphs[0].runs[0].font.color
    assert (color.type, color.theme_color) == (rdocx.MSO_COLOR_TYPE.THEME, rdocx.MSO_THEME_COLOR.ACCENT_1)
    doc.paragraphs[0].runs[0].font.color = "auto"
    assert doc.paragraphs[0].runs[0].font.color.type == rdocx.MSO_COLOR_TYPE.AUTO
    doc.paragraphs[0].runs[0].font.color = None                        # removes the colour
    assert doc.paragraphs[0].runs[0].font.color.type is None
    with pytest.raises(TypeError, match="RGBColor or a hex string"):
        doc.paragraphs[0].runs[0].font.color = 0x7B1E3A
    with pytest.raises(ValueError, match="six hexadecimal digits"):
        doc.paragraphs[0].runs[0].font.color = "GG0000"


def test_table_and_cell_borders_margins_grid_and_rows(tmp_path):
    doc = rdocx.Document()
    doc.add_table(2, 2)
    assert doc.tables[0].border("top") is None and doc.tables[0].cell_margins is None
    doc.tables[0].set_border("top", "double", size=8, color="FF0000")
    doc.tables[0].set_cell_margins(top=rdocx.Pt(2), right=rdocx.Pt(4), bottom=rdocx.Pt(2), left=rdocx.Pt(4))
    doc.tables[0].grid_widths = [rdocx.Inches(1), rdocx.Inches(2)]
    t = doc.tables[0]
    assert t.border("top") == ("double", 8, (255, 0, 0)) and t.cell_margins == (rdocx.Pt(2), rdocx.Pt(4), rdocx.Pt(2), rdocx.Pt(4))
    assert t.grid_widths == (rdocx.Inches(1), rdocx.Inches(2))
    t.cell(0, 0).set_border("left", "single", size=4, color="#00FF00")
    doc.tables[0].cell(0, 0).set_margins(top=rdocx.Pt(1), right=0, bottom=rdocx.Pt(1), left=0)
    cell = doc.tables[0].cell(0, 0)
    assert cell.border("left") == ("single", 4, (0, 255, 0)) and cell.margins == (rdocx.Pt(1), 0, rdocx.Pt(1), 0)
    assert isinstance(cell.margins[0], rdocx.Length)
    row = doc.tables[0].rows[0]
    row.is_header, row.cant_split = True, True
    doc.tables[0].rows[1].height = rdocx.Pt(20)
    doc.tables[0].rows[1].height_rule = rdocx.WD_ROW_HEIGHT_RULE.EXACTLY
    assert (doc.tables[0].rows[0].is_header, doc.tables[0].rows[0].cant_split) == (True, True)
    doc.save(tmp_path / "t.docx")
    xml = part(tmp_path / "t.docx", "word/document.xml").decode()
    assert "<w:tblHeader/>" in xml and "<w:cantSplit/>" in xml and 'w:hRule="exact"' in xml
    assert re.search(r'<w:top w:val="double" w:sz="8"[^>]*w:color="FF0000"', xml)


# ---------------------------------------------------------------- gaps reported 2026-10-09, closed by the 2026-10-10 pin
def test_validate_does_not_count_a_picture_paragraph_as_empty(rdocx_cli, tmp_path):
    doc = rdocx.Document()
    doc.add_paragraph("Before")
    for _ in range(3):
        doc.add_picture(png(), "red.png", rdocx.Inches(1), rdocx.Inches(0.5))
    doc.add_paragraph("After")
    doc.save(tmp_path / "p.docx")
    assert len(rdocx.Document.open(tmp_path / "p.docx").paragraphs) == 5
    res = run([rdocx_cli, "validate", tmp_path / "p.docx"])
    assert res.returncode == 0 and "empty paragraph" not in res.stdout, res.stdout


def test_every_colour_argument_takes_an_rgbcolor_and_a_hex_string(tmp_path):
    doc = rdocx.Document()
    doc.add_paragraph("text")
    table = doc.add_table(1, 1)
    red = rdocx.RGBColor(0xFF, 0x00, 0x00)
    for colour in (red, "FF0000", "#FF0000", (255, 0, 0)):
        doc.set_style("Normal", color=colour)
        doc.paragraphs[0].runs[0].font.color = colour
        doc.paragraphs[0].runs[0].font.shading = colour
        doc.tables[0].set_borders("single", size=4, color=colour)
        doc.tables[0].set_border("top", "single", size=4, color=colour)
        doc.tables[0].cell(0, 0).set_border("left", "single", size=4, color=colour)
        doc.tables[0].cell(0, 0).shading = colour
        assert str(doc.paragraphs[0].runs[0].font.color.rgb) == "FF0000"
        assert doc.paragraphs[0].runs[0].font.shading == doc.tables[0].cell(0, 0).shading == red
    doc.paragraphs[0].runs[0].font.color.rgb = rdocx.RGBColor(0x00, 0x00, 0xFF)   # python-docx
    assert str(doc.paragraphs[0].runs[0].font.color.rgb) == "0000FF"


def test_table_indent_negative_included(tmp_path):
    doc = rdocx.Document()
    doc.add_table(1, 2)
    doc.tables[0].indent = rdocx.Inches(0.25)
    assert doc.tables[0].indent == rdocx.Inches(0.25)
    doc.tables[0].indent = rdocx.Inches(-0.1)
    doc.save(tmp_path / "t.docx")
    assert re.search(rb'<w:tblInd w:w="-144" w:type="dxa"/>', part(tmp_path / "t.docx", "word/document.xml"))
    doc.tables[0].indent = None
    assert doc.tables[0].indent is None
