"""docx: every public mutating call of rdocx costs about the same in a document of N paragraphs and in one of 4N, as
in python-docx, where an edit touches its element and, for a picture or a note, its own part. A call whose cost grows
with the document is a strict xfail (tests/gaps.py, docx-edit-cost-grows-with-body), so a fix of some calls of the
family shows up as XPASSes at the next build. `test_every_mutating_call_is_timed` keeps the list complete as the API
grows; the reads tensorbee/rdocx#407 names are timed the same way (READS_TIMED).

Each paragraph of the two documents holds 6 runs with the run properties Google Docs writes on every run; every 10th
paragraph also holds a picture and is followed by a 2x2 table, and the documents have two sections, the first with a
header: the body's XML, its tables and the bytes of the package grow fourfold. Making a target costs about as much as
a timed call, so the documents end with targets made in advance (builders.edit_cost_docx), and each timed call opens
its documents again, so that no call sees the targets another one changed."""
import ast
import datetime
import io
import itertools
import re
import time
from pathlib import Path

import docx
import pytest
import rdocx
from PIL import ImageFont
from rdocx import (WD_ALIGN_PARAGRAPH, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT,
                   XL_CHART_TYPE, CategoryChartData, Inches, Pt, RGBColor)

from builders import edit_cost_docx, noise_png

N = 25
CALLS = 3  # fresh targets a call is timed on in each document, the fastest counting; over 50 ms, two are enough
# A cost independent of the document measures about 1 time, one that grows with it 2 to 4 times. Between the two
# bounds, the result stays: a call of GROWING counts as fixed under FIXED times, any other one as grown over RATIO
# times, so that the noise of a loaded machine does not turn a call near 2 times into an XPASS or a failure.
RATIO, FIXED, FLOOR_MS = 2.5, 1.5, 5
GROWS = "docx-edit-cost-grows-with-body"
STAMP = "2026-10-01T12:00:00Z"
LICENSE = rdocx.FontEmbeddingLicense(authorized=True, identity="CC0-1.0: Aileron")
FONT = ImageFont.load_default(12).font_bytes
KINDS = ["regular", "bold", "italic"]
RGB = RGBColor(0x7B, 0x1E, 0x3A)
URL = "https://example.com/a"
IDS = itertools.count(1)


# ---------------------------------------------------------------- targets: each gives the k-th fresh object
def document(doc, k):
    return doc


def last(items, pick):
    """The last of `items` that pick() accepts: the targets sit at the end of the body."""
    return next(items[i] for i in range(len(items) - 1, -1, -1) if pick(items[i]))


def para(doc, k):
    """The paragraph "alpha beta k", of the runs "alpha " and "beta k"."""
    return last(doc.paragraphs, lambda p: p.text == f"alpha beta {k}")


def run(doc, k):
    return para(doc, k).runs[0]


def table(doc, k):
    """A 2x2 table with "alpha k" in its first cell."""
    return last(doc.tables, lambda t: t.cell(0, 0).text == f"alpha {k}")


def cell(doc, k):
    return table(doc, k).cell(0, 0)


def row(doc, k):
    return table(doc, k).rows[0]


def picture(doc, k):
    """The picture of the k-th of the paragraphs that hold one, after the "alpha beta" ones."""
    return doc.pictures()[k - CALLS]


def section(doc, k):
    return doc.section(0)


def footnote(doc, k):
    return doc.add_footnote(para(doc, k), f"note {k}")


def endnote(doc, k):
    return doc.add_endnote(para(doc, k), f"note {k}")


def tab_stops(doc, k):
    stops = para(doc, k).paragraph_format.tab_stops
    stops.add_tab_stop(Inches(1))
    return stops


def bordered(doc, k):
    fmt = para(doc, k).paragraph_format
    fmt.set_border("bottom", "single", size=8, color="000000")
    return fmt


def merged(doc, k):
    t = table(doc, k)
    return t.cell(0, 0).merge(t.cell(0, 1))


def index_of(doc, k):
    return doc.find_content_index(para(doc, k))


def item_of(doc, p):
    """The body StoryItem of paragraph p."""
    bi = doc.find_content_index(p)
    return next(it for it in doc.story_items if it.story.kind == "body" and it.kind == "paragraph"
                and it.direct_body_index == bi)


def location_of(doc, p):
    it = item_of(doc, p)
    return rdocx.ContentLocation(it.story, it.kind, it.index_path)


def location(doc, k):
    return location_of(doc, para(doc, k))


def body_story(doc):
    return next(s for s in doc.stories if s.kind == "body")


def run_range(doc, k):
    bi = index_of(doc, k)
    return rdocx.RunRange(start=rdocx.RunPosition(body_index=bi, run_index=0),
                          end=rdocx.RunPosition(body_index=bi, run_index=1))


def story_range(doc, k):
    it = item_of(doc, para(doc, k))
    return rdocx.StoryRunRange(start=rdocx.StoryRunPosition(item=it, run_index=0),
                               end=rdocx.StoryRunPosition(item=it, run_index=1))


def comment(doc, k):
    return next(c.id for c in doc.comments if c.text == f"comment {k}")


def found(doc, pick):
    it = next(it for it in doc.story_items if pick(it))
    return rdocx.ContentLocation(it.story, it.kind, it.index_path)


def tagged(tag):
    return lambda it: it.kind == "content_control" and f'<w:tag w:val="{tag}"/>'.encode() in it.xml


def control(doc, k):
    """The location of the block content control probe_k of the document (edit_cost_docx)."""
    return found(doc, tagged(f"probe{k}"))


def controls(doc, k):
    """(the location of the content control probe_k, the location of the paragraph after it)."""
    return control(doc, k), found(doc, lambda it: it.kind == "paragraph" and it.text == f"spot {k}")


def hyperlink(doc, k):
    text = f"link {next(IDS)}"
    para(doc, k).add_hyperlink(text, URL)
    return next(h for h in doc.hyperlinks if h.text == text)


def revision(doc, k):
    """The id of the tracked insertion in the paragraph "tracked k"."""
    return 901 + k


def style(doc, k, kind="paragraph"):
    return doc.add_style(f"Probe style {next(IDS)}", kind).style_id


def numbering_definition(doc, k):
    return doc.add_numbering_definition([rdocx.ListLevel(format="decimal", text="%1.")])


def chart_data():
    data = CategoryChartData()
    data.categories = ["a", "b"]
    data.add_series("s", (1, 2))
    return data


def small_document(k, paragraphs=1):
    """A document from python-docx's template, as the timed ones: a new rdocx.Document() does not go into those
    (RdocxError "more than one default")."""
    other = docx.Document()
    for i in range(paragraphs):
        other.add_paragraph(f"inserted {k}.{i}")
    buffer = io.BytesIO()
    other.save(buffer)
    return rdocx.Document.from_bytes(buffer.getvalue())


def fragment(k):
    other = small_document(k, 2)
    story = body_story(other)
    return rdocx.DocumentFragment.from_range(other, rdocx.ContentLocation(story, "paragraph", [0]),
                                             rdocx.ContentLocation(story, "paragraph", [1]))


def with_xml(target, edit):
    """(target, its XML edited): for replace_xml."""
    def make(doc, k):
        obj = target(doc, k)
        return obj, edit(obj.xml, k)
    return make


# ---------------------------------------------------------------- setups
# A setup takes the document and returns ready(k), which makes the k-th fresh target untimed and returns the timed
# action; the targets are made one at a time, as the timing needs them.
class FromDoc:
    """A setter value read from (or made in) the document being timed."""

    def __init__(self, read):
        self.read = read


def setter(target, attr, value):
    """`target(...).attr = value` (value[k] for a list) on fresh targets."""
    def setup(doc):
        resolved = value.read(doc) if isinstance(value, FromDoc) else value
        def ready(k):
            obj, v = target(doc, k), resolved[k] if isinstance(resolved, list) else resolved
            return lambda: setattr(obj, attr, v)
        return ready
    return setup


def toggled(target, attr):
    """`target(...).attr = not target(...).attr` on fresh targets."""
    def setup(doc):
        def ready(k):
            obj = target(doc, k)
            v = not getattr(obj, attr)
            return lambda: setattr(obj, attr, v)
        return ready
    return setup


def method(target, name, *args, **kwargs):
    """`target(...).name(*args, **kwargs)` on fresh targets; a function argument is called with (doc, k)."""
    def setup(doc):
        def ready(k):
            obj = target(doc, k)
            values = [a(doc, k) if callable(a) else a for a in args]
            named = {key: v(doc, k) if callable(v) else v for key, v in kwargs.items()}
            return lambda: getattr(obj, name)(*values, **named)
        return ready
    return setup


def on_made(make, act):
    """act(doc, obj) on objects that make(doc, k) creates."""
    def setup(doc):
        def ready(k):
            obj = make(doc, k)
            return lambda: act(doc, obj)
        return ready
    return setup


def each(setup):
    """A setup written out: `setup(doc)` makes every target first and returns call(k)."""
    def eager(doc):
        call = setup(doc)
        return lambda k: lambda: call(k)
    return eager


def path(*names):
    """A target reached from another one by attribute names: path(run, "font", "color")."""
    def make(doc, k):
        obj = names[0](doc, k)
        for name in names[1:]:
            obj = getattr(obj, name)
        return obj
    return make


def per_call(*values):
    """values[k], as a method argument."""
    return lambda doc, k: values[k]


def fonts_embedded(doc, names, kinds=("regular",)):
    for name in names:
        for kind in kinds:
            doc.embed_font(name, FONT, kind, license=LICENSE)


def bookmarked(doc, k):
    """The paragraph "alpha beta k", holding the bookmark target_k."""
    doc.add_bookmark(f"target_{k}", run_range(doc, k))
    return para(doc, k)


def section_xml(xml, k):
    return re.sub(r'w:top="\d+"', f'w:top="{1500 + 10 * k}"', xml, count=1)


VARIANTS = ["first", "even", "default"]
VARIANT_PARTS = ["header", "footer", "even_page_header"]
TIMED = {
    # Document: fonts, settings, defaults
    "Document.set_font": method(document, "set_font", lambda doc, k: rdocx.FontDefinition(f"Probe {next(IDS)}")),
    "Document.embed_font": each(lambda doc: lambda k: doc.embed_font("Aileron", FONT, KINDS[k], license=LICENSE)),
    "Document.remove_embedded_font": each(lambda doc: (fonts_embedded(doc, ["Aileron"], KINDS),
                                                       lambda k: doc.remove_embedded_font("Aileron", KINDS[k]))[1]),
    "Document.remove_font": on_made(lambda doc, k: (lambda name: (doc.set_font(rdocx.FontDefinition(name)), name)[1])(
        f"Gone {k}"), lambda doc, name: doc.remove_font(name)),
    "Document.embed_true_type_fonts": setter(document, "embed_true_type_fonts", [True, False, True]),
    "Document.save_subset_fonts": setter(document, "save_subset_fonts", [True, False, True]),
    "Document.embed_system_fonts": setter(document, "embed_system_fonts", [True, False, True]),
    "Document.default_font_name": setter(document, "default_font_name", ["Georgia", "Verdana", "Tahoma"]),
    "Document.default_font_size": setter(document, "default_font_size", [Pt(13), Pt(14), Pt(15)]),
    "Document.update_fields_on_open": setter(document, "update_fields_on_open", [True, False, True]),
    "Document.set_default_paragraph_format": method(document, "set_default_paragraph_format",
                                                    space_after=lambda doc, k: Pt(k + 3)),
    "Document.set_page_color": method(document, "set_page_color", per_call("FFF2CC", "E2EFDA", "DDEBF7")),
    "Document.set_text_watermark": method(document, "set_text_watermark", per_call("DRAFT", "COPY", "VOID")),
    "Document.set_image_watermark": method(document, "set_image_watermark", lambda doc, k: noise_png(8, 600 + k),
                                           "mark.png", Inches(2), Inches(2)),
    "Document.set_note_policy": method(document, "set_note_policy", "footnote", start=lambda doc, k: k + 2),
    "Document.set_header": method(document, "set_header", per_call("header 0", "header 1", "header 2")),
    "Document.set_footer": method(document, "set_footer", per_call("footer 0", "footer 1", "footer 2")),
    "Settings.track_revisions": setter(path(document, "settings"), "track_revisions", [True, False, True]),
    "Settings.odd_and_even_pages_header_footer": setter(path(document, "settings"), "odd_and_even_pages_header_footer",
                                                        [True, False, True]),
    # Styles and lists
    "Document.add_style": each(lambda doc: lambda k: doc.add_style(f"Probe style {next(IDS)}", bold=True)),
    "Document.set_style": on_made(style, lambda doc, s: doc.set_style(s, bold=True)),
    "Document.remove_style": on_made(style, lambda doc, s: doc.remove_style(s)),
    "Document.set_default_style": on_made(style, lambda doc, s: doc.set_default_style(s)),
    "Document.add_numbering_definition": each(lambda doc: lambda k: numbering_definition(doc, k)),
    "Document.add_numbering_instance": on_made(numbering_definition, lambda doc, d: doc.add_numbering_instance(d)),
    "Document.link_style_to_numbering": on_made(
        lambda doc, k: (style(doc, k), doc.add_numbering_instance(numbering_definition(doc, k))),
        lambda doc, made: doc.link_style_to_numbering(made[0], made[1], 0)),
    "Document.add_bullet_list_item": each(lambda doc: lambda k: doc.add_bullet_list_item(f"item {k}")),
    "Document.add_numbered_list_item": each(lambda doc: lambda k: doc.add_numbered_list_item(f"item {k}")),
    "Document.restart_numbering": on_made(lambda doc, k: doc.add_numbered_list_item(f"item {k}"),
                                          lambda doc, p: doc.restart_numbering(p, start=3)),
    # Sections, headers, footers, notes
    "Document.update_section": method(document, "update_section", 0, margin_top=lambda doc, k: Inches(1 + k / 10)),
    "Document.insert_section": on_made(lambda doc, k: len(doc.sections), lambda doc, i: doc.insert_section(i)),
    "Document.remove_section": on_made(lambda doc, k: (doc.insert_section(len(doc.sections)), len(doc.sections) - 1)[1],
                                       lambda doc, s: doc.remove_section(s)),
    "Document.create_section_story": each(lambda doc: lambda k: doc.create_section_story(1, "footer", VARIANTS[k])),
    "Document.unlink_section_story": on_made(lambda doc, k: (doc.create_section_story(0, "footer", VARIANTS[k]),
                                                             VARIANTS[k])[1],
                                             lambda doc, variant: doc.unlink_section_story(1, "footer", variant)),
    "Document.link_section_story": on_made(lambda doc, k: doc.create_section_story(0, "footer", VARIANTS[k]),
                                           lambda doc, story: doc.link_section_story(1, "footer", "default", story)),
    "SectionHandle.replace_xml": on_made(with_xml(section, section_xml),
                                         lambda doc, made: made[0].replace_xml(made[1])),
    "SectionHandle.set_text_watermark": method(section, "set_text_watermark", per_call("DRAFT", "COPY", "VOID"),
                                               width=Inches(4), height=Inches(1), font_size=Pt(36)),
    "SectionHandle.set_image_watermark": method(section, "set_image_watermark", lambda doc, k: noise_png(8, 700 + k),
                                                "mark.png", width=Inches(2), height=Inches(2)),
    "SectionHandle.remove_watermark": on_made(lambda doc, k: (doc.section(0).set_text_watermark(
        f"DRAFT {k}", width=Inches(4), height=Inches(1), font_size=Pt(36)), doc.section(0))[1],
        lambda doc, s: s.remove_watermark()),
    "SectionHandle.set_page_border": method(section, "set_page_border", per_call("top", "left", "bottom")),
    "SectionHandle.set_note_policy": method(section, "set_note_policy", "footnote", start=lambda doc, k: k + 2),
    "SectionHandle.different_first_page_header_footer": setter(section, "different_first_page_header_footer",
                                                               [True, False, True]),
    "SectionHandle.orientation": setter(section, "orientation", ["landscape", "portrait", "landscape"]),
    "HeaderFooter.is_linked_to_previous": setter(lambda doc, k: getattr(doc.section(1), VARIANT_PARTS[k]),
                                                 "is_linked_to_previous", False),
    "HeaderFooter.add_paragraph": method(lambda doc, k: doc.section(0).header, "add_paragraph", "header text"),
    "HeaderFooter.add_table": method(lambda doc, k: doc.section(0).header, "add_table", 1, 1),
    "Document.add_footnote": on_made(para, lambda doc, p: doc.add_footnote(p, "note")),
    "Document.add_endnote": on_made(para, lambda doc, p: doc.add_endnote(p, "note")),
    "Document.remove_footnote": on_made(lambda doc, k: footnote(doc, k).id, lambda doc, i: doc.remove_footnote(i)),
    "Document.remove_endnote": on_made(lambda doc, k: endnote(doc, k).id, lambda doc, i: doc.remove_endnote(i)),
    "FootnoteEndnote.add_paragraph": method(footnote, "add_paragraph", "more"),
    "FootnoteEndnote.add_table": method(footnote, "add_table", 1, 1),
    "FootnoteEndnote.delete": method(footnote, "delete"),
    # Body content
    "Document.add_paragraph": each(lambda doc: lambda k: doc.add_paragraph("x")),
    "Document.add_heading": each(lambda doc: lambda k: doc.add_heading("x", level=1)),
    "Document.add_page_break": each(lambda doc: lambda k: doc.add_page_break()),
    "Document.add_table": each(lambda doc: lambda k: doc.add_table(2, 2)),
    "Document.insert_paragraph": each(lambda doc: lambda k: doc.insert_paragraph(10, "x")),
    "Document.insert_table": each(lambda doc: lambda k: doc.insert_table(10, 2, 2)),
    "Document.remove_content": on_made(index_of, lambda doc, i: doc.remove_content(i)),
    "Document.remove_content_range": on_made(index_of, lambda doc, i: doc.remove_content_range(i, i + 1)),
    "Document.pop_content": on_made(index_of, lambda doc, i: doc.pop_content(i)),
    "Document.insert_content": on_made(lambda doc, k: doc.pop_content(index_of(doc, k)),
                                       lambda doc, f: doc.insert_content(10, f)),
    "Document.clone_content": on_made(para, lambda doc, p: doc.clone_content(p, 10)),
    "Document.move_content": on_made(para, lambda doc, p: doc.move_content(p, 10)),
    "Document.insert_document": on_made(lambda doc, k: small_document(k), lambda doc, other: doc.insert_document(
        other, at=10)),
    "Document.import_fragment": on_made(lambda doc, k: (rdocx.ContentLocation.end(body_story(doc)), fragment(k)),
                                        lambda doc, made: doc.import_fragment(*made)),
    "Document.insert_html": on_made(para, lambda doc, p: doc.insert_html("<p>inserted</p>", after=p)),
    "Document.split_run": on_made(para, lambda doc, p: doc.split_run(p, 0, 3)),
    "Document.add_picture": each(lambda doc: lambda k: doc.add_picture(noise_png(8, 800 + k), "probe.png", Inches(1),
                                                                       Inches(1))),
    "Document.replace_image": on_made(lambda doc, k: (picture(doc, k).relationship_id, noise_png(8, 900 + k)),
                                      lambda doc, made: doc.replace_image(*made)),
    "Document.replace_image_for_story": on_made(
        lambda doc, k: (lambda rid: (body_story(doc), rid, noise_png(8, 1000 + k)))(picture(doc, k).relationship_id),
        lambda doc, made: doc.replace_image_for_story(*made)),
    "Document.set_picture_size": on_made(picture, lambda doc, pic: doc.set_picture_size(pic, Inches(2), Inches(2))),
    "Document.set_picture_sizes_for_relationship": on_made(
        lambda doc, k: picture(doc, k).relationship_id,
        lambda doc, rid: doc.set_picture_sizes_for_relationship(rid, Inches(2), Inches(2))),
    "Document.add_chart": on_made(lambda doc, k: (para(doc, k), chart_data()), lambda doc, made: doc.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, made[1], Inches(4), Inches(3), after=made[0])),
    "Document.set_chart_legend": on_made(
        lambda doc, k: doc.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, chart_data(), Inches(4), Inches(3)),
        lambda doc, p: doc.set_chart_legend(p, visible=True, position="b")),
    "Document.add_heading_hyperlink": on_made(lambda doc, k: (para(doc, k), doc.add_heading(f"Heading {k}", level=1)),
                                              lambda doc, made: doc.add_heading_hyperlink(*made, "see")),
    # Comments, bookmarks, revisions
    "Document.add_comment": on_made(run_range, lambda doc, r: doc.add_comment(r, author="Reviewer", text="c",
                                                                              date=STAMP)),
    "Document.move_comment": on_made(lambda doc, k: (comment(doc, k), story_range(doc, k)),
                                     lambda doc, made: doc.move_comment(*made)),
    "Document.reply_to": on_made(comment, lambda doc, c: doc.reply_to(c, author="Author", text="r", date=STAMP)),
    "Document.resolve_comment": on_made(comment, lambda doc, c: doc.resolve_comment(c)),
    "Document.remove_comment": on_made(comment, lambda doc, c: doc.remove_comment(c)),
    "Document.add_bookmark": on_made(lambda doc, k: (f"probe_{next(IDS)}", run_range(doc, k)),
                                     lambda doc, made: doc.add_bookmark(*made)),
    "Document.accept_revision_id": on_made(revision, lambda doc, rid: doc.accept_revision_id(rid)),
    "Document.reject_revision_id": on_made(revision, lambda doc, rid: doc.reject_revision_id(rid)),
    # Text and hyperlinks by story
    "Document.replace_text_at": on_made(lambda doc, k: item_of(doc, para(doc, k)),
                                        lambda doc, it: doc.replace_text_at(it, "alpha", "omega")),
    "Document.set_story_text": on_made(lambda doc, k: item_of(doc, para(doc, k)),
                                       lambda doc, it: doc.set_story_text(it, "omega")),
    "Document.add_hyperlink_to_story": on_made(lambda doc, k: body_story(doc),
                                               lambda doc, story: doc.add_hyperlink_to_story(story, "link", URL)),
    "Document.set_hyperlink_url": on_made(hyperlink, lambda doc, h: doc.set_hyperlink_url(h, URL + "b")),
    "Document.remove_hyperlink": on_made(hyperlink, lambda doc, h: doc.remove_hyperlink(h)),
    # Content controls
    "Document.insert_content_control": on_made(location, lambda doc, loc: doc.insert_content_control(
        loc, "block", "plain_text", tag="new")),
    "Document.set_content_control_tag": on_made(control, lambda doc, loc: doc.set_content_control_tag(loc, "renamed")),
    "Document.set_content_control_alias": on_made(control, lambda doc, loc: doc.set_content_control_alias(
        loc, "Alias")),
    "Document.set_content_control_lock": on_made(control, lambda doc, loc: doc.set_content_control_lock(
        loc, "sdtLocked")),
    "Document.set_content_control_appearance": on_made(control, lambda doc, loc: doc.set_content_control_appearance(
        loc, "hidden")),
    "Document.set_content_control_placeholder": on_made(control, lambda doc, loc: doc.set_content_control_placeholder(
        loc, True)),
    "Document.set_content_control_type": on_made(control, lambda doc, loc: doc.set_content_control_type(
        loc, rdocx.ContentControlTypePayload.plain_text())),
    "Document.remove_content_control": on_made(control, lambda doc, loc: doc.remove_content_control(loc)),
    "Document.unwrap_content_control": on_made(control, lambda doc, loc: doc.unwrap_content_control(loc)),
    "Document.move_content_control": on_made(controls, lambda doc, made: doc.move_content_control(*made)),
    "Document.clone_content_control": on_made(controls, lambda doc, made: doc.clone_content_control(*made)),
    "Document.repeat_content_control": on_made(lambda doc, k: found(doc, tagged("rows")),
                                               lambda doc, loc: doc.repeat_content_control(loc)),
    "Document.add_custom_xml_store": each(lambda doc: lambda k: doc.add_custom_xml_store(
        f"<data><v>{k}</v></data>".encode())),
    "Document.set_content_control_binding": on_made(
        lambda doc, k: (lambda store: (control(doc, k), store))(doc.add_custom_xml_store(b"<data><v>x</v></data>")),
        lambda doc, made: doc.set_content_control_binding(made[0], made[1], "/data/v")),
    # Paragraphs
    "Paragraph.insert_paragraph_before": method(para, "insert_paragraph_before", "x"),
    "Paragraph.replace_xml": on_made(with_xml(para, lambda xml, k: xml.replace("alpha", "omega")),
                                     lambda doc, made: made[0].replace_xml(made[1])),
    "Paragraph.add_equation": method(para, "add_equation", "x^2", format="latex"),
    "Paragraph.replace_text": method(para, "replace_text", "alpha", "omega"),
    "Paragraph.text": setter(para, "text", "x"),
    "Paragraph.alignment": setter(para, "alignment", WD_ALIGN_PARAGRAPH.CENTER),
    "Paragraph.add_run": method(para, "add_run", "x"),
    "Paragraph.add_hyperlink": method(para, "add_hyperlink", "x", URL),
    "Paragraph.add_internal_hyperlink": method(bookmarked, "add_internal_hyperlink", "x", lambda doc, k: f"target_{k}"),
    "Paragraph.style": setter(para, "style", "Title"),
    "Paragraph.numbering": setter(para, "numbering", FromDoc(lambda doc: (
        doc.add_numbering_instance(numbering_definition(doc, 0)), 0))),
    # Runs
    "Run.bold": setter(run, "bold", True),
    "Run.italic": setter(run, "italic", True),
    "Run.underline": setter(run, "underline", True),
    "Run.replace_xml": on_made(with_xml(run, lambda xml, k: xml.replace("alpha", "omega")),
                               lambda doc, made: made[0].replace_xml(made[1])),
    "Run.text": setter(run, "text", "x"),
    "Run.remove": method(lambda doc, k: para(doc, k).runs[1], "remove"),
    "Run.style_id": setter(run, "style_id", FromDoc(lambda doc: style(doc, 0, "character"))),
    "Run.add_tab": method(run, "add_tab"),
    "Run.add_break": method(run, "add_break"),
    "Run.add_picture": method(run, "add_picture", lambda doc, k: noise_png(8, 1100 + k), "probe.png", Inches(1),
                              Inches(1)),
    "Run.add_field": method(run, "add_field", "PAGE", "1"),
    "ColorFormat.rgb": setter(path(run, "font", "color"), "rgb", RGB),
    "ColorFormat.theme_color": setter(path(run, "font", "color"), "theme_color", "accent2"),
    "Font.highlight_color": setter(path(run, "font"), "highlight_color", rdocx.enum.text.WD_COLOR_INDEX.YELLOW),
    "Font.superscript": setter(path(run, "font"), "superscript", True),
    "Font.subscript": setter(path(run, "font"), "subscript", True),
    "Font.all_caps": setter(path(run, "font"), "all_caps", True),
    "Font.small_caps": setter(path(run, "font"), "small_caps", True),
    "Font.double_strike": setter(path(run, "font"), "double_strike", True),
    "Font.hidden": setter(path(run, "font"), "hidden", True),
    "Font.rtl": setter(path(run, "font"), "rtl", True),
    "Font.language": setter(path(run, "font"), "language", "fr-FR"),
    "Font.language_east_asia": setter(path(run, "font"), "language_east_asia", "ja-JP"),
    "Font.language_bidi": setter(path(run, "font"), "language_bidi", "ar-SA"),
    "Font.name_east_asia": setter(path(run, "font"), "name_east_asia", "MS Gothic"),
    "Font.name_complex_script": setter(path(run, "font"), "name_complex_script", "Arial"),
    "Font.character_spacing": setter(path(run, "font"), "character_spacing", Pt(1)),
    "Font.name": setter(path(run, "font"), "name", "Georgia"),
    "Font.size": setter(path(run, "font"), "size", Pt(14)),
    "Font.color": setter(path(run, "font"), "color", RGB),
    "Font.bold": setter(path(run, "font"), "bold", True),
    "Font.italic": setter(path(run, "font"), "italic", True),
    "Font.underline": setter(path(run, "font"), "underline", True),
    "Font.strike": setter(path(run, "font"), "strike", True),
    "Font.highlight": setter(path(run, "font"), "highlight", "yellow"),
    "Font.shading": setter(path(run, "font"), "shading", RGB),
    # Paragraph format and tab stops
    "ParagraphFormat.set_border": method(path(para, "paragraph_format"), "set_border", "bottom", "single", size=8,
                                         color="000000"),
    "ParagraphFormat.clear_border": method(bordered, "clear_border", "bottom"),
    "ParagraphFormat.clear_borders": method(bordered, "clear_borders"),
    "ParagraphFormat.shading": setter(path(para, "paragraph_format"), "shading", RGB),
    "ParagraphFormat.outline_level": setter(path(para, "paragraph_format"), "outline_level", 1),
    "ParagraphFormat.rtl": setter(path(para, "paragraph_format"), "rtl", True),
    "ParagraphFormat.alignment": setter(path(para, "paragraph_format"), "alignment", WD_ALIGN_PARAGRAPH.RIGHT),
    "ParagraphFormat.space_before": setter(path(para, "paragraph_format"), "space_before", Pt(6)),
    "ParagraphFormat.space_after": setter(path(para, "paragraph_format"), "space_after", Pt(6)),
    "ParagraphFormat.left_indent": setter(path(para, "paragraph_format"), "left_indent", Inches(0.5)),
    "ParagraphFormat.right_indent": setter(path(para, "paragraph_format"), "right_indent", Inches(0.5)),
    "ParagraphFormat.first_line_indent": setter(path(para, "paragraph_format"), "first_line_indent", Inches(0.2)),
    "ParagraphFormat.line_spacing": setter(path(para, "paragraph_format"), "line_spacing", 1.5),
    "ParagraphFormat.keep_with_next": setter(path(para, "paragraph_format"), "keep_with_next", True),
    "ParagraphFormat.keep_together": setter(path(para, "paragraph_format"), "keep_together", True),
    "ParagraphFormat.page_break_before": setter(path(para, "paragraph_format"), "page_break_before", True),
    "ParagraphFormat.widow_control": setter(path(para, "paragraph_format"), "widow_control", False),
    "TabStops.add_tab_stop": method(path(para, "paragraph_format", "tab_stops"), "add_tab_stop", Inches(1)),
    "TabStops.__delitem__": method(tab_stops, "__delitem__", 0),
    "TabStops.clear_all": method(tab_stops, "clear_all"),
    # Tables
    "Table.replace_xml": on_made(with_xml(table, lambda xml, k: xml.replace("alpha", "omega")),
                                 lambda doc, made: made[0].replace_xml(made[1])),
    "Table.style": setter(table, "style", "TableGrid"),
    "Table.indent": setter(table, "indent", Inches(0.2)),
    "Table.alignment": setter(table, "alignment", WD_TABLE_ALIGNMENT.CENTER),
    "Table.width": setter(table, "width", Inches(5)),
    "Table.set_borders": method(table, "set_borders", "single", size=8, color="000000"),
    "Table.set_border": method(table, "set_border", "top", "single", size=8, color="000000"),
    "Table.set_cell_margins": method(table, "set_cell_margins", top=Pt(2), right=Pt(2), bottom=Pt(2), left=Pt(2)),
    "Table.grid_widths": setter(table, "grid_widths", [[Inches(2), Inches(3)]] * CALLS),
    "Table.set_column_width": method(table, "set_column_width", 0, Inches(2)),
    "Table.set_cell_grid_span": method(table, "set_cell_grid_span", 0, 0, 2),
    "Table.set_cell_vertical_merge": method(table, "set_cell_vertical_merge", 0, 0, "restart"),
    "Table.add_row": method(table, "add_row"),
    "Table.insert_column": method(table, "insert_column", 0),
    "Table.delete_column": method(table, "delete_column", 0),
    "Table.autofit": toggled(table, "autofit"),
    "Table.first_row": toggled(table, "first_row"),
    "Table.last_row": toggled(table, "last_row"),
    "Table.first_column": toggled(table, "first_column"),
    "Table.last_column": toggled(table, "last_column"),
    "Table.banded_rows": toggled(table, "banded_rows"),
    "Table.banded_columns": toggled(table, "banded_columns"),
    "Table.clone_row": method(table, "clone_row", 0),
    "Table.remove_row": method(table, "remove_row", 0),
    "Row.height": setter(row, "height", Inches(0.5)),
    "Row.height_rule": setter(lambda doc, k: (lambda r: (setattr(r, "height", Inches(0.5)), r)[1])(row(doc, k)),
                              "height_rule", WD_ROW_HEIGHT_RULE.EXACTLY),
    "Row.cant_split": setter(row, "cant_split", True),
    "Row.is_header": setter(row, "is_header", True),
    "Cell.replace_xml": on_made(with_xml(cell, lambda xml, k: xml.replace("alpha", "omega")),
                                lambda doc, made: made[0].replace_xml(made[1])),
    "Cell.add_picture": method(cell, "add_picture", lambda doc, k: noise_png(8, 1200 + k), "probe.png", Inches(1),
                               Inches(1)),
    "Cell.replace_text": method(cell, "replace_text", "alpha", "omega"),
    "Cell.text": setter(cell, "text", "x"),
    "Cell.add_table": method(cell, "add_table", 1, 1),
    "Cell.merge": on_made(lambda doc, k: (lambda t: (t.cell(0, 0), t.cell(0, 1)))(table(doc, k)),
                          lambda doc, cells: cells[0].merge(cells[1])),
    "Cell.split": method(merged, "split"),
    "Cell.add_paragraph": method(cell, "add_paragraph", "x"),
    "Cell.width": setter(cell, "width", Inches(2)),
    "Cell.vertical_alignment": setter(cell, "vertical_alignment", WD_CELL_VERTICAL_ALIGNMENT.CENTER),
    "Cell.shading": setter(cell, "shading", RGB),
    "Cell.set_border": method(cell, "set_border", "top", "single", size=8, color="000000"),
    "Cell.set_margins": method(cell, "set_margins", top=Pt(2), right=Pt(2), bottom=Pt(2), left=Pt(2)),
    # Document properties
    "CustomProperties.__setitem__": each(lambda doc: lambda k: doc.custom_properties.__setitem__(f"key{k}", "value")),
    "CustomProperties.set": each(lambda doc: lambda k: doc.custom_properties.set(f"key{k}", "value")),
    "CustomProperties.__delitem__": each(lambda doc: ([doc.custom_properties.set(f"gone{k}", "v")
                                                       for k in range(CALLS)],
                                                      lambda k: doc.custom_properties.__delitem__(f"gone{k}"))[1]),
    "CustomProperties.delete": each(lambda doc: ([doc.custom_properties.set(f"gone{k}", "v") for k in range(CALLS)],
                                                 lambda k: doc.custom_properties.delete(f"gone{k}"))[1]),
}
DATES = [datetime.datetime(2026, 10, k + 1, 12, tzinfo=datetime.timezone.utc) for k in range(CALLS)]
CORE = {"created": DATES, "modified": DATES, "last_printed": DATES, "revision": [k + 2 for k in range(CALLS)]}
for name in ("author", "category", "comments", "content_status", "created", "identifier", "keywords", "language",
             "last_modified_by", "last_printed", "modified", "revision", "subject", "title", "version"):
    TIMED[f"CoreProperties.{name}"] = setter(path(document, "core_properties"), name,
                                             CORE.get(name, [f"{name} {k}" for k in range(CALLS)]))
for name in ("page_width", "page_height", "margin_top", "margin_right", "margin_bottom", "margin_left", "top_margin",
             "right_margin", "bottom_margin", "left_margin", "gutter", "header_distance", "footer_distance"):
    TIMED[f"SectionHandle.{name}"] = setter(section, name, [Inches(8 + k / 10) if name.startswith("page") else
                                                            Inches(0.5 + k / 10) for k in range(CALLS)])
for name in ("company", "manager", "template", "application", "application_version", "presentation_format"):
    TIMED[f"AppProperties.{name}"] = setter(path(document, "app_properties"), name,
                                            [f"{name} {k}" if name != "application_version" else f"16.{k}000"
                                             for k in range(CALLS)])

# Reads that tensorbee/rdocx#407 names: their cost grows with the body too.
READS_TIMED = {
    "Paragraph.xml": on_made(para, lambda doc, p: p.xml),
    "ParagraphFormat.line_spacing (read)": on_made(path(para, "paragraph_format"), lambda doc, fmt: fmt.line_spacing),
}

# Calls that search or work on every story by design: their cost grows with the document in python-docx too.
WHOLE_DOCUMENT = {
    "Document.sign": "signs every part of the package",
    "Document.compare": "compares every story",
    "Document.accept_all": "every revision", "Document.reject_all": "every revision",
    "Document.accept_revisions_by_author": "searches every revision",
    "Document.reject_revisions_by_author": "searches every revision",
    "Document.accept_revisions_in_date_range": "searches every revision",
    "Document.reject_revisions_in_date_range": "searches every revision",
    "Document.try_replace_text": "searches every story", "Document.replace_all": "searches every story",
    "Document.replace_all_regex": "searches every story",
    "Document.replace_all_regex_expected": "searches every story",
    "Document.add_comment_on_text": "searches the text", "Document.move_comment_to_text": "searches the text",
    "Document.bind_content_controls": "searches every content control by tag",
    "Document.fill_content_control_tags": "searches every content control by tag",
    "Document.set_content_control_value_by_tag": "searches every content control by tag",
    "Document.set_content_control_value_by_alias": "searches every content control by alias",
    "Document.set_legacy_form_field_value": "searches every story for the field's name",
    "Document.update_fields": "every field", "Document.update_all_fields": "every field",
    "Document.update_page_fields": "lays out every page",
    "Document.update_layout_backed_fields": "lays out every page",
    "Document.rebuild_toc": "lists every heading", "Document.insert_toc": "lists every heading",
}
# Names that raise an error naming the rdocx way (python-docx names, tensorbee/rdocx#379, and setters of snapshots),
# and public methods that change nothing in the document.
RAISES = {"Document.add_section", "Table.add_column", "Section.left_margin",
          "Section.different_first_page_header_footer", "TabStop.alignment"}
READS = {"Document.verify_signatures", "Document.chart_legend", "Document.text", "Document.text_report",
         "Document.word_count", "Document.character_count", "Document.page_count", "Document.validate",
         "Document.section", "Document.save", "Document.image_data", "Document.pictures", "Document.part_names",
         "Document.part_xml", "Document.legacy_form_fields", "Document.content_controls", "Document.layout",
         "Document.layout_page", "Document.find_content_index", "Document.find_content_indices", "Paragraph.equations",
         "ParagraphFormat.border", "Table.cell", "Table.border", "Cell.border", "CustomProperties.items",
         "CustomProperties.get", "FragmentConflictPolicy.with_style_reuse",
         "FragmentConflictPolicy.with_numbering_reuse", "FragmentConflictPolicy.with_related_part_reuse"}


def mutating_surface():
    """Class.member for every public method, property setter, __setitem__ and __delitem__ of the rdocx stubs, but the
    outputs (render*, to_*), READS and RAISES."""
    tree = ast.parse((Path(rdocx.__file__).parent / "_rdocx.pyi").read_text())
    surface = set()
    for node in tree.body:
        if not isinstance(node, ast.ClassDef):
            continue
        for f in node.body:
            if not isinstance(f, ast.FunctionDef):
                continue
            decorators = [ast.unparse(d) for d in f.decorator_list]
            setter_ = any(d.endswith(".setter") for d in decorators)
            if f.name in ("__setitem__", "__delitem__") or setter_ or (
                    not f.name.startswith("_") and not set(decorators) & {"property", "staticmethod", "classmethod"}
                    and not f.name.startswith(("render", "to_"))):
                surface.add(f"{node.name}.{f.name}")
    return surface - READS - RAISES


def costs_ms(paths, setup):
    """The fastest of CALLS calls on fresh targets in each of the documents opened from `paths`, or of two calls once
    each one takes over 50 ms. The documents take turns, so that a change of the machine's load hits all of them."""
    readies = [setup(rdocx.Document.open(path)) for path in paths]
    best = [float("inf")] * len(paths)
    for k in range(CALLS):
        for i, ready in enumerate(readies):
            action = ready(k)
            start = time.perf_counter()
            action()
            best[i] = min(best[i], time.perf_counter() - start)
        if k and min(best) > 0.05:
            break
    return [b * 1000 for b in best]


@pytest.fixture(scope="module")
def documents(tmp_path_factory):
    folder = tmp_path_factory.mktemp("scaling")
    return (edit_cost_docx(folder / "n.docx", N, picture_side=100),
            edit_cost_docx(folder / "4n.docx", 4 * N, picture_side=100))


# Calls whose cost grows with the document on the pin, in the order of TIMED: 145 of the 252 on rdocx 0.16.3, of which
# tensorbee/rdocx#415 and #417 to #420 fixed 70, Paragraph.xml included; `--runxfail` shows the numbers. An XPASS
# means a build fixed that call: remove it from here once it XPASSes on every CI platform, and the gap's entry once
# this set is empty. The passing calls are regression guards: keep them when the pin moves to an upstream release.
GROWING = {
    "Document.set_page_color", "Document.set_header", "Document.set_footer", "Document.add_bullet_list_item",
    "Document.add_numbered_list_item", "Document.restart_numbering", "Document.insert_section",
    "Document.remove_section", "Document.create_section_story", "Document.unlink_section_story",
    "Document.link_section_story", "SectionHandle.replace_xml", "SectionHandle.set_page_border",
    "SectionHandle.set_note_policy", "HeaderFooter.is_linked_to_previous", "Document.add_footnote",
    "Document.add_endnote", "Document.remove_footnote", "Document.remove_endnote", "FootnoteEndnote.add_paragraph",
    "FootnoteEndnote.add_table", "FootnoteEndnote.delete", "Document.pop_content", "Document.insert_content",
    "Document.clone_content", "Document.move_content", "Document.insert_document", "Document.import_fragment",
    "Document.insert_html", "Document.add_picture", "Document.set_picture_size",
    "Document.set_picture_sizes_for_relationship", "Document.set_chart_legend", "Document.add_heading_hyperlink",
    "Document.move_comment", "Document.remove_comment", "Document.accept_revision_id", "Document.reject_revision_id",
    "Document.replace_text_at", "Document.set_story_text", "Document.add_hyperlink_to_story",
    "Document.set_hyperlink_url", "Document.remove_hyperlink", "Document.insert_content_control",
    "Document.set_content_control_tag", "Document.set_content_control_alias", "Document.set_content_control_lock",
    "Document.set_content_control_appearance", "Document.set_content_control_placeholder",
    "Document.set_content_control_type", "Document.remove_content_control", "Document.unwrap_content_control",
    "Document.move_content_control", "Document.clone_content_control", "Document.repeat_content_control",
    "Document.add_custom_xml_store", "Document.set_content_control_binding", "Paragraph.replace_xml",
    "Paragraph.replace_text", "Paragraph.add_internal_hyperlink", "Run.replace_xml", "Run.text", "Run.add_picture",
    "Table.replace_xml", "Table.add_row", "Table.insert_column", "Table.delete_column", "Table.clone_row",
    "Table.remove_row", "Cell.replace_xml", "Cell.add_picture", "Cell.replace_text", "Cell.text", "Cell.merge",
    "Cell.split",
}


@pytest.mark.parametrize("name", [pytest.param(name, marks=pytest.mark.gap(GROWS)) if name in GROWING else name
                                  for name in TIMED | READS_TIMED])
def test_call_cost_does_not_grow_with_the_document(documents, name):
    small, large = costs_ms(documents, (TIMED | READS_TIMED)[name])
    bound = (FIXED if name in GROWING else RATIO) * small + FLOOR_MS
    assert large < bound, f"{name}: {large:.1f} ms at {4 * N} paragraphs, {small:.1f} ms at {N}"


def test_every_mutating_call_is_timed():
    assert sorted(GROWING - set(TIMED) - set(READS_TIMED)) == []
    timed = {name.split(" (")[0] for name in TIMED}
    assert sorted(mutating_surface() - timed - set(WHOLE_DOCUMENT)) == []
    assert sorted((timed | set(WHOLE_DOCUMENT)) - mutating_surface()) == []
