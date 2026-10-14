"""docx, fields, table of contents, layout, rendering and CLI robustness."""
import io
import re
import subprocess

import docx
import pytest
import rdocx
from docx.shared import Pt
from PIL import Image

from conftest import FREE_FONT_LICENSE, MONOSPACE_FONTS, STAMP, VARIABLE_FONTS, digest, font_dir, poppler, run


def test_toc_rebuild_adds_no_empty_paragraph(report_docx):
    """The fixture's TOC field begins in its first entry and ends in its last, as Google Docs and Word write it;
    a rebuild with as many entries keeps as many paragraphs, none of them empty."""
    doc = rdocx.Document.open(report_docx)
    before = len(doc.paragraphs)
    assert doc.rebuild_toc().entry_count == 21
    toc = [p.text for p in doc.paragraphs if (p.style or "").startswith("TOC") and p.style != "TOCHeading"]
    assert (len(doc.paragraphs), toc.count("")) == (before, 0)


def redline(path):
    """A word-level redline: "OLDWORD" deleted, "NEWWORD" inserted."""
    a, b = rdocx.Document(), rdocx.Document()
    a.add_paragraph("Keep OLDWORD here.")
    b.add_paragraph("Keep NEWWORD here.")
    a.compare(b, "Reviewer", STAMP, granularity="word")
    a.save(path)
    return a


# ---------------------------------------------------------------- caller fonts
def font_docx(path, family, bold=False, text="Sample text in a font from a folder"):
    d = docx.Document()
    r = d.add_paragraph().add_run(text)
    r.font.name, r.font.bold, r.font.size = family, bold, Pt(24)
    d.save(path)
    return path


def test_font_dir_reaches_the_images(rdocx_cli, tmp_path):
    fonts = font_dir(tmp_path, MONOSPACE_FONTS, "Monofamily")
    src = font_docx(tmp_path / "f.docx", "Monofamily")
    run([rdocx_cli, "convert", src, "--to", "png", "-o", tmp_path / "plain.png"], check=True)
    run([rdocx_cli, "convert", src, "--to", "png", "--font-dir", fonts, "-o", tmp_path / "fonts.png"], check=True)
    assert digest((tmp_path / "fonts.png").read_bytes()) != digest((tmp_path / "plain.png").read_bytes())
    res = run([rdocx_cli, "render", src, "--font-dir", fonts, "-o", tmp_path / "pages"])
    assert res.returncode == 0, res.stderr
    doc = rdocx.Document.open(src)
    assert digest(doc.render_page_to_png(0, 150.0, font_dir=fonts)) != digest(doc.render_page_to_png(0, 150.0))
    assert digest(doc.render_pages(dpi=100, pages=[0], font_dir=fonts)[0]) != digest(doc.render_pages(dpi=100, pages=[0])[0])
    assert run([rdocx_cli, "layout", "--json", src, "--font-dir", fonts]).returncode == 0
    assert run([rdocx_cli, "convert", src, "--to", "md", "--font-dir", fonts, "-o", tmp_path / "f.md"]).returncode != 0
    missing = run([rdocx_cli, "convert", src, "--to", "pdf", "--font-dir", tmp_path / "nowhere", "-o", tmp_path / "m.pdf"])
    assert missing.returncode != 0 and not (tmp_path / "m.pdf").exists()


def test_a_font_dir_file_named_after_the_font_wins_over_an_embedded_face(rdocx_cli, tmp_path, free_font):
    """A font_dir / --font-dir file counts as the family its file name gives, and wins over the embedded face of
    that font even when its own family differs: the folder holds a monospace face saved as Aileron.ttf."""
    doc = rdocx.Document()
    doc.add_paragraph("Embedded face text").runs[0].font.name = "Aileron"
    doc.embed_font("Aileron", free_font, license=rdocx.FontEmbeddingLicense(authorized=True, identity=FREE_FONT_LICENSE))
    doc.save(tmp_path / "f.docx")
    plain = rdocx.Document.open(tmp_path / "f.docx")
    plain.remove_embedded_font("Aileron", "regular")
    plain.save(tmp_path / "plain.docx")
    fonts = font_dir(tmp_path, MONOSPACE_FONTS, "Aileron")
    embedded = rdocx.Document.open(tmp_path / "f.docx")
    assert digest(embedded.render_page_to_png(0, 50.0)) != digest(plain.render_page_to_png(0, 50.0))
    assert digest(plain.render_page_to_png(0, 50.0, font_dir=fonts)) != digest(plain.render_page_to_png(0, 50.0))
    assert digest(embedded.render_page_to_png(0, 50.0, font_dir=fonts)) == digest(
        plain.render_page_to_png(0, 50.0, font_dir=fonts))
    for name in ("f", "plain"):
        run([rdocx_cli, "convert", tmp_path / f"{name}.docx", "--to", "png", "--dpi", "50", "--font-dir", fonts, "-o",
             tmp_path / f"{name}.png"], check=True)
    assert digest((tmp_path / "f.png").read_bytes()) == digest((tmp_path / "plain.png").read_bytes())


def ink(png):
    with Image.open(io.BytesIO(png)) as im:
        return sum(1 for v in im.convert("L").getdata() if v < 128)


def test_a_variable_font_from_a_folder_renders_bold(tmp_path):
    pdftoppm = poppler("pdftoppm")
    if not pdftoppm:
        pytest.skip("Poppler's pdftoppm is not installed")
    fonts = font_dir(tmp_path, VARIABLE_FONTS, "Varfamily")
    weights = []
    for bold in (False, True):
        pdf = tmp_path / f"{bold}.pdf"
        pdf.write_bytes(rdocx.Document.open(font_docx(tmp_path / f"{bold}.docx", "Varfamily", bold)).to_pdf(font_dir=fonts))
        subprocess.run([pdftoppm, "-png", "-r", "100", "-singlefile", pdf, tmp_path / f"{bold}"], check=True)
        weights.append(ink((tmp_path / f"{bold}.png").read_bytes()))
    assert weights[1] > 1.15 * weights[0], weights                     # a bold run has visibly heavier strokes


def test_a_font_without_a_bold_face_gets_a_synthetic_bold(tmp_path):
    """A bold run in a family that has only a regular face draws heavier, as Word does (no Poppler needed)."""
    fonts = font_dir(tmp_path, MONOSPACE_FONTS, "Monofamily")
    weights = [ink(rdocx.Document.open(font_docx(tmp_path / f"{bold}.docx", "Monofamily", bold)).render_page_to_png(
        0, 100.0, font_dir=fonts)) for bold in (False, True)]
    assert weights[1] > 1.15 * weights[0], weights


# ---------------------------------------------------------------- tracked view in colour, raw XML in headers (#373, through #375)
def blue_and_red(png):
    """How many pixels of the image are clearly blue, and how many clearly red."""
    pixels = list(Image.open(io.BytesIO(png)).convert("RGB").getdata())
    return (sum(1 for r, g, b in pixels if b > 150 and r < 90 and g < 90),
            sum(1 for r, g, b in pixels if r > 150 and g < 90 and b < 90))


def test_a_tracked_render_can_be_asked_in_colour(rdocx_cli, tmp_path):
    """As legal comparison tools show a redline: insertions blue, deletions red."""
    doc = redline(tmp_path / "r.docx")
    assert blue_and_red(doc.render_pages(dpi=100, pages=[0], revision_view="tracked")[0]) == (0, 0)
    blue, red = blue_and_red(doc.render_pages(dpi=100, pages=[0], revision_view="tracked", revision_style="legal")[0])
    assert blue > 0 and red > 0
    run([rdocx_cli, "convert", tmp_path / "r.docx", "--to", "png", "--dpi", "100", "--revision-view", "tracked",
         "--revision-style", "legal", "-o", tmp_path / "r.png"], check=True)
    blue, red = blue_and_red((tmp_path / "r.png").read_bytes())
    assert blue > 0 and red > 0


def test_raw_xml_reaches_header_paragraphs():
    """Needed to colour the revisions of a header or footer by hand."""
    doc = rdocx.Document()
    doc.add_paragraph("Body.")
    header = doc.sections[0].header.paragraphs[0]
    header.text = "Header text."
    header.replace_xml(header.xml.replace(b"Header text.", b"New header."))
    assert doc.sections[0].header.paragraphs[0].text == "New header."


def test_replace_xml_takes_the_svg_namespace_declared_on_the_paragraph(tmp_path):
    """A picture paragraph with an SVG twin, its asvg namespace declared on w:p rather than on svgBlip: the same
    XML infoset, which xml.etree writes that way."""
    image = io.BytesIO()
    Image.new("RGB", (4, 4), (255, 0, 0)).save(image, "PNG")
    (tmp_path / "dot.png").write_bytes(image.getvalue())
    doc = rdocx.Document()
    doc.add_picture(str(tmp_path / "dot.png"), width=rdocx.Inches(0.5))
    doc.add_picture(str(tmp_path / "dot.png"), width=rdocx.Inches(0.5))  # its image part stands in for the SVG one
    first, second = (p.xml.decode() for p in doc.paragraphs[:2])
    rid = re.search(r'r:embed="([^"]+)"', second).group(1)
    ns = 'xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main"'
    ext = ('<a:extLst><a:ext uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}"><asvg:svgBlip r:embed="%s"/></a:ext>'
           '</a:extLst>' % rid)
    xml = (re.sub(r"(<a:blip\b[^>]*?)/>", lambda m: m.group(1) + ">" + ext + "</a:blip>", first, count=1)
           if "</a:blip>" not in first else first.replace("</a:blip>", ext + "</a:blip>", 1))
    doc.paragraphs[0].replace_xml(re.sub(r"^<w:p\b", "<w:p " + ns, xml, count=1))
    assert b"svgBlip" in doc.paragraphs[0].xml
